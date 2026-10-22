"""Build the Week 6 design review inside the course's own Word template.

Since 26 September 2026 the submission follows the course template
(`References/Week_6_Design_Review_Template(1).docx`) exactly: its title,
Submission Detail table, Purpose and How-to boxes, Team Role table,
Submission Instructions, and sections 1-9 with their instruction lines,
tables, picture boxes and explanation boxes, in the template's own styles
(Arial, US Letter). There is no page limit (team decision of 26 September
2026). The template's toaster-oven examples are replaced with the team's
content; section 6 repeats the template's statechart block for all five
subsystems ("Add more sub-systems if needed").

The content itself is not duplicated here: it is taken from the report
function of Tools/build_week6_submission.py (requirements, pin table,
tasks, statechart fields, assumptions, AI evidence, consistency checks and
the software identity), so both builds always say the same thing.

Usage
  python Tools/build_week6_template_report.py [--entries FILE.json]
         [--out DOCX] [--check-pdf PDF]

  --entries    team entries (default: local/week6_entries.json when it
               exists, which adds the names, student IDs, group ID and
               repository URL, else Part1_Week6_Report/report_entries.json).
               Any identity field left blank stays a grey Word field.
  --out        output DOCX (default Part1_Week6_Report/INF2004_Week6_Design_Review.docx)
  --check-pdf  check a rendered PDF of the output: every template heading,
               no template example or "Enter ..." text left, eight figures,
               text at least 8.5 pt (the template's table size; spans under
               8.4 pt fail) inside the page; prints the page count.

The output is a private submission file; a named build may be committed to
this private repository only (team decision of 25 September 2026).
"""
import argparse
import copy
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_week6_submission as b  # noqa: E402  (content source)

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "References" / "Week_6_Design_Review_Template(1).docx"
OUT = ROOT / "Part1_Week6_Report" / "INF2004_Week6_Design_Review.docx"
TITLE = "Design Review: Autonomous Robot Car (INF2004)"
CAPTION_PT = 9
# Template text that must not survive in the filled report.
LEFTOVERS = ("oven", "toaster", "Enter explanation here", "Enter justification here",
             "Example:", "Example Requirement", "Click to enter B")
HEADINGS = ["Team Role and Subsystem Ownership", "Submission Instructions",
            "1. Requirements", "1.1 Functional Requirements",
            "1.2 Non-Functional Requirements", "2. Hardware Integration Pin Layout",
            "3. RTOS Task Table", "4. Task Interaction Diagram",
            "5. System-Level Statechart", "6. Subsystem Statecharts",
            "6.1 Subsystem Statechart 1", "6.2 Subsystem Statechart 2",
            "6.3 Subsystem Statechart 3", "6.4 Subsystem Statechart 4",
            "6.5 Subsystem Statechart 5", "7. Assumptions",
            "8. AI-Supported Design Evidence", "8.1 Pre-Submission AI Ideation",
            "8.2 Team-Reviewed Architecture Freeze", "8.3 Final AI-Vetted Gap Check",
            "9. Design Consistency"]


# ----------------------------------------------------------- content capture --
def capture(final=True):
    """Run the content generator with its layout helpers wrapped by
    recorders (the compact document it builds is discarded); return the
    calls, each tagged with the heading it follows."""
    calls = []
    section = [""]
    saved = {n: getattr(b, n) for n in ("h", "table", "cap", "p", "picture", "subsystem")}

    def h(d, text, level=1, mark=True):
        section[0] = text
        calls.append(("h", text, level))
        return saved["h"](d, text, level, mark)

    def table(d, headers, rows, widths, *a, **k):
        calls.append(("table", section[0], headers, rows))
        return saved["table"](d, headers, rows, widths, *a, **k)

    def cap(d, text, *a, **k):
        calls.append(("cap", section[0], text))
        return saved["cap"](d, text, *a, **k)

    def p(d, text, style=None, *a, **k):
        calls.append(("p", section[0], text, style))
        return saved["p"](d, text, style, *a, **k)

    def picture(d, name, caption):
        calls.append(("picture", section[0], name, caption))
        return saved["picture"](d, name, caption)

    def subsystem(d, heading, name, caption, text):
        section[0] = heading
        calls.append(("subsystem", heading, name, caption, text))
        return saved["subsystem"](d, heading, name, caption, text)

    try:
        for n, f in (("h", h), ("table", table), ("cap", cap), ("p", p),
                     ("picture", picture), ("subsystem", subsystem)):
            setattr(b, n, f)
        b.report(final=final)
    finally:
        for n, f in saved.items():
            setattr(b, n, f)
    return calls


def find(calls, kind, section=None, start=None):
    for c in calls:
        if c[0] == kind and (section is None or c[1] == section) \
                and (start is None or str(c[2]).startswith(start)):
            return c
    raise SystemExit(f"content not found: {kind} {section!r} {start!r}")


def text_of(part):
    """Plain text of a cell value (str or list of str / Field)."""
    if isinstance(part, (list, tuple)):
        return "".join(text_of(x) for x in part)
    if isinstance(part, b.Field):
        return part.value
    return str(part)


def strip_bold_lead(text):
    """'**Main paths:** rest' -> 'rest'."""
    return re.sub(r"^\*\*[^*]+\*\*\s*", "", text)


# ------------------------------------------------------------- docx editing --
W_P, W_R, W_TBL, W_TR = qn("w:p"), qn("w:r"), qn("w:tbl"), qn("w:tr")
_SDT = [900]


def body_items(d):
    return list(d.element.body.iterchildren())


def para_text(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def heading_el(d, text):
    for el in body_items(d):
        if el.tag == W_P and para_text(el).strip() == text:
            return el
    raise SystemExit(f"template heading not found: {text!r}")


def next_tables(el, n):
    """The next n tables after body element el (stopping at nothing)."""
    out, cur = [], el.getnext()
    while cur is not None and len(out) < n:
        if cur.tag == W_TBL:
            out.append(cur)
        cur = cur.getnext()
    return out


def table_obj(d, tbl_el):
    for t in d.tables:
        if t._tbl is tbl_el:
            return t
    raise SystemExit("table not found")


# Schema order of w:rPr children (CT_RPr); inserted properties keep it.
RPR_ORDER = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps",
             "strike", "dstrike", "outline", "shadow", "emboss", "imprint",
             "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing",
             "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect",
             "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang",
             "eastAsianLayout", "specVanish", "oMath"]


def base_rpr(par_el):
    """A copy of the first run's properties (font, size, bold) without
    italic or colour (the template's grey example styling)."""
    r = par_el.find(qn("w:r"))
    rpr = r.find(qn("w:rPr")) if r is not None else None
    rpr = copy.deepcopy(rpr) if rpr is not None else parse_xml(f"<w:rPr {nsdecls('w')}/>")
    for tag in ("w:i", "w:iCs", "w:color", "w:highlight"):
        for x in rpr.findall(qn(tag)):
            rpr.remove(x)
    return rpr


def rpr_set(rpr, tag, val=None):
    """Set (replace) child w:tag of rPr in schema order."""
    for x in rpr.findall(qn("w:" + tag)):
        rpr.remove(x)
    attr = f' w:val="{val}"' if val is not None else ""
    el = parse_xml(f"<w:{tag} {nsdecls('w')}{attr}/>")
    rank = RPR_ORDER.index(tag)
    for child in rpr:
        name = child.tag.split("}")[1]
        if name in RPR_ORDER and RPR_ORDER.index(name) > rank:
            child.addprevious(el)
            return
    rpr.append(el)


def add_runs(par_el, text, rpr, italic=False, size=None):
    """Append runs for text with **bold** markup."""
    for k, seg in enumerate(re.split(r"\*\*", text)):
        if not seg:
            continue
        r = parse_xml(f"<w:r {nsdecls('w')}/>")
        props = copy.deepcopy(rpr)
        if k % 2:
            rpr_set(props, "b")
        if italic:
            rpr_set(props, "i")
        if size:
            rpr_set(props, "sz", int(size * 2))
            rpr_set(props, "szCs", int(size * 2))
        r.append(props)
        t = parse_xml(f'<w:t {nsdecls("w")} xml:space="preserve"/>')
        t.text = seg
        r.append(t)
        par_el.append(r)


def field_sdt(key, rpr):
    """Plain-text Word content control showing grey 'Click to enter ...'."""
    _SDT[0] += 1
    props = "".join(_xml(child) for child in rpr)
    hint = b.PLACEHOLDER + b.FIELDS[key]
    return parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtPr><w:rPr>{props}</w:rPr>'
        f'<w:alias w:val="{b.FIELDS[key]}"/><w:tag w:val="{key}"/>'
        f'<w:id w:val="{_SDT[0]}"/><w:showingPlcHdr/><w:text/></w:sdtPr>'
        f'<w:sdtContent><w:r><w:rPr><w:rStyle w:val="PlaceholderText"/>{props}</w:rPr>'
        f'<w:t xml:space="preserve">{hint}</w:t></w:r></w:sdtContent></w:sdt>')


def _xml(el):
    from lxml import etree
    return re.sub(r' xmlns:\w+="[^"]*"', "", etree.tostring(el, encoding="unicode"))


def set_par_text(par_el, text):
    """Replace a paragraph's text, keeping its first run's formatting."""
    ts = list(par_el.iter(qn("w:t")))
    if not ts:
        add_runs(par_el, text, base_rpr(par_el))
        return
    ts[0].text = text
    for t in ts[1:]:
        t.text = ""


def fill_cell(cell, value, italic=False, size=None):
    """Replace a template cell's text with value (str with **bold** and
    blank-line paragraph breaks, or a list with Field objects), keeping the
    cell's paragraph and run formatting."""
    pars = cell._tc.findall(W_P)
    first = pars[0]
    rpr = base_rpr(first)
    for extra in pars[1:]:
        cell._tc.remove(extra)
    for child in list(first):
        if child.tag != qn("w:pPr"):
            first.remove(child)
    if isinstance(value, (list, tuple)):
        for part in value:
            if isinstance(part, b.Field) and not part.value:
                first.append(field_sdt(part.key, rpr))
            else:
                add_runs(first, text_of(part), rpr, italic, size)
        return
    chunks = str(value).split("\n\n")
    add_runs(first, chunks[0], rpr, italic, size)
    prev = first
    for chunk in chunks[1:]:
        par = copy.deepcopy(first)
        for child in list(par):
            if child.tag != qn("w:pPr"):
                par.remove(child)
        add_runs(par, chunk, rpr, italic, size)
        prev.addnext(par)
        prev = par


def set_rows(t, rows, keep_header=True):
    """Fill a template table's data rows, adding or removing rows (copies of
    the last data row) to match."""
    data = t.rows[1:] if keep_header else t.rows
    proto = copy.deepcopy(data[-1]._tr)
    while len(t.rows) - (1 if keep_header else 0) < len(rows):
        t._tbl.append(copy.deepcopy(proto))
    while len(t.rows) - (1 if keep_header else 0) > len(rows):
        t._tbl.remove(t.rows[-1]._tr)
    for row, values in zip(t.rows[1:] if keep_header else t.rows, rows):
        if len(values) < len(row.cells):          # e.g. a closing row that
            span_row(row, len(values))           # spans the whole table
        for cell, value in zip(row.cells, values):
            fill_cell(cell, value)


def span_row(row, n):
    """Merge a row's cells so it has n cells, the last one spanning the
    rest (w:gridSpan), and drop the merged cells' template text."""
    tcs = row._tr.findall(qn("w:tc"))
    keep, extra = tcs[n - 1], tcs[n:]
    span = len(tcs) - n + 1
    tcpr = keep.get_or_add_tcPr()
    for x in tcpr.findall(qn("w:gridSpan")):
        tcpr.remove(x)
    width = sum(int(tc.find(qn("w:tcPr")).find(qn("w:tcW")).get(qn("w:w")))
                for tc in tcs[n - 1:] if tc.find(qn("w:tcPr")) is not None
                and tc.find(qn("w:tcPr")).find(qn("w:tcW")) is not None)
    tcw = tcpr.find(qn("w:tcW"))
    if tcw is not None and width:
        tcw.set(qn("w:w"), str(width))
    gs = parse_xml(f'<w:gridSpan {nsdecls("w")} w:val="{span}"/>')
    (tcw.addnext(gs) if tcw is not None else tcpr.insert(0, gs))
    for tc in extra:
        row._tr.remove(tc)


def new_par_after(anchor_el, proto_par_el, text, italic=False, size=None):
    """Insert a paragraph (formatted like proto) after anchor; return it."""
    par = copy.deepcopy(proto_par_el)
    rpr = base_rpr(proto_par_el)
    for child in list(par):
        if child.tag != qn("w:pPr"):
            par.remove(child)
    add_runs(par, text, rpr, italic, size)
    anchor_el.addnext(par)
    return par


def picture_into(par_el, d, name, max_width):
    """Put figure `name` into paragraph par_el (clearing it), at its native
    size or narrower, centred."""
    from docx.text.paragraph import Paragraph
    for child in list(par_el):
        if child.tag != qn("w:pPr"):
            par_el.remove(child)
    par = Paragraph(par_el, None)
    par.alignment = 1
    wd, _ = b.fig_size(name, template=True)
    run = par.add_run()
    inline = d.part.new_pic_inline(str(b.FIG_T / (name + ".png")),
                                   Inches(min(wd, max_width)), None)
    run._r.add_drawing(inline)


def picture_box(d, tbl_el, name, caption):
    """Template picture box (1x1 table): our figure and its caption."""
    t = table_obj(d, tbl_el)
    cell = t.rows[0].cells[0]
    pars = cell._tc.findall(W_P)
    for extra in pars[1:]:
        cell._tc.remove(extra)
    width = cell_width(t) - cell_padding(cell) - 0.02   # inside the padding
    picture_into(pars[0], d, name, width)
    cap_par = copy.deepcopy(pars[0])
    for child in list(cap_par):
        if child.tag != qn("w:pPr"):
            cap_par.remove(child)
    rpr = parse_xml(f"<w:rPr {nsdecls('w')}/>")
    add_runs(cap_par, caption_markup(caption), rpr, italic=False, size=CAPTION_PT)
    pars[0].addnext(cap_par)


def caption_markup(caption):
    """'Fig. 2. Text' -> '**Fig. 2.** Text'."""
    m = re.match(r"(Fig\. \d+\.)\s*(.*)", caption, re.S)
    return f"**{m.group(1)}** {m.group(2)}" if m else caption


def cell_padding(cell):
    """Left plus right cell margin in inches (the template's picture boxes
    set 160 twips each; Word's table default is 108)."""
    mar = cell._tc.find(qn("w:tcPr") + "/" + qn("w:tcMar"))
    total = 0
    for side in ("left", "right"):
        el = mar.find(qn("w:" + side)) if mar is not None else None
        total += int(el.get(qn("w:w"))) if el is not None else 108
    return total / 1440


def cell_width(t):
    grid = t._tbl.find(qn("w:tblGrid"))
    cols = grid.findall(qn("w:gridCol")) if grid is not None else []
    if cols:
        return sum(int(c.get(qn("w:w"))) for c in cols) / 1440
    return 7.0


# ------------------------------------------------------------------- build --
def build(entries_path, out):
    entries = json.loads(Path(entries_path).read_text(encoding="utf-8"))
    for i in range(1, 6):                        # legacy bN_identity keys
        legacy = str(entries.pop(f"b{i}_identity", "")).strip()
        if legacy and not (entries.get(f"b{i}_name") or entries.get(f"b{i}_id")):
            name, _, sid = legacy.rpartition("·") if "·" in legacy else legacy.rpartition(" ")
            entries[f"b{i}_name"], entries[f"b{i}_id"] = name.strip(), sid.strip()
    b.ENTRIES = entries
    b.LONG_TEXT = True                            # no page limit: full wording
    b.ensure_figures()
    calls = capture(final=True)
    missing = [k for k in b.USED if k not in b.FIELDS and not str(entries.get(k, "")).strip()]
    if missing:
        raise SystemExit("pending entries: " + ", ".join(missing))

    d = Document(str(TEMPLATE))
    b._styles(d)                                  # PlaceholderText style

    # Title and Submission Detail (GROUP ID; public repository link added,
    # as the lecturer asked, Week 4 lecture 01:28:15-01:29:55).
    for par in d.paragraphs:
        if par.style.name == "Title":
            set_par_text(par._p, TITLE)
            break
    sub = d.tables[0]
    fill_cell(sub.rows[1].cells[1], [b.Field("group_id")])
    sub._tbl.append(copy.deepcopy(sub.rows[1]._tr))
    fill_cell(sub.rows[2].cells[0], "Public GitHub repository")
    fill_cell(sub.rows[2].cells[1], [b.Field("repo_url")])

    # Team Role and Subsystem Ownership: Student ID column holds ID (name).
    roles = find(calls, "table", "Team Role and Subsystem Ownership")[3]
    t = table_obj(d, next_tables(heading_el(d, "Team Role and Subsystem Ownership"), 1)[0])
    rows = []
    for i, (_, ident, resp) in enumerate(roles, 1):
        sid, name = b.Field(f"b{i}_id").value, b.Field(f"b{i}_name").value
        who = f"{sid} ({name})" if sid and name else [b.Field(f"b{i}_id")]
        rows.append((f"Buddy {i}", who, resp))
    set_rows(t, rows)
    set_widths(t, [0.9, 2.6])                    # names fit on one line

    # 1. Requirements: sources line after the instruction paragraph.
    h1 = heading_el(d, "1. Requirements")
    intro = h1.getnext()
    new_par_after(intro, intro, find(calls, "cap", "1. Requirements")[2])
    for head, subject in (("1.1 Functional Requirements", "robot"),
                          ("1.2 Non-Functional Requirements", "system")):
        _, _, headers, rows = find(calls, "table", head)
        t = table_obj(d, next_tables(heading_el(d, head), 1)[0])
        set_par_text(t.rows[0].cells[1]._tc.find(W_P), "Requirement")
        out_rows = []
        for r in rows:
            rid = text_of(r[0]).replace(" · ", " (") + (")" if " · " in text_of(r[0]) else "")
            req = f"The {subject} shall {text_of(r[1])}"
            req = req if req.endswith(".") else req + "."
            out_rows.append((rid, f"{req}\n\n**Verification:** {text_of(r[2])}"))
        set_rows(t, out_rows)

    # 2. Pin layout: sketch (Fig. 1) before the table, supply notes after it.
    h2 = heading_el(d, "2. Hardware Integration Pin Layout")
    intro = h2.getnext()
    _, _, name, caption = find(calls, "picture", "2. Hardware Integration Pin Layout")
    pic = copy.deepcopy(intro)
    intro.addnext(pic)
    picture_into(pic, d, name, 7.0)
    new_par_after(pic, intro, caption_markup(caption), size=CAPTION_PT)
    pins = find(calls, "table", "2. Hardware Integration Pin Layout")[3]
    t = table_obj(d, next_tables(h2, 1)[0])
    set_rows(t, [[text_of(c) for c in r] for r in pins])
    set_widths(t, [1.30, 0.62, 0.98, 0.66, 0.58, 0.72])
    new_par_after(t._tbl, intro, find(calls, "cap", "2. Hardware Integration Pin Layout")[2],
                  size=CAPTION_PT)

    # 3. RTOS task table and justification box.
    tasks = find(calls, "table", "3. RTOS Task Table")[3]
    h3 = heading_el(d, "3. RTOS Task Table")
    t_tasks, t_just = next_tables(h3, 2)
    set_rows(table_obj(d, t_tasks), [[text_of(c) for c in r] for r in tasks])
    fill_cell(table_obj(d, t_just).rows[0].cells[0],
              strip_bold_lead(find(calls, "p", "3. RTOS Task Table")[2]))
    align_left(table_obj(d, t_just).rows[0].cells[0])

    # 4 and 5: picture box, then explanation box.
    for head in ("4. Task Interaction Diagram", "5. System-Level Statechart"):
        el = heading_el(d, head)
        box, expl = next_tables(el, 2)
        _, _, name, caption = find(calls, "picture", head)
        picture_box(d, box, name, caption)
        fill_cell(table_obj(d, expl).rows[0].cells[0],
                  strip_bold_lead(find(calls, "p", head)[2]))
        align_left(table_obj(d, expl).rows[0].cells[0])

    # 6. One template block per subsystem (the template has three; copy
    # the last block for 6.4 and 6.5).
    subs = [c for c in calls if c[0] == "subsystem"]
    first = heading_el(d, "6.3 Subsystem Statechart 3")
    block, cur = [], first
    stop = heading_el(d, "7. Assumptions")
    while cur is not None and cur is not stop:
        block.append(cur)
        cur = cur.getnext()
    # Keep the block's own elements up to its explanation box; trailing
    # empty paragraphs (page breaks) stay once, before section 7.
    last_tbl = max(i for i, el in enumerate(block) if el.tag == W_TBL)
    proto = block[:last_tbl + 1]
    anchor = proto[-1]
    for n in (4, 5):
        for el in proto:
            clone = copy.deepcopy(el)
            anchor.addnext(clone)
            anchor = clone
            if el is proto[0]:
                set_par_text(clone, f"6.{n} Subsystem Statechart {n}")
    for n, (_, heading, name, caption, text) in enumerate(subs, 1):
        el = heading_el(d, f"6.{n} Subsystem Statechart {n}")
        m = re.match(r"6\.\d (.+?) · Related RTOS task\(s\): (.+)$", heading)
        subname, related = m.group(1).replace(" — ", " (") + ")", m.group(2)
        purpose = re.search(r"\*\*Purpose:\*\* (.+?) \*\*Reason", text).group(1)
        reason = re.search(r"\*\*Reason a statechart is needed:\*\* (.+?) \*\*States",
                           text).group(1)
        states = re.search(r"\*\*States, events and fault handling:\*\* (.+)$",
                           text, re.S).group(1)
        fields_t, box, expl = next_tables(el, 3)
        ft = table_obj(d, fields_t)
        for row, value in zip(ft.rows[1:], (subname, related, cap_first(purpose),
                                            cap_first(reason))):
            fill_cell(row.cells[1], value)
        picture_box(d, box, name, caption)
        fill_cell(table_obj(d, expl).rows[0].cells[0], cap_first(states))
        align_left(table_obj(d, expl).rows[0].cells[0])
        set_par_text(el, f"6.{n} Subsystem Statechart {n}: {subname}")

    # 7. Assumptions.
    rows = find(calls, "table", "7. Assumptions")[3]
    set_rows(table_obj(d, next_tables(heading_el(d, "7. Assumptions"), 1)[0]),
             [[text_of(c) for c in r] for r in rows])

    # 8.1: the template has separate Buddy and prompt columns.
    h81 = heading_el(d, "8.1 Pre-Submission AI Ideation")
    rows = find(calls, "table", "8.1 Pre-Submission AI Ideation")[3]
    out_rows = []
    for i, r in enumerate(rows, 1):
        prompt = re.sub(r"^\*\*B\d:\*\*\s*", "", text_of(r[0]))
        out_rows.append((f"Buddy {i}", prompt, text_of(r[1]), text_of(r[2]), text_of(r[3])))
    set_rows(table_obj(d, next_tables(h81, 1)[0]), out_rows)
    new_par_after(h81, heading_el(d, "8. AI-Supported Design Evidence").getnext(),
                  find(calls, "cap", "8.1 Pre-Submission AI Ideation")[2])

    # 8.2.
    h82 = heading_el(d, "8.2 Team-Reviewed Architecture Freeze")
    rows = find(calls, "table", "8.2 Team-Reviewed Architecture Freeze")[3]
    t82 = table_obj(d, next_tables(h82, 1)[0])
    set_rows(t82, [[text_of(c) for c in r] for r in rows])
    set_widths(t82, [1.75, 2.55, 1.55])

    # 8.3: the prompt box's right cell gets our prompt; provenance after it.
    h83 = heading_el(d, "8.3 Final AI-Vetted Gap Check")
    prompt_t, gap_t = next_tables(h83, 2)
    capt = find(calls, "cap", "8.3 Final AI-Vetted Gap Check")[2]
    prompt = re.search(r"“(.+?)”", capt, re.S).group(1)
    tail = capt[capt.index("”") + 1:].strip()
    tail = tail[:tail.index("Team Response options")].strip() if "Team Response options" in tail else tail
    fill_cell(table_obj(d, prompt_t).rows[0].cells[1], prompt)
    new_par_after(prompt_t, h83.getnext(), tail)
    rows = find(calls, "table", "8.3 Final AI-Vetted Gap Check")[3]
    set_rows(table_obj(d, gap_t), [[text_of(c) for c in r] for r in rows])

    # 9. Design consistency, then the closing identity note.
    h9 = heading_el(d, "9. Design Consistency")
    rows = find(calls, "table", "9. Design Consistency")[3]
    t9 = table_obj(d, next_tables(h9, 1)[0])
    set_rows(t9, [[text_of(c) for c in r] for r in rows])
    closing = [c for c in calls if c[0] == "cap" and c[1] == "9. Design Consistency"][-1][2]
    new_par_after(t9._tbl, h9.getnext(), closing, size=8.5)

    tidy_spacing(d)
    keep_layout(d)
    scrub_ids(d)
    b._clean_package(d, "INF2004 Week 6 Design Review (course template)",
                     "Week 6 design review in the course template; no instructor "
                     "approval or hardware result is claimed.")
    d.save(str(out))
    text = "\n".join(para_text(el) for el in d.element.body.iter(W_P))
    left = [w for w in LEFTOVERS if w.lower() in text.lower()]
    left += [x[:60] for x in template_examples() if x in text]
    if left:
        raise SystemExit(f"template text left in {out}: {left}")
    print(f"{out}: built from {entries_path}; figures "
          f"{len(d.inline_shapes)}; blank fields: "
          + (", ".join(k for k in ('group_id', 'repo_url') if not b.Field(k).value) or "none"))


def template_examples():
    """Example text of the template's data rows (every table except the
    instruction boxes and field labels we keep), for the leftover check."""
    t = Document(str(TEMPLATE))
    keep = ("Purpose", "How to use this template", "Use one of these Team Response",
            "AI reviewer prompt used", "Subsystem name", "Related RTOS task",
            "Reason a statechart is needed", "Field", "Student Entry", "GROUP ID",
            "Submission Detail")
    out = set()
    for tbl in t.tables:
        if len(tbl.rows) < 2:        # instruction and picture/explanation boxes
            continue                 # ("Enter ..." is in LEFTOVERS)
        for row in tbl.rows[1:]:
            for cell in row.cells:
                for par in cell.paragraphs:
                    x = par.text.strip()
                    # section 9's check questions are kept on purpose
                    if len(x) > 20 and not x.startswith(keep) and not x.endswith("?"):
                        out.add(x)
    return out


def is_blank(el):
    """An empty body paragraph (no text, picture or section break)."""
    return (el.tag == W_P and not para_text(el).strip()
            and el.find(".//" + qn("w:drawing")) is None
            and el.find(".//" + qn("w:sectPr")) is None)


def tidy_spacing(d):
    """The template spaces its sections with runs of blank lines (and one
    page break) sized for its short examples; with real content they push
    blank space onto page tops. Collapse each run: nothing between text,
    one small blank paragraph where it separates two tables (adjacent
    tables would merge)."""
    items = body_items(d)
    i = 0
    while i < len(items):
        if not is_blank(items[i]):
            i += 1
            continue
        j = i
        while j < len(items) and is_blank(items[j]):
            j += 1
        prev_t = i > 0 and items[i - 1].tag == W_TBL
        next_t = j < len(items) and items[j].tag == W_TBL
        run = items[i:j]
        keep = run[0] if (prev_t and next_t) else None
        for el in run:
            if el is not keep:
                el.getparent().remove(el)
        if keep is not None:
            for child in list(keep):
                if child.tag != qn("w:pPr"):
                    keep.remove(child)
            ppr = keep.find(qn("w:pPr"))
            if ppr is None:
                ppr = parse_xml(f"<w:pPr {nsdecls('w')}/>")
                keep.insert(0, ppr)
            for x in ppr.findall(qn("w:spacing")) + ppr.findall(qn("w:pStyle")):
                ppr.remove(x)
            ppr.insert(0, parse_xml(f'<w:spacing {nsdecls("w")} w:before="0" '
                                    'w:after="0" w:line="240" w:lineRule="auto"/>'))
        i = j


def set_widths(t, first):
    """Set the first column widths (inches); the last column takes the rest
    of the table's total width."""
    grid = t._tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol"))
    total = sum(int(g.get(qn("w:w"))) for g in grid)
    tw = [int(w * 1440) for w in first]
    tw.append(total - sum(tw))
    for g, w in zip(grid, tw):
        g.set(qn("w:w"), str(w))
    for row in t.rows:
        for cell, w in zip(row.cells, tw):
            tcpr = cell._tc.get_or_add_tcPr()
            for x in tcpr.findall(qn("w:tcW")):
                tcpr.remove(x)
            tcw = parse_xml(f'<w:tcW {nsdecls("w")} w:w="{w}" w:type="dxa"/>')
            cnf = tcpr.find(qn("w:cnfStyle"))      # schema: cnfStyle, then tcW
            if cnf is not None:
                cnf.addnext(tcw)
            else:
                tcpr.insert(0, tcw)


def scrub_ids(d):
    """Copied template rows and blocks repeat Word's paragraph IDs and
    bookmarks; drop them (Word regenerates paragraph IDs) so every ID in
    the saved file is unique."""
    w14 = "http://schemas.microsoft.com/office/word/2010/wordml"
    for el in d.element.body.iter():
        for attr in ("{%s}paraId" % w14, "{%s}textId" % w14):
            if attr in el.attrib:
                del el.attrib[attr]
    for tag in ("w:bookmarkStart", "w:bookmarkEnd"):
        for el in list(d.element.body.iter(qn(tag))):
            el.getparent().remove(el)


def keep_layout(d):
    """Word/LibreOffice pagination aids: a lead-in line or figure paragraph
    stays on the page of the table or caption that follows it, and table
    rows are not split across pages."""
    items = body_items(d)
    for el, nxt in zip(items, items[1:]):
        if el.tag == W_P and para_text(el).strip() and (
                nxt.tag == W_TBL or el.find(".//" + qn("w:drawing")) is not None):
            ppr = el.find(qn("w:pPr"))
            if ppr is None:
                ppr = parse_xml(f"<w:pPr {nsdecls('w')}/>")
                el.insert(0, ppr)
            if ppr.find(qn("w:keepNext")) is None:
                kn = parse_xml(f"<w:keepNext {nsdecls('w')}/>")
                style = ppr.find(qn("w:pStyle"))
                (style.addnext(kn) if style is not None else ppr.insert(0, kn))
        elif el.tag == W_P and el.find(".//" + qn("w:drawing")) is not None:
            ppr = el.find(qn("w:pPr"))
            if ppr is not None and ppr.find(qn("w:keepNext")) is None:
                kn = parse_xml(f"<w:keepNext {nsdecls('w')}/>")
                style = ppr.find(qn("w:pStyle"))
                (style.addnext(kn) if style is not None else ppr.insert(0, kn))
    for tr in d.element.body.iter(W_TR):
        trpr = tr.find(qn("w:trPr"))
        if trpr is None:
            trpr = parse_xml(f"<w:trPr {nsdecls('w')}/>")
            tblpex = tr.find(qn("w:tblPrEx"))
            (tblpex.addnext(trpr) if tblpex is not None else tr.insert(0, trpr))
        if trpr.find(qn("w:cantSplit")) is None:
            cs = parse_xml(f"<w:cantSplit {nsdecls('w')}/>")
            before = [c for c in trpr if c.tag.split("}")[1] in
                      ("cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter")]
            (before[-1].addnext(cs) if before else trpr.insert(0, cs))


def align_left(cell):
    for par in cell._tc.findall(W_P):
        ppr = par.find(qn("w:pPr"))
        if ppr is None:
            continue
        for jc in ppr.findall(qn("w:jc")):
            jc.set(qn("w:val"), "left")


def cap_first(s):
    s = s.strip()
    return s[:1].upper() + s[1:] if s else s


# ------------------------------------------------------------------- check --
def check_pdf(path):
    import pymupdf
    doc = pymupdf.open(str(path))
    bad, sizes = [], []
    text = "\n".join(pg.get_text() for pg in doc)
    flat = re.sub(r"\s+", " ", text)
    for hd in HEADINGS:
        if hd.split(":")[0] not in flat:
            bad.append(f"missing heading {hd!r}")
    for w in LEFTOVERS:
        if w.lower() in flat.lower():
            bad.append(f"template text left: {w!r}")
    for n, pg in enumerate(doc, 1):
        for blk in pg.get_text("rawdict")["blocks"]:
            for line in blk.get("lines", []):
                for span in line["spans"]:
                    if not "".join(c["c"] for c in span["chars"]).strip():
                        continue
                    sizes.append(span["size"])
                    if span["size"] < 8.4:          # template tables: 8.5 pt
                        bad.append(f"page {n}: {span['size']:.1f} pt text")
                    if not pg.rect.contains(pymupdf.Rect(span["bbox"])):
                        bad.append(f"page {n}: text outside the page")
        for img in pg.get_image_info():
            if not pg.rect.contains(pymupdf.Rect(img["bbox"])):
                bad.append(f"page {n}: image outside the page")
    figs = sum(len(pg.get_image_info()) for pg in doc)
    if figs != 8:
        bad.append(f"{figs} images, expected the 8 figures")
    blanks = sorted(set(re.findall(b.PLACEHOLDER + r"(group ID|public repository URL)", flat)))
    if blanks:
        print("blank fields (type them in Word before submitting): " + ", ".join(blanks))
    print(f"{path}: {len(doc)} pages; smallest text {min(sizes):.1f} pt; {figs} figures; "
          f"{len(bad)} problems")
    for x in bad:
        print("  " + x)
    return len(bad)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    default = ROOT / "local" / "week6_entries.json"
    ap.add_argument("--entries", type=Path,
                    default=default if default.exists()
                    else ROOT / "Part1_Week6_Report" / "report_entries.json")
    ap.add_argument("--out", type=Path, default=OUT)
    ap.add_argument("--check-pdf", type=Path, metavar="PDF")
    args = ap.parse_args(argv)
    if args.check_pdf:
        return 1 if check_pdf(args.check_pdf) else 0
    build(args.entries, args.out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
