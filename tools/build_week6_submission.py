"""Hold the Week 6 report content; build the five-page alternative and the
buddy review forms.

The submission is Part1_Week6_Report/INF2004_Week6_Design_Review.docx, the
course's own Word template filled with this module's content by
Tools/build_week6_template_report.py (team decision, 26 September 2026: follow
the template, no page limit), with the figures of Template_Figures.  This
script's own report is the compact alternative, written to
Part1_Week6_Report/Five_Page_Version/: the template's sections on at most five
numbered pages (the module overview's "3-5 pages document"), for use only if
the lecturer enforces the page limit.  --cover adds an SIT cover page and a
contents page (front matter, unnumbered) before those pages.  Requires Python 3,
python-docx and matplotlib (PyMuPDF only for --check-pdf and --update-toc).
Render with Microsoft Word (Tools/render_week6_word.ps1) and inspect every
page after each content edit: nothing here proves the page count.  No public
export input or firmware file is modified by this generator.

Usage
  python Tools/build_week6_submission.py            draft build (default)
  python Tools/build_week6_submission.py --final    no draft wording; refuses
                                  while any entry is missing, a Team response
                                  is not a template option, or
                                  software_status.json is not finalised; it
                                  also drops the "proposal" qualifiers of the
                                  ownership table and FR10, so run it only
                                  after the team has confirmed both
  --cover               add the SIT cover and contents pages (7 pages in all);
                        without it the report is the template's five pages
  --blank-fields        with --final: allow the identity fields (group ID,
                        repository URL, names, student IDs) to stay blank.
                        Each is a Word content control showing grey "Click to
                        enter ..." text; the cover copy is bound to the body
                        copies, so a name typed once on the cover fills the
                        ownership table too
  --entries FILE.json   team entries keyed as --list-entries prints (the
                        report coordinator keeps it at the git-ignored
                        local/week6_entries.json).  The build writes them
                        into the TRACKED report DOCX; the team agreed on
                        25 September 2026 that the named report may be
                        committed to this private repository, never to a
                        public one
  --list-entries        print every entry key, its state and hint, and exit
  --fill-test OUT.docx  write a scratch report whose every pending entry holds
                        100-120 characters of dummy text (a Team response the
                        longest option; identity fields realistic long
                        dummies); capacity test only, refuses repository paths
  --check-pdf PDF       check a Word-rendered five-page alternative: at most
                        five A4 body pages (after a cover and contents page
                        only for a --cover build); every text
                        span >= 9.5 pt, all text inside the page body,
                        figures at full size, all template headings, and the
                        contents page numbers equal to the rendered pages;
                        with --final also no [ENTER, DRAFT, FILL-TEST, the
                        word "proposal", blank identity field (unless
                        --blank-fields) and a 64-hex digest
  --update-toc PDF      record the body page of every heading and figure of
                        a Word render in Part1_Week6_Report/toc_pages.json;
                        rebuild and render again afterwards

Software identity (digest, checks, CI, hardware status) is
read from Part1_Week6_Report/software_status.json.  Refresh it from the final
verification of the exact revision being submitted before --final: the
generator checks the fields are final, not that they match the tree.
Figures are drawn with Windows Calibri when it is installed, otherwise with
the metric-compatible Carlito (registered automatically on Linux), and
only as a last resort with DejaVu Sans, which changes text widths: draw
them on Windows or on Linux with Carlito installed.  diagrams() draws both
sets: fig_<name>() for the compact five-page build (Compact_Figures/) and,
with TEMPLATE True, tpl_<name>() for the submission in the course
template (Template_Figures/, TPL_W wide).  On 26 September 2026 all eight
template figures were drawn this way on Linux with Carlito, and the
compact Figs. 2-5, 7 and 8 were redrawn at unchanged sizes to fix the
errors found by the figure review; compact Figs. 1 and 6 re-render
byte-identical.  To keep the committed figures in a rebuild, import the
module, set diagrams = lambda: None, then call main().
"""
import argparse
import copy
import datetime
import hashlib
import io
import json
import re
import sys
import zipfile
import zlib
from pathlib import Path
from xml.sax.saxutils import escape

import matplotlib
matplotlib.use("Agg")
from matplotlib import font_manager  # noqa: E402
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch  # noqa: E402
from matplotlib.path import Path as MPath  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.opc.packuri import PackURI  # noqa: E402
from docx.opc.part import Part  # noqa: E402
from docx.oxml import OxmlElement, parse_xml  # noqa: E402
from docx.oxml.ns import nsdecls, qn  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Part1_Week6_Report"
FIG = OUT / "Compact_Figures"
# Since 26 September 2026 the submission is the course template, whose
# picture boxes are full width: Template_Figures holds the figures redrawn
# for that layout (larger, with the extra labels the compact pages cannot
# fit). The compact five-page build keeps Compact_Figures.
FIG_T = OUT / "Template_Figures"
TPL_W = 6.90                     # template picture box (7.1 in) less padding
TEMPLATE = False                 # True while diagrams() draws FIG_T
# The template build (no page limit) prints the longer wording of a few
# notes; the five-page build keeps the short forms. Set by
# build_week6_template_report.py before it reads report().
LONG_TEXT = False
STATUS = OUT / "software_status.json"
TOC_PAGES = OUT / "toc_pages.json"
TEMPLATE_DOCX = ROOT / "References" / "Week_6_Design_Review_Template(1).docx"
LOGO = ROOT / "Tools" / "templates" / "sit_logo.png"
# Since 26 September 2026 the submission is the course-template build
# (Tools/build_week6_template_report.py writes OUT/INF2004_Week6_Design_Review.docx
# from this module's content); this generator's compact five-page layout
# is kept as an alternative in its own folder.
REPORT = OUT / "Five_Page_Version" / "INF2004_Week6_Design_Review_5pages.docx"
FORMS = OUT / "Buddy_Reviews" / "Buddy_Review_Sheets.docx"
DOC_DATE = datetime.datetime(2026, 9, 25)  # fixed: reproducible package
FORMS_DATE = datetime.datetime(2026, 9, 23)  # buddy forms unchanged since

# A4 with 0.45 in side and 0.42 in top/bottom margins.
PAGE_W, PAGE_H, SIDE, TOPBOT = 8.27, 11.69, 0.45, 0.42
WIDTH = round(PAGE_W - 2 * SIDE, 2)          # 7.37 in text width
SIDE_FIG = 4.30                                # subsystem chart width
MIN_PT = 9.5                                   # smallest text anywhere (PDF)
FS = 9.5                                       # figure label size (print)
BOXH = 0.345                                   # two-line state box height
BODY_PT, TABLE_PT = 10.0, 10.0                 # Word PDF export writes 9.96 pt
LEAD = Pt(11.4)                                # exact leading, 10 pt text
MAX_PAGES = 5                                  # numbered body pages
FRONT_PAGES = 2                                # cover + contents, unnumbered
COVER = False                                  # --cover: add those front pages
BLUE = "244A66"                                  # buddy forms (unchanged)
# Report colours from the course template: navy title, Heading 1 and table
# header rows (white bold text), lighter blue Heading 2 and title rule.
NAVY, SKY, RULE_BLUE, GRID = "17365D", "2F75B5", "4F81BD", "B8C4CE"
HEAD_FOOT = "666666"                             # running header and footer
# Front matter: 0.75 in top, 0.8 in bottom and 1 in side margins.
FRONT_TOP, FRONT_BOTTOM, FRONT_SIDE = 0.75, 0.8, 1.0
FRONT_WIDTH = round(PAGE_W - 2 * FRONT_SIDE, 2)  # 6.27 in
COVER_FONT = "Calibri"             # one typeface from cover to body
COVER_TITLE = "Autonomous Robot Car Week 6 Design Review Report"
COVER_LINES = ("Autonomous Robot Car", "Week 6 Design Review Report")
DEGREE = ("BACHELOR OF ENGINEERING (WITH HONOURS)",
          "IN INFORMATION AND COMMUNICATIONS TECHNOLOGY",
          "(INFORMATION SECURITY)")
MODULE = "INF2004 Embedded Systems Programming AY 2026/2027 Trimester 1"
# Cover table: navy header row as in the body tables, light banding.
COVER_HEAD, COVER_ALT, COVER_RULE = NAVY, "F7F9FB", GRID
RTOS_TAIL = " · Related RTOS task"   # 6.x heading tail (regular weight)
SUBSYSTEMS = ("Communications", "Motion Control", "Optical Navigation",
              "Terrain and Hump Height", "Obstacle Handling")

for _name in ("calibri.ttf", "calibrib.ttf", "calibrii.ttf"):
    _font = Path("C:/Windows/Fonts") / _name
    if _font.exists():
        font_manager.fontManager.addfont(str(_font))
# Off Windows, Carlito (metric-compatible with Calibri) keeps every label
# the same width, so the layout checks give the same answer.
for _font in Path("/usr/share/fonts/truetype/crosextra").glob("Carlito*.ttf"):
    font_manager.fontManager.addfont(str(_font))
_have = {f.name for f in font_manager.fontManager.ttflist}
FONT = "Calibri" if "Calibri" in _have else \
    "Carlito" if "Carlito" in _have else "DejaVu Sans"
# Deterministic SVG: fixed hash salt for clip-path ids; no date metadata.
matplotlib.rcParams.update({"font.family": FONT, "svg.hashsalt": "inf2004-week6",
                            "svg.fonttype": "path", "path.simplify": False})

# ---------------------------------------------------------------- entries --
ENTRIES = {}      # key -> genuine team text (from --entries)
FILL = False      # --fill-test: every entry becomes dummy text
USED = {}         # key -> hint, in document order
CHOICE = {}       # key -> allowed values (entries that must be one option)
_WORDS = ("accepted after peer challenge; keeps the single owner and the "
          "bounded retry; evidence follows in the dated review record with "
          "the measured bench values and the reviewer's own notes").split()
# Template 8.3: "Use one of these Team Response options".
RESPONSES = ("Fixed", "Partially fixed", "Accepted as risk", "Deferred",
             "Rejected with reason", "Needs validation")


def dummy(key):
    """Deterministic 100-120 character stand-in for a genuine entry."""
    n = 101 + zlib.crc32(key.encode()) % 20          # 101..120
    text, i = "FILL-TEST " + key, 0
    while len(text) < n:
        text += " " + _WORDS[i % len(_WORDS)]
        i += 1
    return text[:n - 1].rstrip(" ;,") + "."            # 100..120 characters


def E(key, hint="", choices=None):
    """A pending team entry: genuine text, dummy text or a labelled gap.

    With choices, a genuine entry must be exactly one of them and the fill
    test uses the longest option (a choice is not free text).
    """
    USED.setdefault(key, hint)
    if choices:
        CHOICE[key] = choices
    if FILL:
        return max(choices, key=len) if choices else dummy(key)
    value = str(ENTRIES.get(key, "")).strip()
    if value:
        if choices and value not in choices:
            raise SystemExit(f"entry {key!r} = {value!r}: use one of "
                             + ", ".join(choices))
        return value
    return "[ENTER]" if not hint else "[ENTER " + hint + "]"


# Identity fields.  Each is a plain-text Word content control bound to one
# node of a custom XML part, so every copy (cover, submission detail,
# ownership table) shows the same value and a name typed once in Word fills
# them all.  A blank field shows grey placeholder text starting with
# PLACEHOLDER; --final refuses blanks unless --blank-fields is given.
PLACEHOLDER = "Click to enter "
FIELDS = {"group_id": "group ID", "repo_url": "public repository URL"}
for _i in range(1, 6):
    FIELDS[f"b{_i}_name"] = f"B{_i} name"
    FIELDS[f"b{_i}_id"] = f"B{_i} student ID"
# Fill test: realistic worst-case lengths, never real people.
FIELD_DUMMY = {"group_id": "P3-13",
               "repo_url": "https://github.com/inf2004-p3-13/autonomous-robot-car"}
for _i in range(1, 6):
    FIELD_DUMMY[f"b{_i}_name"] = "Muhammad Abdul Rahman Bin Mohamed Iskandar"
    FIELD_DUMMY[f"b{_i}_id"] = f"250{_i}999"
TEAM_NS = "urn:inf2004:week6:team"
TEAM_STORE = "{3B7C1E52-9A4D-4F0B-8C2E-5D6A7B8C9D01}"
_SDT_IDS = [0]


class Field:
    """A bound identity field; its value comes from --entries or stays blank."""

    def __init__(self, key, bold=False, size=None, font=None):
        USED.setdefault(key, FIELDS[key])
        self.key, self.bold, self.size, self.font = key, bold, size, font

    @property
    def value(self):
        if FILL:
            return FIELD_DUMMY[self.key]
        return str(ENTRIES.get(self.key, "")).strip()

    def element(self):
        """The w:sdt element: bound text, or grey placeholder text when blank.

        The same run properties (schema order: rFonts, b, sz) are given to
        the control, so Word keeps them when it replaces the bound text.
        """
        _SDT_IDS[0] += 1
        value = self.value
        props = ((f'<w:rFonts w:ascii="{self.font}" w:hAnsi="{self.font}" '
                  f'w:cs="{self.font}" w:eastAsia="{self.font}"/>' if self.font else "")
                 + ("<w:b/>" if self.bold else "")
                 + (f'<w:sz w:val="{int(self.size * 2)}"/>' if self.size else ""))
        grey = "" if value else '<w:rStyle w:val="PlaceholderText"/>'
        shown = escape(value or PLACEHOLDER + FIELDS[self.key])
        return parse_xml(
            f'<w:sdt {nsdecls("w")}><w:sdtPr><w:rPr>{props}</w:rPr>'
            f'<w:alias w:val="{escape(FIELDS[self.key])}"/>'
            f'<w:tag w:val="{self.key}"/><w:id w:val="{_SDT_IDS[0]}"/>'
            + ("" if value else "<w:showingPlcHdr/>") +
            f'<w:dataBinding w:prefixMappings="xmlns:ns0=\'{TEAM_NS}\'" '
            f'w:xpath="/ns0:team[1]/ns0:{self.key}[1]" w:storeItemID="{TEAM_STORE}"/>'
            f'<w:text/></w:sdtPr><w:sdtContent><w:r><w:rPr>{grey}{props}</w:rPr>'
            f'<w:t xml:space="preserve">{shown}</w:t></w:r></w:sdtContent></w:sdt>')


def team_store(d):
    """Add the custom XML part that the identity content controls bind to."""
    xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
           f'<team xmlns="{TEAM_NS}">'
           + "".join(f"<{k}>{escape(Field(k).value)}</{k}>" for k in FIELDS)
           + "</team>").encode("utf-8")
    props = ('<?xml version="1.0" encoding="UTF-8" standalone="no"?>'
             f'<ds:datastoreItem ds:itemID="{TEAM_STORE}" xmlns:ds="http://schemas.'
             'openxmlformats.org/officeDocument/2006/customXml"><ds:schemaRefs/>'
             '</ds:datastoreItem>').encode("utf-8")
    pkg = d.part.package
    used = {str(part.partname) for part in pkg.iter_parts()}
    k = 1
    while f"/customXml/item{k}.xml" in used:
        k += 1
    item = Part(PackURI(f"/customXml/item{k}.xml"), "application/xml", xml, pkg)
    prop = Part(PackURI(f"/customXml/itemProps{k}.xml"), "application/vnd."
                "openxmlformats-officedocument.customXmlProperties+xml", props, pkg)
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
    item.relate_to(prop, rel + "customXmlProps")
    d.part.relate_to(item, rel + "customXml")


STATUS_NEED =["build_input_digest", "local_checks_date", "local_checks_summary",
               "ci_status_text", "hardware_status_text"]


def software_status(final=False):
    """Identity caption from software_status.json.

    --final refuses unless every printed field is finalised (no PENDING
    text) and the digest is a full 64-hex build-input digest.  The file must
    be refreshed from the final verification of the exact revision being
    submitted; nothing here can prove that it was.  In the fill test,
    pending values are replaced by realistic-length dummies.
    """
    data = json.loads(STATUS.read_text(encoding="utf-8"))
    missing = [k for k in STATUS_NEED if not str(data.get(k, "")).strip()]
    if missing:
        raise SystemExit(f"{STATUS}: missing {missing}")
    data = {k: str(v).strip().rstrip(".") for k, v in data.items()}
    if FILL:
        for k in STATUS_NEED:
            if "PENDING" in data[k].upper():
                data[k] = (hashlib.sha256(k.encode()).hexdigest()
                           if k == "build_input_digest" else dummy(k).rstrip("."))
    if final:
        bad = [k for k in STATUS_NEED if "PENDING" in data[k].upper()]
        if not re.fullmatch(r"[0-9a-f]{64}", data["build_input_digest"]):
            bad.append("build_input_digest (not 64 hex characters)")
        if bad:
            raise SystemExit(f"--final refused: {STATUS} is not finalised: "
                             + ", ".join(sorted(set(bad))))
    digest = data["build_input_digest"]
    if re.fullmatch(r"[0-9a-f]{64}", digest):
        digest = digest[:12] + "…" + digest[-6:]
    nb = str.maketrans("-", "‑")          # keep dates and the path unbroken
    return (f"**Software identity:** private build-input digest "
            f"{digest}; local checks {data['local_checks_date'].translate(nb)}: "
            f"{data['local_checks_summary']}. CI: {data['ci_status_text']}. "
            f"{data['hardware_status_text']}.")


# ---------------------------------------------------------------- figures --
KIND = {  # fill, edge, line style
    "state": ("#edf3f7", "#244a66", "-"),
    "fault": ("#fae3de", "#a3402f", "-"),
    "soft": ("#fdf1e3", "#a3402f", "--"),
    "task": ("#e3eef9", "#1f4e79", "-"),
    "isr": ("#fde8b8", "#8a5a00", "-"),
    "hw": ("#efefef", "#555555", "--"),
    "kobj": ("#ffffff", "#2e2e2e", "-"),
    "store": ("#ece7f6", "#5b4a91", "-"),
    "good": ("#e4f3ed", "#387b60", "-"),
}
EDGE = "#244a66"
# Template-style palette (the course template's example diagrams): pastel
# fill and a darker edge per role, and one colour per kind of flow.
ROLE = {
    "green": ("#e8f5e9", "#2e7d32"), "yellow": ("#fff8e1", "#c79100"),
    "orange": ("#fff3e0", "#e65100"), "purple": ("#f3e5f5", "#6a1b9a"),
    "red": ("#ffebee", "#c62828"), "blue": ("#e3f2fd", "#1565c0"),
    "teal": ("#e0f2f1", "#00796b"), "grey": ("#eeeeee", "#616161"),
    "note": ("#fafafa", "#9e9e9e"),
}
FLOW = {"data": "#1f4fbf", "command": "#6a1b9a", "status": "#00897b",
        "fault": "#c62828", "hw": "#222222", "event": "#c62828"}
TITLE_PT = 12.5


class Fig:
    """Fixed-size figure whose text is FS points when inserted at full size.

    save() refuses clipped artists, labels touching arrows, boxes or other
    labels, box text overflowing its box, arrows passing through boxes and
    (unless allowed) arrows crossing each other.
    """

    def __init__(self, name, w, h):
        self.name, self.w, self.h = name, w, h
        self.fig = plt.figure(figsize=(w, h), dpi=100)
        self.ax = self.fig.add_axes([0, 0, 1, 1])
        self.ax.set_xlim(0, w)
        self.ax.set_ylim(0, h)
        self.ax.axis("off")
        self.boxes, self.arrows, self.labels, self.frames = [], [], [], []
        self.swatches = []

    def box(self, x, y, w, h, label, kind="state", bold=None):
        fill, edge, ls = KIND[kind]
        patch = FancyBboxPatch((x, y), w, h,
                               boxstyle="round,pad=0,rounding_size=0.045",
                               lw=0.9, ec=edge, fc=fill, ls=ls, zorder=2)
        self.ax.add_patch(patch)
        text = self.ax.text(x + w / 2, y + h / 2, label, ha="center", va="center",
                            fontsize=FS, linespacing=1.3, zorder=4,
                            fontweight="bold" if bold else "normal")
        self.boxes.append((patch, text))
        return (x, y, w, h)

    def frame(self, x, y, w, h, title=None, tx=0.0, ty=0.0, ha="left"):
        patch = FancyBboxPatch((x, y), w, h,
                               boxstyle="round,pad=0,rounding_size=0.08",
                               lw=0.9, ec="#5b6f82", fc="none", ls=(0, (4, 2)),
                               zorder=1)
        self.ax.add_patch(patch)
        self.frames.append(patch)
        if title:
            self.label(tx, ty, title, ha=ha, va="center", style="italic")

    def arrow(self, pts, label=None, at=None, ha="left", va="center",
              both=False, dashed=False, color=None, lw=0.9, lcolor=None):
        """Polyline arrow; color is a FLOW key or a colour (default EDGE);
        the label takes the arrow's colour unless lcolor is given."""
        col = FLOW.get(color, color) if color else EDGE
        path = MPath(pts, [MPath.MOVETO] + [MPath.LINETO] * (len(pts) - 1))
        self.ax.add_patch(FancyArrowPatch(
            path=path, arrowstyle="<|-|>" if both else "-|>",
            mutation_scale=8 + 2 * (lw - 0.9), lw=lw, color=col,
            ls=(0, (3, 2)) if dashed else "-", zorder=3))
        self.arrows.append(pts)
        if label:
            self.label(at[0], at[1], label, ha=ha, va=va,
                       color=lcolor or (col if color else None))

    def label(self, x, y, s, ha="left", va="center", **kw):
        kw = {k: v for k, v in kw.items() if v is not None}
        text = self.ax.text(x, y, s, ha=ha, va=va, fontsize=kw.pop("fontsize", FS),
                            linespacing=kw.pop("linespacing", 1.0), zorder=5, **kw)
        self.labels.append(text)
        return text

    # -- template-style primitives (course template look) -------------------
    def _rounded(self, x, y, w, h, fill, edge, lw=1.1, ls="-", r=0.06, z=2):
        patch = FancyBboxPatch((x, y), w, h,
                               boxstyle=f"round,pad=0,rounding_size={r}",
                               lw=lw, ec=edge, fc=fill, ls=ls, zorder=z)
        self.ax.add_patch(patch)
        return patch

    def _boxed_text(self, patch, x, y, s, ha="center", va="center", **kw):
        """Text that must stay inside patch (checked like box text)."""
        t = self.ax.text(x, y, s, ha=ha, va=va, fontsize=kw.pop("fontsize", FS),
                         linespacing=kw.pop("linespacing", 1.25), zorder=4, **kw)
        self.boxes.append((patch, t))
        return t

    def card(self, x, y, w, h, title, sub=None, bullets=(), role="blue",
             dashed=False):
        """Template task/device card: bold title, optional (subtitle) and
        left-aligned bullets, in the role's pastel colour."""
        fill, edge = ROLE[role]
        patch = self._rounded(x, y, w, h, fill, edge, ls=(0, (4, 2)) if dashed else "-")
        top = y + h - 0.07
        if bullets:
            self._boxed_text(patch, x + w / 2, top, title, va="top", fontweight="bold")
            yy = top - 0.19 * title.count("\n") - 0.19
            if sub:
                self._boxed_text(patch, x + w / 2, yy, sub, va="top")
                yy -= 0.19 * (sub.count("\n") + 1)
            self._boxed_text(patch, x + 0.09, yy - 0.03,
                             "\n".join("• " + b for b in bullets), ha="left", va="top",
                             linespacing=1.3)
        else:
            nt = title.count("\n") + 1
            ns = sub.count("\n") + 1 if sub else 0
            step = 0.175                         # one line at FS, with leading
            top = y + h / 2 + step * (nt + ns) / 2
            self._boxed_text(patch, x + w / 2, top - step * nt / 2, title,
                             fontweight="bold")
            if sub:
                self._boxed_text(patch, x + w / 2, top - step * (nt + ns / 2), sub)
        return (x, y, w, h)

    def state(self, x, y, w, h, name, desc=None, role="blue"):
        """Template statechart state: bold NAME over a (description)."""
        fill, edge = ROLE[role]
        patch = self._rounded(x, y, w, h, fill, edge)
        if desc:
            self._boxed_text(patch, x + w / 2, y + h * 0.66, name, fontweight="bold")
            self._boxed_text(patch, x + w / 2, y + h * 0.30, desc)
        else:
            self._boxed_text(patch, x + w / 2, y + h / 2, name, fontweight="bold")
        return (x, y, w, h)

    def queue(self, x, y, color="data", cells=5, w=0.56, h=0.16):
        """Message buffer drawn as a row of cells (template queue icon)."""
        col = FLOW.get(color, color)
        patch = FancyBboxPatch((x, y), w, h, boxstyle="square,pad=0", lw=1.2,
                               ec=col, fc="white", zorder=2)
        self.ax.add_patch(patch)
        for k in range(1, cells):
            self.ax.plot([x + w * k / cells] * 2, [y, y + h], color=col, lw=1.0,
                         zorder=3)
        self.swatches.append(patch)
        return (x, y, w, h)

    def start(self, x, y, r=0.06):
        """Initial state: filled dot centred at (x, y)."""
        from matplotlib.patches import Circle
        c = Circle((x, y), r, color="#111111", zorder=3)
        self.ax.add_patch(c)
        self.swatches.append(c)

    def final(self, x, y, r=0.075):
        """Final state: ring with a dot, centred at (x, y)."""
        from matplotlib.patches import Circle
        ring = Circle((x, y), r, fc="white", ec="#111111", lw=1.2, zorder=3)
        self.ax.add_patch(ring)
        self.ax.add_patch(Circle((x, y), r * 0.55, color="#111111", zorder=4))
        self.swatches.append(ring)

    def panel(self, x, y, w, h, title, sub=None, role="blue", head=0.42):
        """Subsystem panel with a tinted header band (template statecharts)."""
        fill, edge = ROLE[role]
        outer = self._rounded(x, y, w, h, "white", edge, lw=1.0, r=0.08, z=0)
        band = self._rounded(x, y + h - head, w, head, fill, edge, lw=1.0, r=0.08, z=0)
        self.frames += [outer, band]
        self.ax.text(x + w / 2, y + h - head / 2 + (0.08 if sub else 0), title,
                     ha="center", va="center", fontsize=FS + 1.5, fontweight="bold",
                     color=edge, zorder=5)
        if sub:
            self.ax.text(x + w / 2, y + h - head / 2 - 0.10, sub, ha="center",
                         va="center", fontsize=FS, zorder=5)
        return y + h - head

    def title(self, text, sub=None):
        """Figure title at the top centre (template banner)."""
        self.label(self.w / 2, self.h - 0.14, text, ha="center", va="center",
                   fontsize=TITLE_PT, fontweight="bold")
        if sub:
            self.label(self.w / 2, self.h - 0.34, sub, ha="center", va="center",
                       color="#555555")

    def keybox(self, x, y, w, h, rows, title="Legend", cols=1):
        """Legend box. rows: (kind, spec, text) with kind 'role' (spec = ROLE
        key), 'line' (spec = (FLOW key, dashed)), 'queue' (spec = FLOW key),
        'start' or 'final'."""
        patch = self._rounded(x, y, w, h, ROLE["note"][0], ROLE["note"][1], lw=0.9)
        self.frames.append(patch)
        self.label(x + 0.10, y + h - 0.13, title, fontweight="bold")
        per = -(-len(rows) // cols)
        cw = (w - 0.10) / cols
        for i, (kind, spec, text) in enumerate(rows):
            cx = x + 0.10 + (i // per) * cw
            cy = y + h - 0.34 - (i % per) * 0.19
            if kind == "role":
                fill, edge = ROLE[spec]
                sw = self._rounded(cx, cy - 0.055, 0.26, 0.11, fill, edge, lw=0.9,
                                   r=0.02, z=3)
                self.swatches.append(sw)
            elif kind == "line":
                col, dashed = spec
                self.ax.annotate("", xy=(cx + 0.26, cy), xytext=(cx, cy),
                                 arrowprops=dict(arrowstyle="-|>", color=FLOW.get(col, col),
                                                 lw=1.0, ls="--" if dashed else "-",
                                                 mutation_scale=8), zorder=3)
            elif kind == "queue":
                self.queue(cx, cy - 0.05, spec, cells=4, w=0.26, h=0.10)
            elif kind == "start":
                self.start(cx + 0.13, cy, r=0.045)
            elif kind == "final":
                self.final(cx + 0.13, cy, r=0.055)
            self.label(cx + 0.34, cy, text)

    def notes(self, x, y, w, h, lines, title="Notes"):
        """Notes box with bullet lines (template notes panel)."""
        patch = self._rounded(x, y, w, h, ROLE["note"][0], ROLE["note"][1], lw=0.9)
        self._boxed_text(patch, x + 0.10, y + h - 0.08, title, ha="left", va="top",
                         fontweight="bold")
        self._boxed_text(patch, x + 0.10, y + h - 0.30,
                         "\n".join("• " + t for t in lines), ha="left", va="top",
                         linespacing=1.3)

    def legend(self, x, y, items, gap=0.16):
        """One legend row of kind swatches; returns the end x."""
        for kind, text in items:
            fill, edge, ls = KIND[kind]
            sw = FancyBboxPatch((x, y), 0.24, 0.13,
                                boxstyle="round,pad=0,rounding_size=0.02",
                                lw=0.8, ec=edge, fc=fill, ls=ls, zorder=2)
            self.ax.add_patch(sw)
            self.swatches.append(sw)
            t = self.label(x + 0.30, y + 0.065, text)
            self.fig.canvas.draw()
            bb = t.get_window_extent(self.fig.canvas.get_renderer())
            x = self.ax.transData.inverted().transform((bb.x1, bb.y1))[0] + gap
        return x

    # -- checks -----------------------------------------------------------
    def _check(self, allow_cross):
        fig, ax = self.fig, self.ax
        fig.canvas.draw()
        r = fig.canvas.get_renderer()
        fb = fig.bbox
        errors = []
        tr = ax.transData.transform

        def inside(bb, other, tol=0.5):
            return (bb.x0 >= other.x0 - tol and bb.y0 >= other.y0 - tol and
                    bb.x1 <= other.x1 + tol and bb.y1 <= other.y1 + tol)

        def overlap(a, b, tol=0.0):
            return (a.x0 < b.x1 - tol and b.x0 < a.x1 - tol and
                    a.y0 < b.y1 - tol and b.y0 < a.y1 - tol)

        texts = [t for _, t in self.boxes] + self.labels
        for t in texts:
            if t.get_fontsize() < MIN_PT:
                errors.append(f"font below {MIN_PT}: {t.get_text()!r}")
            if not inside(t.get_window_extent(r), fb):
                errors.append(f"text clipped: {t.get_text()!r}")
        for p in list({id(b): b for b, _ in self.boxes}.values()) + self.frames + self.swatches:
            if not inside(p.get_window_extent(r), fb):
                errors.append(f"patch clipped at {p.get_x():.2f},{p.get_y():.2f}")
        for patch, t in self.boxes:
            inner = patch.get_window_extent(r).padded(-1.5)
            if not inside(t.get_window_extent(r), inner, tol=0):
                errors.append(f"box text overflows box: {t.get_text()!r}")
        bbs = [(t.get_text(), t.get_window_extent(r)) for t in texts]
        for i, (na, a) in enumerate(bbs):
            for nb, b in bbs[i + 1:]:
                if overlap(a, b):
                    errors.append(f"text overlap: {na!r} / {nb!r}")
        uniq = list({id(p): p for p, _ in self.boxes}.values())   # a card has
        pb = [p.get_window_extent(r) for p in uniq]                  # several texts
        for i, a in enumerate(pb):
            for b in pb[i + 1:]:
                if overlap(a, b, tol=0.5):
                    errors.append("box overlap")
        for t in self.labels:
            tb = t.get_window_extent(r)
            for b in pb + [s.get_window_extent(r) for s in self.swatches]:
                if overlap(tb, b):
                    errors.append(f"label on box: {t.get_text()!r}")
        # sample every arrow at ~0.5 px
        samples = []
        for k, pts in enumerate(self.arrows):
            d = tr(pts)
            if not all(fb.x0 - .5 <= x <= fb.x1 + .5 and fb.y0 - .5 <= y <= fb.y1 + .5
                       for x, y in d):
                errors.append(f"arrow {k} leaves figure")
            seq = []
            for (x0, y0), (x1, y1) in zip(d[:-1], d[1:]):
                n = max(2, int(((x1 - x0) ** 2 + (y1 - y0) ** 2) ** .5 * 2))
                seq += [(x0 + (x1 - x0) * i / n, y0 + (y1 - y0) * i / n)
                        for i in range(n + 1)]
            samples.append((k, d, seq))
        for t in texts:
            tb = t.get_window_extent(r).padded(1.0)
            for k, _, seq in samples:
                if any(tb.contains(x, y) for x, y in seq):
                    errors.append(f"arrow {k} crosses text {t.get_text()!r}")
                    break
        for k, d, seq in samples:
            ends = (d[0], d[-1])
            for b in pb + [s.get_window_extent(r) for s in self.swatches]:
                inner = b.padded(-2.0)
                hit = [(x, y) for x, y in seq if inner.contains(x, y) and
                       min(abs(x - e[0]) + abs(y - e[1]) for e in ends) > 4]
                if hit:
                    errors.append(f"arrow {k} passes through a box")
                    break
        if not allow_cross:
            segs = []
            for k, d, _ in samples:
                segs += [(k, tuple(a), tuple(b)) for a, b in zip(d[:-1], d[1:])]

            def cross(p, q, s, t):
                def o(a, b, c):
                    return (b[0] - a[0]) * (c[1] - a[1]) - (b[1] - a[1]) * (c[0] - a[0])
                d1, d2, d3, d4 = o(s, t, p), o(s, t, q), o(p, q, s), o(p, q, t)
                return (d1 * d2 < -1e-6) and (d3 * d4 < -1e-6)
            for i, (ka, a0, a1) in enumerate(segs):
                for kb, b0, b1 in segs[i + 1:]:
                    if ka != kb and cross(a0, a1, b0, b1):
                        errors.append(f"arrows {ka} and {kb} cross")
        return sorted(set(errors))

    def save(self, allow_cross=False):
        errors = self._check(allow_cross)
        if errors:
            plt.close(self.fig)
            raise ValueError(f"Figure {self.name}:\n  " + "\n  ".join(errors))
        out = FIG_T if TEMPLATE else FIG
        self.fig.savefig(out / (self.name + ".png"), dpi=220, facecolor="white")
        svg = out / (self.name + ".svg")
        self.fig.savefig(svg, facecolor="white", metadata={"Date": None})
        svg.write_text("\n".join(line.rstrip() for line in
                                 svg.read_text(encoding="utf-8").splitlines()) + "\n",
                       encoding="utf-8", newline="\n")
        plt.close(self.fig)


def fig_hardware():
    f = Fig("hardware", WIDTH, 1.56)
    cw, h = 1.93, 0.345
    lw = round((WIDTH - cw - 2 * 0.28 - 0.04) / 2, 3)
    lx, cx = 0.02, 0.02 + lw + 0.28
    rx = cx + cw + 0.28
    ys = [1.19, 0.805, 0.42, 0.035]
    left = ["Encoders L / R · Grove 3, Grove 4\nGP4/GP5, GP16/GP17 · GPIO IRQ",
            "IR line L, line R, barcode · Grove 5/6/7\nGP26 / GP27 / GP28 · ADC0–2",
            "GY-511 IMU · Grove 2, addr 0x19\nGP2 SDA / GP3 SCL · I2C1",
            "Buttons (on board) · GPIO in, pull-up\nSTOP GP20 · START/RESET GP21"]
    right = ["Motor driver M1 / M2 (on board)\nGP8/GP9, GP10/GP11 · PWM 20 kHz",
             "Scan servo (servo port 1)\nGP12 · PWM 50 Hz, 1–2 ms",
             "HC-SR04 (servo ports 3/4, Vcc 5 V)\nTRIG GP14 · ECHO 2.2k/3.3k → GP15",
             "Debug console (UART adapter)\nGP0 TX / GP1 RX · UART0 115200"]
    for y, s in zip(ys, left):
        f.box(lx, y, lw, h, s, "hw")
    for y, s in zip(ys, right):
        f.box(rx, y, lw, h, s, "hw")
    f.box(cx, 0.035, cw, 1.50,
          "Pico W on Robo Pico\nRP2040 · micro T-Kernel\nCYW43 Wi-Fi on board (SPI)\n\n"
          "Power bank 5 V → VIN →\nmotor driver, servo V+,\nHC-SR04 Vcc; 3V3: logic\n"
          "and sensors; common GND", "state")
    for i, y in enumerate(ys):
        yc = y + h / 2
        f.arrow([(lx + lw + 0.02, yc), (cx - 0.02, yc)], both=(i == 2))
        f.arrow([(cx + cw + 0.02, yc), (rx - 0.02, yc)], both=(i in (2, 3)))
    f.save()


def tpl_hardware():
    """Fig. 1 for the course template: device cards around the Pico W, one
    arrow per connection, coloured by interface and labelled with its pins."""
    gpio, adc, pwm, i2c = "#1565c0", "#2e7d32", "#e65100", "#6a1b9a"
    uart, spi, power = "#00796b", "#6d4c41", "#c62828"
    q, pad = 0.18, 0.05                        # card line pitch, inner margin
    m = 0.02                                   # outer margin
    lx, lw, gap = m, 1.70, 0.80
    cx = lx + lw + gap
    cw = TPL_W - 2 * (lx + lw + gap)
    rx = cx + cw + gap
    leg_h = 1.04
    bottom = m + leg_h + 0.14                  # main area: device columns
    top = bottom + 3.24
    f = Fig("hardware", TPL_W, top + 0.32)
    f.title("AUTONOMOUS ROBOT CAR – HARDWARE PIN LAYOUT")

    def card(x, y_top, w, lines, dashed=False):
        """Grey device card: bold first line, then centred lines at pitch q;
        returns the y of each line (arrows align with them) and the height."""
        h = 2 * pad + q * len(lines)
        fill, edge = ROLE["grey"]
        patch = f._rounded(x, y_top - h, w, h, fill, edge,
                           ls=(0, (4, 2)) if dashed else "-")
        ys = [y_top - pad - q * (i + 0.5) for i in range(len(lines))]
        for i, (s, y) in enumerate(zip(lines, ys)):
            f._boxed_text(patch, x + w / 2, y, s,
                          fontweight="bold" if i == 0 else "normal")
        return ys, h

    def column(x, cards):
        """Stack cards from top to bottom with equal gaps; return line ys."""
        heights = [2 * pad + q * len(lines) for lines, _ in cards]
        g = (top - bottom - sum(heights)) / (len(cards) - 1)
        y, out = top, []
        for (lines, dashed), h in zip(cards, heights):
            out.append(card(x, y, lw, lines, dashed)[0])
            y -= h + g
        return out

    def link(side, y, pins, col, way):
        """Arrow across the gap at height y; way: 'in' (to the Pico),
        'out' (from the Pico) or 'both'; pin label above the arrow."""
        a, b = (lx + lw + 0.02, cx - 0.02) if side == "L" else \
            (rx - 0.02, cx + cw + 0.02)
        pts = [(a, y), (b, y)] if way != "out" else [(b, y), (a, y)]
        f.arrow(pts, pins, ((a + b) / 2, y + 0.075), ha="center",
                both=way == "both", color=col, lw=1.3)

    # Sensors and buttons on the left, actuators and links on the right;
    # dashed cards are parts of the Robo Pico or the Pico W itself.
    enc, ir, imu, btn = column(lx, [
        (["Encoders (GPIO interrupt)", "Left wheel A/B · Grove 3",
          "Right wheel A/B · Grove 4"], False),
        (["IR Sensors ×3 (A0 → ADC)", "Left line sensor · Grove 5",
          "Right line sensor · Grove 6", "Barcode sensor · Grove 7"], False),
        (["GY-511 IMU (I2C1)", "Accelerometer at 0x19", "Grove 2 · I2C 100 kHz"], False),
        (["Buttons (on board, GPIO)", "STOP · active-low", "START / RESET · active-low"],
         True)])
    mot, servo, hc, probe, wifi = column(rx, [
        (["Motor Driver (on board)", "PWM 20 kHz · VIN 5 V", "M1 → left motor",
          "M2 → right motor"], True),
        (["SG90 Scan Servo (PWM)", "Servo port 1 · 50 Hz, 1–2 ms"], False),
        (["HC-SR04 (GPIO, Vcc 5 V)", "TRIG · servo port 3",
          "ECHO · port 4 via divider"], False),
        (["Raspberry Pi Debug Probe", "UART0 · 115200 8N1"], False),
        (["CYW43 Wi-Fi (on Pico W)", "SPI · MQTT to the broker"], True)])
    link("L", enc[1], "GP4, GP5", gpio, "in")
    link("L", enc[2], "GP16, GP17", gpio, "in")
    for y, s in zip(ir[1:], ("GP26 ADC0", "GP27 ADC1", "GP28 ADC2")):
        link("L", y, s, adc, "in")
    link("L", imu[1], "GP2 SDA", i2c, "both")
    link("L", imu[2], "GP3 SCL", i2c, "out")
    link("L", btn[1], "GP20", gpio, "in")
    link("L", btn[2], "GP21", gpio, "in")
    link("R", mot[2], "GP8, GP9", pwm, "out")
    link("R", mot[3], "GP10, GP11", pwm, "out")
    link("R", sum(servo) / 2, "GP12", pwm, "out")
    link("R", hc[1], "GP14 TRIG", gpio, "out")
    link("R", hc[2], "GP15 ECHO", gpio, "in")
    link("R", probe[0], "GP0 TX", uart, "out")
    link("R", probe[1], "GP1 RX", uart, "in")
    link("R", sum(wifi) / 2, "GP23–25, 29", spi, "both")

    # Controller: the Pico W on its carrier board, and the power it passes on.
    fill, edge = ROLE["blue"]
    pico = f._rounded(cx, bottom, cw, top - bottom, fill, edge, lw=1.4)
    xm, tx = cx + cw / 2, cx + 0.08

    def block(y, head, lines):
        f._boxed_text(pico, xm, y, head, fontweight="bold")
        f._boxed_text(pico, tx, y - 0.15, "\n".join(lines), ha="left", va="top",
                      linespacing=1.3)

    f._boxed_text(pico, xm, top - 0.18, "Raspberry Pi Pico W", fontsize=FS + 1.5,
                  fontweight="bold", color=edge)
    f._boxed_text(pico, xm, top - 0.38, "on Cytron Robo Pico", fontweight="bold")
    f._boxed_text(pico, tx, top - 0.54, "• RP2040 at 125 MHz\n"
                  "• micro T-Kernel, firmware in C\n• 3.3 V I/O, not 5 V tolerant",
                  ha="left", va="top", linespacing=1.3)
    block(top - 1.20, "Pins not used", [
        "• Free: GP6, GP7, GP13, GP19",
        "• On board: GP18 NeoPixel,", "   GP22 buzzer"])
    ydiv = bottom + 1.36
    f.ax.plot([cx + 0.10, cx + cw - 0.10], [ydiv, ydiv], color=edge, lw=0.7,
              ls=(0, (3, 2)), zorder=3)
    block(ydiv - 0.17, "Power distribution", [
        "• VIN → motor driver and", "   servo + pins (5 V)",
        "• 3V3 → Grove ports (≤ 300 mA)", "• Common ground"])
    f._boxed_text(pico, xm, bottom + 0.14, "VIN terminal (3.6–6 V)", color=power)

    # Supply below the board; legend and notes on either side of it.
    pw = 1.70
    px = xm - pw / 2
    _, ph = card(px, m + 2 * pad + 3 * q, pw,
                 ["Power Bank", "20000 mAh · 5 V / 2.1 A", "red +, black − (meter first)"])
    f.arrow([(xm, m + ph + 0.02), (xm, bottom - 0.02)], "5 V",
            (xm + 0.07, (m + ph + bottom) / 2), color=power, lw=1.8)
    kw = px - 0.12 - m
    f.keybox(m, m, kw, leg_h, [
        ("line", (gpio, False), "GPIO"), ("line", (adc, False), "ADC"),
        ("line", (pwm, False), "PWM"), ("line", (i2c, False), "I2C"),
        ("line", (uart, False), "UART"), ("line", (spi, False), "SPI / Wi-Fi"),
        ("line", (power, False), "Power"), (None, None, "On-board part")], cols=2)
    sx = m + 0.10 + (kw - 0.10) / 2            # dashed swatch of the last row
    sy = m + leg_h - 0.34 - 3 * 0.19
    f.swatches.append(f._rounded(sx, sy - 0.055, 0.26, 0.11, *ROLE["grey"], lw=0.9,
                                 ls=(0, (3, 1.5)), r=0.02, z=3))
    nx = px + pw + 0.12
    f.notes(nx, m, TPL_W - m - nx, leg_h, [
        "Proposed pins: wire only after approval",
        "5 V loads: motors, SG90, HC-SR04",
        "3V3 loads: encoders, IR, GY-511",
        "ECHO via 2.2 kΩ/3.3 kΩ: ≈ 3.0 V at GP15",
        "Arrow = signal direction; ↔ both ways"])
    f.save()


def fig_interaction():
    h, g = BOXH, 0.068
    w = round((WIDTH - 0.04 - 6 * g) / 7, 3)
    xs = [0.02 + i * (w + g) for i in range(7)]
    cx = [x + w / 2 for x in xs]
    rE = 0.14
    rD = rE + h + 0.13
    rC = rD + h + 0.15
    hc = BOXH
    rB = rC + hc + 0.15
    rA = rB + h + 0.15
    f = Fig("interaction", WIDTH, rA + h + 0.02)
    for x, s in zip(xs, ["IR ×3\nADC0–2", "GY-511\nI2C1", "STOP button\nGP20",
                         "Motor driver\nPWM, BRAKE", "Encoders,\nECHO pin",
                         "Servo GP12\nTRIG GP14", "CYW43\nWi-Fi (SPI)"]):
        f.box(x, rA, w, h, s, "hw")
    for x, (s, k) in zip(xs, [("Line · P3\n2 ms", "task"), ("IMU · P5\n10 ms", "task"),
                              ("Safety · P1\n2 ms", "task"), ("Motion · P2\n5 ms", "task"),
                              ("GPIO ISR\nedges, echo", "isr"), ("Scan · P6\n2 ms", "task"),
                              ("Radio owner\nP9 · ≥ 2 ms", "task")]):
        f.box(x, rB, w, h, s, k)
    f.box(xs[0], rC, xs[5] + w - xs[0], hc,
          "Shared snapshots, requests, health stamps · short PRIMASK copies "
          "(app_main_lock)\nISR counters: encoder deltas → Motion, encoder "
          "totals → Line, echo times → Scan", "store")
    f.box(xs[0], rD, w, h, "START/RESET\nGP21", "hw")
    f.box(xs[1], rD, w, h, "Vehicle · P4\n10 ms", "task")
    f.box(xs[3], rD, w * 2 + g, h, "Comm · P7 · 5 ms\nSTOP first; others counted", "task")
    f.box(xs[5], rD, w, h, "Telemetry\nP8 · 200 ms", "task")
    f.box(xs[3], rE, w, h, "RX mbf\n4 × 32 B", "kobj")
    f.box(xs[4], rE, w, h, "STOP mbf\n1 × 4 B", "kobj")
    f.box(xs[5], rE, w, h, "TX mbf\n4 × 256 B", "kobj")
    e = 0.012
    tB, tC, tD, tE = rB + h, rC + hc, rD + h, rE + h
    # hardware <-> tasks and ISR
    f.arrow([(cx[0], rA - e), (cx[0], tB + e)])
    f.arrow([(cx[1], rA - e), (cx[1], tB + e)], both=True)
    f.arrow([(xs[2] + 0.25, rA - e), (xs[2] + 0.25, tB + e)])
    f.arrow([(cx[3], tB + e), (cx[3], rA - e)])
    f.arrow([(cx[4], rA - e), (cx[4], tB + e)])
    f.arrow([(cx[5], tB + e), (cx[5], rA - e)])
    f.arrow([(cx[6], tB + e), (cx[6], rA - e)], both=True)
    # Safety may only force BRAKE on the motor driver
    f.arrow([(xs[2] + w - 0.06, tB + e), (xs[3] + 0.12, rA - e)], "BRAKE",
            (xs[2] + w - 0.11, (tB + rA) / 2), ha="right")
    # tasks and ISR <-> shared data (the ISR also reads the motor direction
    # and echo state that tasks write)
    for i in range(6):
        f.arrow([(cx[i], rB - e), (cx[i], tC + e)], both=True)
    # shared data <-> lower tasks
    f.arrow([(cx[1], rC - e), (cx[1], tD + e)], both=True)
    f.arrow([(cx[3] + 0.30, tD + e), (cx[3] + 0.30, rC - e)])
    # Telemetry also writes back: drop counter and its health stamp
    f.arrow([(cx[5], rC - e), (cx[5], tD + e)], both=True)
    f.arrow([(xs[0] + w + e, rD + h / 2), (xs[1] - e, rD + h / 2)])
    # kernel message buffers
    f.arrow([(cx[3], tE + e), (cx[3], rD - e)])
    f.arrow([(cx[4], tE + e), (cx[4], rD - e)])
    f.arrow([(cx[5], rD - e), (cx[5], tE + e)])
    xr = xs[6] + w - 0.18
    f.arrow([(xr, rB - e), (xr, 0.075), (cx[4], 0.075), (cx[4], rE - e)])
    f.arrow([(cx[4], 0.075), (cx[3], 0.075), (cx[3], rE - e)])
    xt = xs[6] + 0.26
    f.arrow([(xs[5] + w + e, rE + h / 2), (xt, rE + h / 2), (xt, rB - e)])
    f.label(xt + 0.05, (rE + h + rD) / 2, "publish", va="center")
    f.label(xt + 0.03, 0.10, "RX, STOP", va="bottom")
    f.legend(0.02, 0.37, [("task", "RTOS task"), ("isr", "ISR"), ("hw", "hardware")])
    f.legend(0.02, 0.13, [("kobj", "kernel msg buffer"), ("store", "shared data (lock)")])
    f.save()


def tpl_interaction():
    """Fig. 2 for the course template: one card per RTOS task, the GPIO ISR,
    hardware at the edges, the three kernel message buffers and the shared
    data under PRIMASK; flows coloured by kind (legend)."""
    W, e, g = TPL_W, 0.012, 0.12
    ISR_STYLE = ("#fdebc8", "#8a5a00")                # amber, dashed border
    STORE_STYLE = ("#e8eef6", "#1f3a5f")              # slate blue, navy border
    hh, th, sh, bh, lh = 0.36, 0.90, 0.58, 0.90, 0.88  # band heights
    y_leg = 0.02                                      # legend and notes
    y5 = y_leg + lh + 0.08                            # START/RESET, CYW43
    yb = y5 + hh + 0.22                               # lower task row
    ys = yb + bh + 0.20                               # shared store
    yt = ys + sh + 0.20                               # upper task row
    yh = yt + th + 0.22                               # hardware row
    f = Fig("interaction", W, yh + hh + 0.28)
    f.title("AUTONOMOUS ROBOT CAR – TASK INTERACTION DIAGRAM")
    hw = dict(color="hw", lw=1.2)                     # flow kinds
    data = dict(color="data", lw=1.2)

    # upper row: the tasks that own a sensor or actuator, hardware above
    w = round((W - 0.04 - 4 * g) / 5, 3)
    xt = [0.02 + i * (w + g) for i in range(5)]
    ct = [x + w / 2 for x in xt]
    for x, (t, s, b), role in zip(xt, [
            ("Safety Task · B5", "(Priority P1 · 2 ms)",
             ("STOP, tilt, liveness", "Latch fault, BRAKE", "Feed watchdog")),
            ("Motion Task · B2", "(Priority P2 · 5 ms)",
             ("PID wheel speed", "Commit PWM", "Pose and path out")),
            ("Line Task · B3", "(Priority P3 · 2 ms)",
             ("Read IR ADC0–2", "Steering, junctions", "Decode barcode")),
            ("IMU Task · B4", "(Priority P5 · 10 ms)",
             ("Read acceleration", "Tilt, motion, hump", "Collision flag")),
            ("Scan Task · B5", "(Priority P6 · 2 ms)",
             ("Servo sweep, TRIG", "Range from echo", "Obstacle profile"))],
            ("red", "yellow", "green", "purple", "teal")):
        f.card(x, yt, w, th, t, s, b, role=role)
    xh = [ct[0] - 0.06] + ct[1:]           # STOP card clear of the BRAKE line
    for c, (t, s), wh in zip(xh, [("STOP button", "GP20, active low"),
                                  ("Motor driver", "PWM GP8–GP11"),
                                  ("IR sensors ×3", "ADC0–2, GP26–28"),
                                  ("GY-511 IMU", "I2C1, GP2/GP3"),
                                  ("Servo, HC-SR04", "GP12, TRIG GP14")],
                             (0.96, 1.08, 1.08, 1.08, 1.08)):
        f.card(c - wh / 2, yh, wh, hh, t, s, role="grey")

    # middle row: watchdog and the shared data (g_shared, ISR counters); every
    # task except the radio owner, and the ISR, copies in or out of it
    xs0, xs1 = 1.02, W - 0.02
    f.card(0.02, ys + (sh - hh) / 2, 0.88, hh, "Watchdog", "1 s timeout", role="grey")
    patch = f._rounded(xs0, ys, xs1 - xs0, sh, *STORE_STYLE, lw=1.4)
    xm, step = (xs0 + xs1) / 2, 0.175
    f._boxed_text(patch, xm, ys + sh / 2 + step,
                  "Shared data · copied in short PRIMASK sections (app_main_lock)",
                  fontweight="bold", color=STORE_STYLE[1])
    f._boxed_text(patch, xm, ys + sh / 2, "g_shared: snapshots, motion and scan "
                  "requests, health stamps, fault latch, counters")
    f._boxed_text(patch, xm, ys + sh / 2 - step, "ISR counters: encoder deltas → "
                  "Motion · encoder totals → Line · echo time → Scan")

    # lower row: supervisor and communication tasks
    xe, gq = 5.92, 0.70                               # ISR column, buffer gap
    wb = round((xe - 0.02 - 3 * 0.10 - gq) / 4, 3)
    xb = [0.02 + i * (wb + 0.10) for i in range(3)]
    xb.append(xb[2] + wb + gq)                        # radio owner
    for x, (t, s, b), role in zip(xb, [
            ("Vehicle Task · B5/B3", "(Priority P4 · 10 ms)",
             ("Mission states", "Motion requests", "Scan requests")),
            ("Telemetry Task · B1", "(Priority P8 · 200 ms)",
             ("Status, 5 Hz", "Perception, 1 Hz", "Diagnostics, 1 Hz")),
            ("Comm Task · B1", "(Priority P7 · 5 ms)",
             ("Takes STOP first", "STOP → fault latch", "Others counted")),
            ("Radio Task · B1", "(Priority P9 · ≥ 2 ms)",
             ("Owns Wi-Fi, MQTT", "Routes RX, STOP", "Publishes ≤ 10 Hz"))],
            ("orange", "blue", "blue", "blue")):
        f.card(x, yb, wb, bh, t, s, b, role=role)
    # GPIO ISR under the shared data; its two inputs share one hardware card
    hi, xim = 0.54, (xe + W - 0.02) / 2
    isr = f._rounded(xe, yb + bh - hi, W - 0.02 - xe, hi, *ISR_STYLE, ls=(0, (4, 2)))
    f._boxed_text(isr, xim, yb + bh - 0.14, "GPIO ISR", fontweight="bold")
    f._boxed_text(isr, xim, yb + bh - 0.36, "counts edges,\ntimes ECHO",
                  linespacing=1.15)
    ye = yb + bh - hi - 0.18                          # encoder card top
    enc = f._rounded(xe, y5, W - 0.02 - xe, ye - y5, *ROLE["grey"])
    for dy, s, bold in ((1.5, "Encoders ×2", True), (0.5, "GP4/5, GP16/17", False),
                        (-0.55, "HC-SR04", True), (-1.55, "ECHO GP15", False)):
        f._boxed_text(enc, xim, (y5 + ye) / 2 + dy * 0.165, s,
                      fontweight="bold" if bold else "normal")
    f.card(0.02, y5, 0.95, hh, "START/RESET", "button GP21", role="grey")
    xc = xb[3] + 0.32
    f.card(xc, y5, xb[3] + wb - xc, hh, "CYW43 Wi-Fi", "SPI, on board", role="grey")

    # hardware links (black)
    top, bot = yt + th, yb + bh
    f.arrow([(xh[0], yh - e), (xh[0], top + e)], **hw)             # STOP
    f.arrow([(ct[1], top + e), (ct[1], yh - e)], **hw)             # PWM
    f.arrow([(ct[2], yh - e), (ct[2], top + e)], **hw)             # IR
    f.arrow([(ct[3], yh - e), (ct[3], top + e)], both=True, **hw)  # I2C
    f.arrow([(ct[4], top + e), (ct[4], yh - e)], **hw)             # servo, TRIG
    f.arrow([(0.46, yt - e), (0.46, ys + (sh + hh) / 2 + e)], **hw)  # feed
    f.arrow([(xim, ye + e), (xim, bot - hi - e)], **hw)            # edges, echo
    f.arrow([(0.495, y5 + hh + e), (0.495, yb - e)], **hw)         # START/RESET
    xw = (xc + xb[3] + wb) / 2
    f.arrow([(xw, yb - e), (xw, y5 + hh + e)], both=True, **hw)    # SPI
    # BRAKE: any task that latches a fault brakes; drawn from Safety
    xk, yk = xt[0] + w - 0.12, yh + hh / 2
    f.arrow([(xk, top + e), (xk, yk), (ct[1] - 0.54 - e, yk)], "BRAKE",
            (xk + 0.06, (top + yh) / 2), color="fault", lw=1.4)

    # snapshot copies under PRIMASK (blue)
    for x in [1.16, ct[1], ct[2], ct[3], ct[4]]:
        f.arrow([(x, yt - e), (x, ys + sh + e)], both=True, **data)
    f.arrow([(1.16, bot + e), (1.16, ys - e)], both=True, **data)          # Vehicle
    f.arrow([(xb[1] + wb / 2, bot + e), (xb[1] + wb / 2, ys - e)], both=True,
            **data)                                                          # Telemetry
    f.arrow([(xb[2] + wb / 2, bot + e), (xb[2] + wb / 2, ys - e)], **data)  # Comm
    # the ISR also reads the motor direction and echo state tasks write
    f.arrow([(xim, bot + e), (xim, ys - e)], both=True, **data)

    # kernel message buffers
    qw = 0.40
    qx = xb[2] + wb + (gq - qw) / 2
    for i, (name, col) in enumerate([("RX 4×32 B", "command"),
                                     ("STOP 1×4 B", "fault")]):
        qy = bot - 0.36 - i * 0.36
        f.queue(qx, qy, col, cells=4, w=qw, h=0.14)
        f.label(qx + qw / 2, qy + 0.17, name, ha="center", va="bottom",
                color=FLOW[col])
        f.arrow([(xb[3] - e, qy + 0.07), (qx + qw + e, qy + 0.07)], color=col, lw=1.2)
        f.arrow([(qx - e, qy + 0.07), (xb[2] + wb + e, qy + 0.07)], color=col, lw=1.2)
    ytx = yb - 0.12
    qx = xb[2] + 0.30
    f.queue(qx, ytx - 0.07, "status", cells=4, w=qw, h=0.14)
    f.label(qx + qw / 2, ytx - 0.10, "TX 4×256 B", ha="center", va="top",
            color=FLOW["status"])
    x0 = xb[1] + wb - 0.30
    f.arrow([(x0, yb - e), (x0, ytx), (qx - e, ytx)], color="status", lw=1.2)
    f.arrow([(qx + qw + e, ytx), (xb[3] + 0.16, ytx), (xb[3] + 0.16, yb - e)],
            color="status", lw=1.2)

    # legend and notes
    kw = 3.80
    f.keybox(0.02, y_leg, kw, lh, [
        ("role", "green", "RTOS task"), ("isr", None, "Interrupt (ISR)"),
        ("role", "grey", "Hardware"), ("queue", "hw", "Message buffer"),
        ("line", ("hw", False), "Hardware link"), ("line", ("data", False), "Shared data"),
        ("line", ("command", False), "Command"), ("line", ("status", False), "Telemetry"),
        ("line", ("fault", False), "STOP, BRAKE")], cols=3)
    sw = f._rounded(0.12, y_leg + lh - 0.34 - 0.19 - 0.055, 0.26, 0.11, *ISR_STYLE,
                    lw=0.9, ls=(0, (3, 1.5)), r=0.02, z=3)    # keybox row 2
    f.swatches.append(sw)
    f.notes(0.02 + kw + 0.10, y_leg, W - 0.04 - kw - 0.10, lh, [
        "P1 is the highest priority, P9 the lowest",
        "Single core (SMP = 0): short PRIMASK copies only;\n"
        "   no mutex, semaphore or event flag",
        "Tasks self-timed by tk_dly_tsk; no notifications"])
    f.save()


def fig_system():
    w, h = 1.25, BOXH
    c1, c3 = 1.80, WIDTH - 1.71
    c2 = round((c1 + c3) / 2, 3)
    r2 = 0.60
    r1 = r2 + h + 0.41
    top = r1 + h + 0.26
    lx, lw = 0.06, 1.02
    fx = 1.62
    y_init, y_idle, y_fault = top - 0.04 - h, 0.92, 0.10
    fb = y_idle - 0.05                    # frame: INIT and IDLE (not active)
    f = Fig("system", WIDTH, top + 0.02)
    f.frame(fx, 0.02, WIDTH - fx - 0.02, top - 0.02, "ACTIVE mission states",
            fx + 0.10, top - 0.09)
    f.frame(0.02, fb, lw + 0.08, top - fb)
    f.box(lx, y_init, lw, h, "INIT\nstart-up gates")
    f.box(lx, y_idle, lw, h, "IDLE\nbraked")
    f.box(lx, y_fault, lw, h, "FAULT\nbrake + inhibit", "fault")
    f.box(c1, r1, w, h, "RANGE_WAIT\nhalted ≤ 1 s")
    f.box(c2, r1, w, h, "SCAN\nstopped sweep")
    f.box(c3, r1, w, h, "BYPASS\n7-step detour")
    f.box(c1, r2, w, h, "FOLLOW\nline + pending")
    f.box(c2, r2, w, h, "NAVIGATE\napproach, pivot")
    f.box(c3, r2, w, h, "REACQUIRE\n≤ 0.6 m straight")
    e = 0.012
    mid = lx + lw / 2
    gap = (r2 + h + r1) / 2
    f.arrow([(mid, y_init - e), (mid, y_idle + h + e)], "ready",
            (mid + 0.05, (y_init + y_idle + h) / 2))
    ys = y_idle + 0.12
    f.arrow([(lx + lw + e, ys), (1.42, ys), (1.42, r2 + h / 2), (c1 - e, r2 + h / 2)],
            "START", (lx + lw + 0.09, ys + 0.04), va="bottom")
    # STOP / fault enters FAULT from every state: ACTIVE and INIT/IDLE frames
    yf = y_fault + 0.06
    f.arrow([(fx - 0.004, yf), (lx + lw + e, yf)],
            "STOP /\nfault", (lx + lw + 0.06, yf + 0.04), va="bottom")
    ym = (y_fault + h + fb) / 2
    f.arrow([(0.14, fb - 0.004), (0.14, y_fault + h + e)], "STOP /\nfault", (0.19, ym))
    f.arrow([(0.70, y_fault + h + e), (0.70, y_idle - e)], "local RESET\n(soft only)",
            (0.75, ym + 0.03), dashed=True)
    # inside ACTIVE
    f.arrow([(1.98, r2 + h + e), (1.98, r1 - e)], "range\nabsent", (2.03, gap))
    f.arrow([(2.55, r1 - e), (2.55, r2 + h + e)], "clear +\nline", (2.60, gap))
    f.arrow([(c1 + w + e, r1 + 0.23), (c2 - e, r1 + 0.23)], "obstacle",
            ((c1 + w + c2) / 2, r1 + 0.26), ha="center", va="bottom")
    f.arrow([(c1 + w + e, r2 + 0.26), (3.40, r2 + 0.26), (3.40, r1 + 0.09),
             (c2 - e, r1 + 0.09)], "obstacle", (3.45, gap))
    f.arrow([(c1 + w + e, r2 + 0.09), (c2 - e, r2 + 0.09)], "pending +\njunction",
            ((c1 + w + c2) / 2, r2 + 0.05), ha="center", va="top")
    f.arrow([(c2 + w + e, r1 + h / 2), (c3 - e, r1 + h / 2)], "qualified\npath",
            ((c2 + w + c3) / 2, r1 + h / 2 - 0.05), ha="center", va="top")
    f.arrow([(c2 + w + e, r2 + h / 2), (c3 - e, r2 + h / 2)], "action\ndone",
            ((c2 + w + c3) / 2, r2 + h / 2 - 0.05), ha="center", va="top")
    x3 = c3 + 0.62
    f.arrow([(x3, r1 - e), (x3, r2 + h + e)], "detour\ndone", (x3 + 0.05, gap))
    yl = 0.23
    f.arrow([(c3 + 0.34, r2 - e), (c3 + 0.34, yl), (2.40, yl), (2.40, r2 - e)],
            "verified original line (100 ms lock)", (4.05, yl - 0.035),
            ha="center", va="top")
    f.save()


def tpl_system():
    """Fig. 3 for the course template, in its Section 5 statechart style:
    start dot, STANDBY (INIT, IDLE) and ACTIVE MISSION panels, red fault
    transitions, dashed recovery, legend and notes; TPL_W wide."""
    W, bh, bw, e = TPL_W, 0.48, 1.10, 0.012
    k = "hw"                               # normal transitions: near black
    kh, head = 1.02, 0.42                  # legend/notes band; panel headers
    yd = 0.02 + kh + 0.10                  # FAULT and ACTIVE panel bottom
    # left column: FAULT, then the STANDBY panel (IDLE, INIT, start dot)
    sx, sw = 0.02, 1.30
    lx = sx + 0.10
    lw = sw - 0.20
    y_fault = yd
    sb = y_fault + bh + 0.40               # STANDBY panel bottom
    y_idle = sb + 0.10
    y_init = y_idle + bh + 0.48
    y_dot = y_init + bh + 0.20
    pt = y_dot + 0.06 + 0.06 + head        # both panel tops, equal headers
    H = pt + 0.48                          # title and subtitle
    # ACTIVE MISSION panel: 2 x 3 grid; its rows line up with INIT and IDLE.
    # The first column gap is wider: it also holds the FOLLOW -> SCAN riser.
    ax0 = sx + sw + 0.50
    aw = W - 0.02 - ax0
    c1 = ax0 + 0.10
    c2 = c1 + bw + 0.80
    c3 = c2 + bw + 0.76
    r1, r2 = y_init, y_idle

    f = Fig("system", W, H)
    f.title("AUTONOMOUS ROBOT CAR – SYSTEM STATECHART",
            "(Vehicle task mission states; the Safety task has precedence)")
    f.panel(sx, sb, sw, pt - sb, "STANDBY", "(robot halted)", role="green",
            head=head)
    f.panel(ax0, yd, aw, pt - yd, "ACTIVE MISSION",
            "(autonomous run; task health, tilt and collision supervised)",
            role="blue", head=head)
    f.start(lx + lw / 2, y_dot)
    f.state(lx, y_init, lw, bh, "INIT", "(start-up checks)", "green")
    f.state(lx, y_idle, lw, bh, "IDLE", "(halted, ready)", "green")
    f.state(lx, y_fault, lw, bh, "FAULT", "(brake + inhibit)", "red")
    f.state(c1, r1, bw, bh, "RANGE_WAIT", "(halted, ≤ 1 s)")
    f.state(c2, r1, bw, bh, "SCAN", "(stopped sweep)")
    f.state(c3, r1, bw, bh, "BYPASS", "(7-step detour)")
    f.state(c1, r2, bw, bh, "FOLLOW", "(line following)")
    f.state(c2, r2, bw, bh, "NAVIGATE", "(approach + pivot)")
    f.state(c3, r2, bw, bh, "REACQUIRE", "(≤ 0.6 m, ≤ 10 s)")

    mx = lx + lw / 2
    f.arrow([(mx, y_dot - 0.06 - e), (mx, y_init + bh + e)], color=k)
    f.arrow([(mx, y_init - e), (mx, y_idle + bh + e)], "checks\npass",
            (mx + 0.05, (y_init + y_idle + bh) / 2), color=k)
    yc = r2 + bh / 2
    f.arrow([(lx + lw + e, yc), (c1 - e, yc)], "START",
            ((sx + sw + ax0) / 2, yc + 0.03), ha="center", va="bottom", color=k)
    # fault transitions (red) into FAULT and the recovery (dashed) out of it
    ym = (y_fault + bh + sb) / 2
    f.arrow([(0.20, sb - 0.004), (0.20, y_fault + bh + e)], "STOP /\nhard fault",
            (0.25, ym), color="fault")
    xr = 0.86
    f.arrow([(xr, y_fault + bh + e), (xr, y_idle - e)], "local RESET\n[soft fault]",
            (xr + 0.05, ym), color=k, dashed=True)
    yf = y_fault + bh - 0.07
    f.arrow([(ax0 - 0.004, yf), (lx + lw + e, yf)], "STOP,\nfault or\ntimeout",
            (lx + lw + 0.06, yf - 0.03), va="top", color="fault")
    # inside ACTIVE MISSION
    gm = (r2 + bh + r1) / 2
    g1, g2 = (c1 + bw + c2) / 2, (c2 + bw + c3) / 2
    xa, xb, xo = c1 + 0.10, c1 + 0.65, c1 + bw + 0.27
    f.arrow([(xa, r2 + bh + e), (xa, r1 - e)], "no fresh\nrange", (xa + 0.05, gm),
            color=k)
    f.arrow([(xb, r1 - e), (xb, r2 + bh + e)], "clear range\n[line valid]",
            (xb + 0.05, gm), color=k)
    f.arrow([(c1 + bw + e, r1 + 0.36), (c2 - e, r1 + 0.36)], "obstacle\n[at rest]",
            (g1, r1 + 0.39), ha="center", va="bottom", color=k)
    f.arrow([(c1 + bw + e, r2 + 0.36), (xo, r2 + 0.36), (xo, r1 + 0.12),
             (c2 - e, r1 + 0.12)], "obstacle\n≤ 0.25 m", (xo + 0.05, gm), color=k)
    f.arrow([(c1 + bw + e, r2 + 0.14), (c2 - e, r2 + 0.14)],
            "junction\n[barcode\npending]", (g1, r2 + 0.10), ha="center", va="top",
            color=k)
    f.arrow([(c2 + bw + e, r1 + bh / 2), (c3 - e, r1 + bh / 2)],
            "sweep done\n[route\nqualified]", (g2, r1 + bh / 2 - 0.03), ha="center",
            va="top", color=k)
    f.arrow([(c2 + bw + e, r2 + bh / 2), (c3 - e, r2 + bh / 2)], "action\ndone",
            (g2, r2 + bh / 2 - 0.03), ha="center", va="top", color=k)
    x3 = c3 + bw / 2
    f.arrow([(x3, r1 - e), (x3, r2 + bh + e)], "detour\ndone", (x3 + 0.05, gm),
            color=k)
    yl = r2 - 0.58
    f.arrow([(c3 + 0.30, r2 - e), (c3 + 0.30, yl), (c1 + 0.40, yl),
             (c1 + 0.40, r2 - e)], "line verified on route [locked ≥ 100 ms]",
            ((c1 + c3 + bw) / 2, yl - 0.035), ha="center", va="top", color=k)

    # legend and notes
    lgw = 3.04
    f.keybox(0.02, 0.02, lgw, kh, [
        ("start", None, "Initial state"),
        ("role", "green", "Standby state"),
        ("role", "blue", "Active mission state"),
        ("role", "red", "Fault state"),
        ("line", (k, False), "Normal transition"),
        ("line", (k, True), "Conditional (RESET)"),
        ("line", ("fault", False), "Fault transition")], cols=2)
    f.notes(lgw + 0.12, 0.02, W - 0.02 - lgw - 0.12, kh, [
        "Hard faults (e.g. tilt > 0.60 rad, collision) latch until reboot.",
        "Local RESET clears soft faults only, at rest with sensors ready.",
        "Safety has precedence: STOP brakes from every state.",
        "Arrows from a panel apply to all its states; labels: event [guard]."])
    f.save()


def fig_comm():
    w, h = 0.98, BOXH
    g = (SIDE_FIG - 0.04 - 4 * w) / 3
    xs = [0.02 + i * (w + g) for i in range(4)]
    r2 = 0.04
    r1 = r2 + h + 0.37
    f = Fig("comm", SIDE_FIG, r1 + h + 0.20)
    # The 30 s watchdog ends at CONNACK; SUBACK has lwIP's own 15 s request
    # timeout (drawn in tpl_comm, which has the room).
    for x, s in zip(xs, ["SETUP\n≤ 3 attempts", "JOINING\n30 s, then 1–30 s",
                         "CONNECTING\nCONNACK ≤ 30 s",
                         "SUBSCRIBED\nTX ≤ 10/s, RX on"]):
        f.box(x, r1, w, h, s)
    f.box(xs[0], r2, w, h, "UNAVAILABLE\nuntil reboot", "fault")
    f.box(xs[1], r2, xs[3] + w - xs[1], h,
          "RETRY_WAIT · MQTT backoff 1 s → 30 s\n"
          "offline: newest TX frame kept, cmds ignored")
    e = 0.012
    ym = (r1 + r2 + h) / 2
    for i, s in enumerate(["STA up", "link + IP", "SUBACK"]):
        f.arrow([(xs[i] + w + e, r1 + h / 2), (xs[i + 1] - e, r1 + h / 2)],
                s, (xs[i] + w + g / 2, r1 + h + 0.025), ha="center", va="bottom")
    f.arrow([(0.30, r1 - e), (0.30, r2 + h + e)], "3rd fail", (0.35, ym))
    xn = xs[1] + 0.25                    # Wi-Fi link loss: rejoin
    f.arrow([(xn, r2 + h + e), (xn, r1 - e)], "no link", (xn + 0.05, ym))
    xa = xs[2] + 0.22
    f.arrow([(xa, r1 - e), (xa, r2 + h + e)], "refused /\ntimeout", (xa - 0.05, ym),
            ha="right")
    xb = xs[2] + w - 0.20
    f.arrow([(xb, r2 + h + e), (xb, r1 - e)], "retry\ndue", (xb + 0.05, ym))
    xc = xs[3] + w - 0.16
    f.arrow([(xc, r1 - e), (xc, r2 + h + e)], "broker or\nlink lost", (xc - 0.05, ym),
            ha="right")
    f.save()


def tpl_comm():
    """Fig. 4 for the course template: one panel of its subsystem statecharts.
    The radio owner's session states, from the port's robot_radio_setup()
    (SETUP) and the net_mqtt.c flags; notes link Comm P7 and Telemetry P8."""
    W, H = TPL_W, 4.50
    f = Fig("comm", W, H)
    # The picture box shows about 6.88 in of the image: keep the frame inset.
    pad = 0.07
    top = f.panel(pad, 0.03, W - 2 * pad, H - 0.06, "COMMUNICATIONS",
                  "(radio owner P9 · Comm P7 · Telemetry P8 · Buddy 1)", head=0.50)
    L, R, e = pad + 0.10, W - pad - 0.10, 0.012

    def go(pts, text=None, at=None, ha="left", va="center", kind="hw", dashed=False):
        """Transition; its label in the arrow's colour with open leading."""
        f.arrow(pts, color=kind, dashed=dashed, lw=1.0)
        if text:
            f.label(at[0], at[1], text, ha=ha, va=va, color=FLOW[kind],
                    linespacing=1.15)

    # states: one column entered from the initial dot; the trunk's labels
    # sit on its left so the right corridor carries only RETRY_WAIT's edges
    x0, sw, sh, gap = 2.14, 1.92, 0.54, 0.40
    cx, x1 = x0 + sw / 2, x0 + sw
    y_set, y_join, y_conn, y_sub = [top - 0.36 - sh - i * (sh + gap) for i in range(4)]
    f.card(x0, y_set, sw, sh, "SETUP",
           "(start Wi-Fi station interface;\n≤ 3 attempts, 1 s then 2 s apart)")
    f.card(x0, y_join, sw, sh, "JOINING",
           "(join the Wi-Fi access point;\n30 s window per attempt)")
    f.card(x0, y_conn, sw, sh, "CONNECTING",
           "(MQTT connect: CONNACK ≤ 30 s,\nsubscribe: SUBACK ≤ 15 s)")
    f.card(x0, y_sub, sw, sh, "SUBSCRIBED",
           "(publish telemetry ≤ 10 frames/s\nin order; accept commands)", role="green")
    f.start(cx, top - 0.15)
    go([(cx, top - 0.21), (cx, y_set + sh + e)])
    for ya, yb, s in ((y_set, y_join, "Station interface up"),
                      (y_join, y_conn, "Link + IPv4 up\n[backoff elapsed]"),
                      (y_conn, y_sub, "SUBACK received\n/ publish “online”")):
        go([(cx, ya - e), (cx, yb + sh + e)], s, (cx - 0.07, (ya + yb + sh) / 2),
           ha="right")

    # setup exhausted: fault state, left
    uw = 1.08
    f.card(L, y_set, uw, sh, "UNAVAILABLE", "(no radio service\nuntil reboot)",
           role="red")
    yu = y_set + 0.19
    go([(x0 - e, yu), (L + uw + e, yu)], "Third attempt\nfails",
       ((L + uw + x0) / 2, yu + 0.05), ha="center", va="bottom", kind="fault")
    # JOINING: the 30 s window ends, leave and rejoin (self-loop, left)
    xl = x0 - 0.24
    go([(x0 - e, y_join + sh - 0.12), (xl, y_join + sh - 0.12), (xl, y_join + 0.12),
        (x0 - e, y_join + 0.12)], "No link in 30 s / leave,\nrejoin after 1 → 30 s",
       (xl - 0.06, y_join + sh / 2), ha="right", dashed=True)

    # MQTT retry: RETRY_WAIT, right
    rw = 1.54
    rx = R - rw
    f.card(rx, y_conn, rw, sh, "RETRY_WAIT",
           "(MQTT backoff 1 → 30 s,\ndoubling; SUBACK resets)")
    mid = (x1 + rx) / 2
    yb = y_conn + sh - 0.13
    go([(rx - e, yb), (x1 + e, yb)], "Backoff elapsed\n[link + IPv4 up]",
       (mid, yb + 0.05), ha="center", va="bottom", dashed=True)
    ya = y_conn + 0.13
    go([(x1 + e, ya), (rx - e, ya)], "Refused, timeout,\nSUBACK failure,\nlink or IP lost",
       (mid, ya - 0.05), ha="center", va="top", kind="fault")
    xc, yc = R - 0.30, y_sub + sh / 2
    go([(x1 + e, yc), (xc, yc), (xc, y_conn - e)],
       "Broker or keepalive lost,\npublish error, link or IP lost",
       ((x1 + xc) / 2, yc - 0.05), ha="center", va="top", kind="fault")
    xd, yd = rx + 0.36, y_join + sh / 2
    go([(xd, y_conn + sh + e), (xd, yd), (x1 + e, yd)], "[Wi-Fi link down]",
       ((x1 + xd) / 2, yd + 0.05), ha="center", va="bottom", dashed=True)

    # legend (top right) and notes with hanging bullets (bottom left)
    lx = x1 + 0.30
    f.keybox(lx, y_set - 0.26, R - lx, top - 0.08 - (y_set - 0.26),
             [("start", None, "Initial state"), ("role", "blue", "Offline state"),
              ("role", "green", "Online state"), ("role", "red", "Fault state"),
              ("line", ("hw", False), "Transition"),
              ("line", ("fault", False), "Fault / loss"),
              ("line", ("hw", True), "Retry / rejoin")], cols=2)
    # notes box: its right edge stays 0.10 in clear of the trunk labels
    # (at most 0.95 in wide) that sit left of the state column
    ny = 0.13
    nw, nh = cx - 0.07 - 0.95 - 0.10 - L, y_join - 0.07 - ny
    patch = f._rounded(L, ny, nw, nh, *ROLE["note"], lw=0.9)
    f._boxed_text(patch, L + 0.10, ny + nh - 0.08, "Notes", ha="left", va="top",
                  fontweight="bold")
    yy = ny + nh - 0.32
    inv = f.ax.transData.inverted().transform
    for s in ("After SETUP, offline states\nkeep the newest telemetry\n"
              "frame; commands ignored.",
              "Comm P7 acts on STOP\nfirst; other commands\nare only counted.",
              "Telemetry P8 queues\n7 frames/s; only full-queue\ndrops are counted.",
              "Remote STOP is best effort;\nthe GP20 button is primary."):
        f._boxed_text(patch, L + 0.10, yy, "•", ha="left", va="top")
        t = f._boxed_text(patch, L + 0.22, yy, s, ha="left", va="top",
                          linespacing=1.2)
        f.fig.canvas.draw()
        yy = inv(t.get_window_extent(f.fig.canvas.get_renderer()).p0)[1] - 0.065
    f.save()


def fig_motion():
    w, h = 1.00, BOXH
    hb = 0.215                      # one-line INHIBITED box: room for the arc
    xs = [0.02, 1.64, 3.26]
    r2 = 0.03
    r1 = r2 + hb + 0.36
    ytop = r1 + h + 0.10            # HALT arc, clear of the "new ID" labels
    f = Fig("motion", SIDE_FIG, ytop + 0.16)     # same 1.21 in as before
    f.box(xs[0], r1, w, h, "IDLE\nzero demand")
    f.box(xs[1], r1, w, h, "RUNNING\nPID, 50 ms lease")
    f.box(xs[2], r1, w, h, "COMPLETE\nzero demand")
    f.box(xs[0], r2, xs[2] + w - xs[0], hb,
          "INHIBITED · zero duty, latched until an eligible reset", "fault")
    yc = r1 + h / 2
    e = 0.012
    for a, b, top, bot in [(0, 1, "new ID", "HALT"), (1, 2, "reached", "new ID")]:
        xa, xb = xs[a] + w + e, xs[b] - e
        mx = (xs[a] + w + xs[b]) / 2
        f.arrow([(xa, yc + 0.07), (xb, yc + 0.07)], top, (mx, yc + 0.10),
                ha="center", va="bottom")
        f.arrow([(xb, yc - 0.07), (xa, yc - 0.07)], bot, (mx, yc - 0.10),
                ha="center", va="top")
    # HALT also returns COMPLETE to IDLE (motion.c: any unlatched state)
    xl, xr = xs[0] + 0.70, xs[2] + w / 2
    f.arrow([(xr, r1 + h + e), (xr, ytop), (xl, ytop), (xl, r1 + h + e)], "HALT",
            (xs[1] + w / 2, ytop + 0.02), ha="center", va="bottom")
    ym = (r1 + r2 + hb) / 2
    f.arrow([(0.12, r1 - e), (0.12, r2 + hb + e)], "STOP,\nsensor", (0.17, ym))
    f.arrow([(0.72, r2 + hb + e), (0.72, r1 - e)], "eligible\nreset",
            (0.77, ym), dashed=True)
    x1 = xs[1] + w / 2
    f.arrow([(x1, r1 - e), (x1, r2 + hb + e)], "STOP, expiry,\ntimeout, sensor",
            (x1 + 0.05, ym))
    x3 = xs[2] + w / 2
    f.arrow([(x3, r1 - e), (x3, r2 + hb + e)], "STOP,\nsensor", (x3 + 0.05, ym))
    f.save()


def tpl_motion():
    """Fig. 5 for the course template: one statechart panel in the style of
    the template's example statecharts (header band, initial state, choice,
    pastel states with internal transitions, red fault transitions, dashed
    guarded reset, notes and legend)."""
    from matplotlib.patches import Polygon
    W, e, lw = TPL_W, 0.012, 1.1
    hn, hi, hs = 1.18, 0.46, 0.84             # notes, INHIBITED, state heights
    yn = 0.12
    yi = yn + hn + 0.12
    ys = yi + hi + 0.52
    T = ys + hs                                # top edge of the state row
    H = T + 0.72 + 0.46 + 0.02                 # start dot clear of the header
    f = Fig("motion", W, H)
    f.panel(0.02, 0.02, W - 0.04, H - 0.04, "MOTION CONTROL",
            "(Motion task P2 · Buddy 2)", role="yellow", head=0.46)
    # states: IDLE, RUNNING, COMPLETE in a row, INHIBITED below all three
    xs, ws = [0.34, 2.60, 5.10], [1.56, 1.80, 1.66]
    f.card(xs[0], ys, ws[0], hs, "IDLE", "(zero duty, waits\nfor a new ID)",
           ["same ID / refused"], role="yellow")
    f.card(xs[1], ys, ws[1], hs, "RUNNING", "(PID wheel speed, 50 ms lease)",
           ["same ID / renew lease", "newer ID / new primitive"], role="yellow")
    f.card(xs[2], ys, ws[2], hs, "COMPLETE", "(distance or turn done,\nzero duty)",
           ["same ID / no-op, no replay"], role="yellow")
    xi, wi = 0.10, xs[2] + ws[2] - 0.10
    f.card(xi, yi, wi, hi, "INHIBITED",
           "(zero duty, latched: motion requests refused, HALT keeps the latch)",
           role="red")
    # initial state and boot choice: the calibration and wiring gate
    xb, yd = 0.18, T + 0.30
    f.start(xb, T + 0.56)
    dia = Polygon([(xb - 0.10, yd), (xb, yd + 0.08), (xb + 0.10, yd), (xb, yd - 0.08)],
                  closed=True, fc="white", ec="#111111", lw=1.2, zorder=3)
    f.ax.add_patch(dia)
    f.swatches.append(dia)
    f.arrow([(xb, T + 0.50 - e), (xb, yd + 0.08 + e)], color="hw", lw=lw)
    xg = 0.62
    f.arrow([(xb + 0.10 + e, yd), (xg, yd), (xg, T + e)],
            "[calibrated]", (xg + 0.05, T + 0.13), color="hw", lw=lw)
    ym = (yi + hi + ys) / 2
    f.arrow([(xb, yd - 0.08 - e), (xb, yi + hi + e)], "[else]", (xb + 0.05, ym),
            color="hw", lw=lw)
    # normal transitions between the three unlatched states
    yc = ys + hs / 2
    for a, b, top, bot in [(0, 1, "new ID", "HALT"),
                           (1, 2, "target\nreached", "new ID")]:
        xa, xz = xs[a] + ws[a] + e, xs[b] - e
        mx = (xs[a] + ws[a] + xs[b]) / 2
        f.arrow([(xa, yc + 0.10), (xz, yc + 0.10)], top, (mx, yc + 0.13),
                ha="center", va="bottom", color="hw", lw=lw)
        f.arrow([(xz, yc - 0.10), (xa, yc - 0.10)], bot, (mx, yc - 0.13),
                ha="center", va="top", color="hw", lw=lw)
    yh = T + 0.38                              # HALT from COMPLETE, over the top
    xl, xr = xs[0] + ws[0] - 0.14, xs[2] + ws[2] / 2
    f.arrow([(xr, T + e), (xr, yh), (xl, yh), (xl, T + e)], "HALT",
            ((xl + xr) / 2, yh + 0.03), ha="center", va="bottom", color="hw", lw=lw)
    # STOP and fault transitions (red), guarded local reset (dashed)
    ti = yi + hi + e
    f.arrow([(0.72, ys - e), (0.72, ti)], "STOP,\nsensor fault", (0.77, ym),
            color="fault", lw=lw)
    f.arrow([(1.55, ti), (1.55, ys - e)],
            "local RESET\n[stationary, STOP released,\nno hard fault, calibrated]",
            (1.60, ym), dashed=True, color="hw", lw=lw)
    f.arrow([(3.20, ys - e), (3.20, ti)],
            "STOP, lease expiry (50 ms),\ndistance or turn timeout (15 s),\n"
            "sensor fault", (3.25, ym), color="fault", lw=lw)
    f.arrow([(5.45, ys - e), (5.45, ti)], "STOP,\nsensor fault", (5.50, ym),
            color="fault", lw=lw)
    # notes and legend
    wk = 2.50
    xk = xi + wi - wk
    f.notes(xi, yn, xk - 0.10 - xi, hn, [
        "Only Vehicle P4 sends requests; it renews the ID every 10 ms.",
        "Lease expiry, timeout and sensor fault are hard faults (until reboot).",
        "Outside STOP, any refused motion request is a hard fault too.",
        "STOP = STOP button, remote STOP, Vehicle FAULT or any fault.",
        "Sensor fault: step gap over 25 ms, or a non-finite speed or PID value.",
        "Calibrated = calibration and wiring confirmed (not yet, so INHIBITED)."])
    rows = [("start", None, "Initial state"), ("choice", None, "Choice [guard]"),
            ("role", "yellow", "Normal state"), ("role", "red", "Latched state"),
            ("line", ("hw", False), "Transition"),
            ("line", ("fault", False), "STOP / fault"),
            ("line", ("hw", True), "Guarded reset")]
    f.keybox(xk, yn, wk, hn, rows, cols=2)
    cx, cy = xk + 0.10 + 0.13, yn + hn - 0.34 - 0.19      # row 2: the choice
    key = Polygon([(cx - 0.07, cy), (cx, cy + 0.055), (cx + 0.07, cy),
                   (cx, cy - 0.055)], closed=True, fc="white", ec="#111111",
                  lw=1.0, zorder=3)
    f.ax.add_patch(key)
    f.swatches.append(key)
    f.save()


def fig_optical():
    w, h = 0.96, BOXH
    xs = [0.14, 1.72, 3.30]
    r2, yl = 0.36, 0.21
    r1 = r2 + h + 0.25
    f = Fig("optical", SIDE_FIG, r1 + h + 0.03)
    f.box(xs[0], r1, w, h, "TRACKING\nline_ctrl")
    f.box(xs[1], r1, w, h, "FOLLOW +\npending_nav")
    f.box(xs[2], r1, w, h, "NAVIGATE\napproach, pivot")
    f.box(xs[0], r2, w, h, "SEARCHING\nweak signal")
    f.box(xs[1], r2, w, h, "FAULT\nvehicle, brake", "fault")
    f.box(xs[2], r2, w, h, "REACQUIRE\n≤ 0.6 m, ≤ 10 s")
    e = 0.012
    y1, y2 = r1 + h / 2, r2 + h / 2
    f.arrow([(xs[0] + w + e, y1), (xs[1] - e, y1)], "DECODED",
            ((xs[0] + w + xs[1]) / 2, y1 + 0.03), ha="center", va="bottom")
    f.arrow([(xs[1] + w + e, y1), (xs[2] - e, y1)], "junction",
            ((xs[1] + w + xs[2]) / 2, y1 + 0.03), ha="center", va="bottom")
    f.arrow([(xs[0] + w + e, y2), (xs[1] - e, y2)], "LOST",
            ((xs[0] + w + xs[1]) / 2, y2 - 0.03), ha="center", va="top")
    f.arrow([(xs[2] - e, y2), (xs[1] + w + e, y2)], "no line /\ntimeout",
            ((xs[1] + w + xs[2]) / 2, y2 - 0.03), ha="center", va="top")
    xm = xs[1] + w / 2
    f.arrow([(xm, r1 - e), (xm, r2 + h + e)], "2nd code\nor 10 s", (xm + 0.05, (r1 + r2 + h) / 2))
    x0 = xs[0] + w / 2
    f.arrow([(x0 - 0.16, r1 - e), (x0 - 0.16, r2 + h + e)])
    f.arrow([(x0 + 0.16, r2 + h + e), (x0 + 0.16, r1 - e)], "signal",
            (x0 + 0.21, (r1 + r2 + h) / 2))
    x2 = xs[2] + w / 2
    f.arrow([(x2, r1 - e), (x2, r2 + h + e)], "done", (x2 + 0.05, (r1 + r2 + h) / 2))
    f.arrow([(x2 + 0.25, r2 - e), (x2 + 0.25, yl), (0.06, yl), (0.06, y1),
             (xs[0] - e, y1)], "verified original line",
            ((xs[1] + xs[2] + w) / 2, yl - 0.035), ha="center", va="top")
    f.save()


def tpl_optical():
    """Fig. 6 for the course template: one statechart panel in the style of
    the template's example statecharts. FOLLOW is a composite state with two
    concurrent regions, the line evidence (line_ctrl status) and the route
    code (barcode decoder); a pending code is acted on at the next junction
    edge by NAVIGATE, then REACQUIRE must lock onto the line ahead (the
    route anchor is the pose after the manoeuvre)."""
    W, H, e = TPL_W, 4.50, 0.012
    f = Fig("optical", W, H)
    f.panel(0.02, 0.02, W - 0.04, H - 0.04, "OPTICAL NAVIGATION",
            "(Line P3 · Vehicle P4 · Buddy 3)", role="green", head=0.46)

    def go(pts, text=None, at=None, ha="left", va="center", kind="hw"):
        """Transition; its label in the arrow's colour with open leading."""
        f.arrow(pts, color=kind, lw=1.1)
        if text:
            f.label(at[0], at[1], text, ha=ha, va=va, color=FLOW[kind],
                    linespacing=1.15)

    # FOLLOW: composite state, two concurrent regions split by a dashed line
    sx0, sx1, sy0, sy1, xd = 0.16, 4.20, 1.96, 3.72, 2.36
    edge = ROLE["green"][1]
    f.frames.append(f._rounded(sx0, sy0, sx1 - sx0, sy1 - sy0, "#f7fcf7", edge,
                               lw=1.1, r=0.06, z=1))
    t = f.label(0.30, sy1 - 0.12, "FOLLOW", fontweight="bold", color=edge)
    f.fig.canvas.draw()
    x_t = f.ax.transData.inverted().transform(
        t.get_window_extent(f.fig.canvas.get_renderer()).p1)[0]
    f.label(x_t + 0.06, sy1 - 0.12,
            "(Vehicle P4: line and route-code regions run concurrently)")
    h, ya, yb = 0.50, 3.00, 2.08               # state rows A (top) and B
    f.ax.plot([xd, xd], [sy0 + 0.06, ya + h + 0.02], color=edge, lw=1.0,
              ls=(0, (5, 3)), zorder=1)
    lx, lw_ = 0.48, 1.72                       # line evidence region
    rx, rw = 2.70, 1.42                        # route code region
    f.card(lx, ya, lw_, h, "TRACKING / JUNCTION", "(steer on the line edge)",
           role="green")
    f.card(lx, yb, lw_, h, "SEARCHING", "(uncertain line: halt, no fault)",
           role="green")
    f.card(rx, ya, rw, h, "NO CODE", "(junction edge ignored)", role="green")
    f.card(rx, yb, rw, h, "CODE PENDING", "(hold the action ≤ 10 s)",
           role="green")
    # the coded junction: NAVIGATE, then REACQUIRE; FAULT below everything
    cx, cw, ch = 4.84, 1.74, 0.60
    f.card(cx, yb + h - ch, cw, ch, "NAVIGATE",
           "(0.060 m approach, then pivot;\nSTRAIGHT: approach only)", role="blue")
    f.card(cx, ya, cw, ch, "REACQUIRE",
           "(drive straight ≤ 0.6 m, ≤ 10 s,\nuntil the line ahead is found)",
           role="blue")
    fy, fh = 1.16, 0.42
    f.card(sx0, fy, 6.76 - sx0, fh, "FAULT (soft)",
           "(brake and inhibit; an eligible local RESET → IDLE, then a new START)",
           role="red")

    # initial states: the chart's (START from IDLE) and one per region
    xs0, ys0 = 0.40, 3.91
    f.start(xs0, ys0)
    go([(xs0, ys0 - 0.06 - e), (xs0, sy1 + e)])
    f.label(xs0 + 0.12, ys0,
            "from IDLE: START [sensors ready, line valid, path clear]",
            color=FLOW["hw"])
    for x0, xb in ((0.30, lx), (2.52, rx)):
        f.start(x0, ya + h / 2, r=0.05)
        go([(x0 + 0.05 + e, ya + h / 2), (xb - e, ya + h / 2)])

    # line evidence region (line_ctrl status, Line P3)
    ym = (ya + yb + h) / 2
    go([(0.84, ya - e), (0.84, yb + h + e)], "weak, or\nboth black",
       (0.79, ym), ha="right")
    go([(1.04, yb + h + e), (1.04, ya - e)],
       "signal back, or\nboth black 20 ms\n/ junction edge", (1.09, ym))
    # route code region (barcode decoder, only while FOLLOW)
    xr = 2.78
    go([(xr, ya - e), (xr, yb + h + e)], "DECODED: LEFT, RIGHT,\nSTRAIGHT or U-TURN",
       (xr + 0.05, ym))
    # coded junction and the way back into FOLLOW
    yc = yb + h / 2
    go([(rx + rw + e, yc), (cx - e, yc)], "junction\nedge",
       ((sx1 + cx) / 2, yc + 0.04), ha="center", va="bottom")
    xn = cx + cw / 2
    go([(xn, yb + h + e), (xn, ya - e)], "manoeuvre\ncomplete", (xn + 0.05, ym))
    yq, xe = 3.83, 3.96
    go([(xn, ya + ch + e), (xn, yq), (xe, yq), (xe, sy1 + e)],
       "line ahead verified (100 ms lock)", ((xe + xn) / 2 + 0.10, yq + 0.03),
       ha="center", va="bottom")

    # fault transitions (red) into FAULT
    yt = fy + fh + e
    yl = (sy0 + fy + fh) / 2
    go([(0.84, yb - e), (0.84, yt)], "line LOST\n(weak ≥ 100 ms)", (0.89, yl),
       kind="fault")
    go([(xr, yb - e), (xr, yt)], "second code or\n10 s pending", (xr + 0.05, yl),
       kind="fault")
    xv = 5.04
    go([(xv, yb + h - ch - e), (xv, yt)], "obstacle or\nrange lost",
       (xv - 0.05, yl), ha="right", kind="fault")
    xc = 6.68
    go([(cx + cw + e, ya + ch / 2), (xc, ya + ch / 2), (xc, yt)],
       "no line in 0.6 m or 10 s;\nobstacle or range lost", (xc - 0.05, yl),
       ha="right", kind="fault")

    # notes and legend
    ky, kh = 0.08, 1.00
    f.notes(sx0, ky, 3.96, kh, [
        "Line P3: IR every 2 ms. Vehicle P4: Motion P2 requests every 10 ms.",
        "Unknown or rejected barcode frames are only counted (no action).",
        "Bad IR sample or gap > 50 ms while active: hard fault until reboot.",
        "STOP in any state → FAULT; an obstacle in FOLLOW → SCAN."])
    kx, kw = 4.22, 2.54
    f.keybox(kx, ky, kw, kh, [
        ("start", None, "Initial state"), ("role", "green", "FOLLOW state"),
        ("role", "blue", "Junction state"), ("role", "red", "Fault state"),
        ("line", ("hw", False), "Transition"),
        ("line", ("fault", False), "Fault transition"),
        ("divider", None, "Region divider")], cols=2)
    xk, yk = kx + 0.10 + (kw - 0.10) / 2, ky + kh - 0.34 - 2 * 0.19
    f.ax.plot([xk, xk + 0.26], [yk, yk], color=edge, lw=1.0, ls=(0, (5, 3)),
              zorder=3)
    f.save()


def fig_terrain():
    h = BOXH
    fy = 0.03 + h + 0.30
    ty = fy + 0.27
    ft = ty + h + 0.19
    f = Fig("terrain", SIDE_FIG, ft + 0.02)
    f.frame(0.02, fy, 4.26, ft - fy, "FRESH samples (an overrun is the newest sample)",
            0.12, ft - 0.095)
    # Hump phases (terrain.c terrain_hump), not the terrain_motion_t class.
    lvl = f.box(0.10, ty, 1.54, h, "NO HUMP OPEN\nwatching for a rise")
    f.box(1.98, ty, 0.98, h, "HUMP OPEN\nintegrate")
    f.box(3.30, ty, 0.90, h, "DESCENT SEEN\nkeep peak")
    yc = ty + h / 2
    e = 0.012
    f.arrow([(1.64 + e, yc), (1.98 - e, yc)])
    f.arrow([(2.96 + e, yc), (3.30 - e, yc)])
    f.label(1.81, yc + 0.03, "up", ha="center", va="bottom")
    f.label(3.13, yc + 0.03, "down", ha="center", va="bottom")
    yp = ty - 0.08
    f.arrow([(3.75, ty - e), (3.75, yp), (lvl[0] + 0.77, yp), (lvl[0] + 0.77, ty - e)],
            "flat 0.2 s: hump closes; peak 5–300 mm, end ±30 mm → run peak",
            (2.25, yp - 0.03), ha="center", va="top")
    f.box(0.02, 0.03, 1.46, h, "INVALID: I2C error,\nsaturation, status → hard",
          "fault")
    f.box(1.58, 0.03, 1.24, h, "TILT_UNAVAILABLE:\nsoft stop", "soft")
    f.box(2.94, 0.03, 1.32, h, "TILT > 0.60 rad or\nCOLLISION → hard", "fault")
    ym = (fy + 0.03 + h) / 2
    for x, s in [(0.52, "bus / status"), (1.66, "PENDING > 50 ms\nor unusable tilt"),
                 (3.72, "tilt,\nimpact")]:
        f.arrow([(x, fy - 0.004), (x, 0.03 + h + e)], s, (x + 0.05, ym))
    f.arrow([(2.72, 0.03 + h + e), (2.72, fy - 0.004)], "fresh tilt +\nlocal RESET",
            (2.77, ym), dashed=True)
    f.save()


def tpl_terrain():
    """Section 6.4 in the course template's statechart style (one panel)."""
    e, sh = 0.012, 0.46                           # arrow gap, state height
    X0, XW = 0.14, TPL_W - 0.28                   # content left edge, width
    nh = 1.00                                     # notes and legend height
    yf = 0.08 + nh + 0.05                         # fault states: bottom
    yb = yf + sh + 0.52                           # tracking state: bottom, top
    yt = yb + 1.66
    H = yt + 0.24 + 0.42 + 0.04
    f = Fig("terrain", TPL_W, H)
    f.panel(0.02, 0.02, TPL_W - 0.04, H - 0.04, "TERRAIN AND HUMP HEIGHT",
            "(IMU task P5 and Safety task P1 · Buddy 4)", role="teal")
    teal, fill = ROLE["teal"][1], ROLE["teal"][0]

    # Composite normal state: every FRESH sample updates class and hump.
    f.frames.append(f._rounded(X0, yb, XW, yt - yb, "white", teal, lw=1.0, r=0.08,
                               z=1))
    f.label(X0 + 0.12, yt - 0.13, "TRACKING TERRAIN", fontweight="bold", color=teal)
    f.label(X0 + 1.28, yt - 0.13, "(each FRESH sample; an overrun counts as FRESH)")
    xl = X0 + XW - 0.56                           # PENDING self-transition
    f.arrow([(xl + 0.40, yt + 0.004), (xl + 0.40, yt + 0.14), (xl, yt + 0.14),
             (xl, yt + 0.004)], "PENDING ≤ 50 ms: wait (sample age kept)",
            (xl - 0.07, yt + 0.13), ha="right", color="hw", dashed=True)

    # Hump sequence (terrain.c): open, descent seen, close after flat dwell.
    y1 = yt - 0.29 - sh
    y2 = y1 - 0.35 - sh
    no = f.state(0.52, y1, 1.26, sh, "NO HUMP OPEN", "(watching for a rise)", "teal")
    op = f.state(2.82, y1, 1.24, sh, "HUMP OPEN", "(integrate height)", "teal")
    ds = f.state(2.82, y2, 1.24, sh, "DESCENT SEEN", "(keep peak until flat)", "teal")
    m1 = y1 + sh / 2
    f.start(0.33, m1)
    f.arrow([(0.39 + e, m1), (no[0] - e, m1)], color="hw")
    gx = (no[0] + no[2] + op[0]) / 2
    f.arrow([(op[0] - e, m1 + 0.07), (no[0] + no[2] + e, m1 + 0.07)],
            "voided + flat 0.2 s", (gx, m1 + 0.10), ha="center", va="bottom",
            color="hw", dashed=True)
    f.arrow([(no[0] + no[2] + e, m1 - 0.07), (op[0] - e, m1 - 0.07)],
            "pitch ≥ +0.08 rad,\nmoving forward", (gx, m1 - 0.10), ha="center",
            va="top", color="hw")
    xv = op[0] + 0.18
    f.arrow([(xv, y1 - e), (xv, y2 + sh + e)], "pitch < −0.035 rad,\nmoving forward",
            (xv + 0.05, (y1 + y2 + sh) / 2), color="hw")
    # Closing edge: its label sits in the pocket above it, clear of the boxes.
    m2 = y2 + sh / 2
    xn = no[0] + 0.43
    f.arrow([(ds[0] - e, m2), (xn, m2), (xn, y1 - e)],
            "|pitch| ≤ 0.035 rad for 0.2 s:\nhump closes; peak 5–300 mm,\n"
            "end within ±30 mm → run peak", (xn + 0.08, m2 + 0.035), va="bottom",
            color="hw")

    # Motion class (terrain_motion); tested in this order, first match wins.
    cx = op[0] + op[2] + 0.12
    cw = X0 + XW - 0.10 - cx
    ch = yt - yb - 0.14
    card = f._rounded(cx, yb + 0.07, cw, ch, fill, teal)
    f._boxed_text(card, cx + 0.10, yb + 0.07 + ch - 0.07,
                  "MOTION CLASS (first match wins)", ha="left", va="top",
                  fontweight="bold")
    f._boxed_text(card, cx + 0.10, yb + 0.07 + ch - 0.27, "\n".join(
        "• " + s for s in ("INVALID: unusable encoder data or input",
                           "COLLISION_CANDIDATE: < −3 m/s²",
                           "TURNING: |yaw rate| > 0.25 rad/s",
                           "STATIONARY: |speed| ≤ 5 mm/s",
                           "CLIMBING: pitch ≥ +0.08 rad",
                           "DESCENDING: pitch ≤ −0.08 rad",
                           "ACCELERATING / BRAKING: ±0.3 m/s²",
                           "STEADY: none of the above")),
                  ha="left", va="top", linespacing=1.3)

    # Faults (robot_app.c): soft stop recovers by local RESET, hard latches.
    f.state(X0, yf, 1.76, sh, "TILT_UNAVAILABLE", "(soft stop: brake)", "orange")
    hard = f.state(3.04, yf, 1.90, sh, "HARD FAULT", "(IMU, TILT or COLLISION)", "red")
    ym = (yf + sh + yb) / 2
    f.arrow([(0.40, yf + sh + e), (0.40, yb - 0.004)], "fresh valid tilt\n+ local RESET",
            (0.45, ym), color="hw", dashed=True)
    f.arrow([(1.30, yb - 0.004), (1.30, yf + sh + e)],
            "PENDING > 50 ms, unusable tilt\nor tilt older than 100 ms", (1.35, ym),
            color="fault")
    f.arrow([(3.18, yb - 0.004), (3.18, yf + sh + e)],
            "IMU: INVALID sample (I2C error, saturation, inconsistent status)\n"
            "TILT: valid |pitch| or |roll| > 0.60 rad, at any time\n"
            "COLLISION: encoder and accelerometer both < −3 m/s²",
            (3.23, ym), color="fault")
    fx = X0 + XW - 0.45
    f.final(fx, yf + sh / 2)
    f.arrow([(hard[0] + hard[2] + e, yf + sh / 2), (fx - 0.075 - e, yf + sh / 2)],
            "reboot", ((hard[0] + hard[2] + fx - 0.075) / 2, yf + sh / 2 + 0.03),
            ha="center", va="bottom", color="fault")

    # Notes and legend.
    kw = 2.50
    f.notes(X0, 0.08, XW - kw - 0.08, nh, [
        "Height = Σ Δs · sin(pitch): encoder path, filtered pitch, no gyroscope.",
        "Voided (not counted): reverse, turn, peak > 300 mm, bad tilt/encoder.",
        "Soft-stop, IMU and collision faults apply only during a run.",
        "Each START resets the hump and run peak (mount reference kept)."])
    f.keybox(X0 + XW - kw, 0.08, kw, nh, [
        ("start", None, "Initial state"), ("final", None, "Final state"),
        ("role", "teal", "Normal state"), ("role", "orange", "Soft stop"),
        ("role", "red", "Hard fault"), ("line", ("hw", False), "Transition"),
        ("line", ("fault", False), "Fault transition"),
        ("line", ("hw", True), "Conditional")], cols=2)
    f.save()


def fig_scan():
    w, h = 0.97, BOXH
    pad = 0.035                                   # frame around the active phases
    g = (SIDE_FIG - 0.02 - pad - 0.02 - 0.02 - 4 * w) / 3
    xs = [0.02 + i * (w + g) for i in range(4)]
    r2 = 0.04
    r1 = r2 + h + 0.40
    f = Fig("scan", SIDE_FIG, r1 + h + 0.31)
    f.frame(xs[1] - pad, r1 - pad, xs[3] + w - xs[1] + 2 * pad, h + 2 * pad)
    f.box(xs[0], r1, w, h, "IDLE\nawaits request")
    f.box(xs[1], r1, w, h, "POSITION\nset servo")
    f.box(xs[2], r1, w, h, "SETTLE\n180 ms → TRIG")
    f.box(xs[3], r1, w, h, "WAIT_ECHO\necho or 30 ms")
    f.box(xs[0], r2, w + g + w, h, "ABORTED, servo centred;\nmoved → hard SCAN fault",
          "fault")
    f.box(xs[2], r2, w + g + w, h, "COMPLETE, servo centred →\nVehicle: BYPASS or FAULT",
          "good")
    yc = r1 + h / 2
    e = 0.012
    for i in (1, 2):
        f.arrow([(xs[i] + w + e, yc), (xs[i + 1] - e, yc)])
    c0, c3 = xs[0] + w / 2, xs[3] + w / 2
    top = r1 + h
    xp = xs[1] + 0.20
    xn = xs[1] + w - 0.20                         # next beam enters right of the request
    # scanner_start refuses unless calibrated (servo confirmed) and stationary
    f.arrow([(c0, top + e), (c0, top + 0.09), (xp, top + 0.09), (xp, top + e)],
            "request, calibrated, stationary", (0.03, top + 0.115), ha="left",
            va="bottom")
    f.arrow([(c3, top + e), (c3, top + 0.12), (xn, top + 0.12), (xn, top + e)],
            "next beam (10° coarse, 2° fine)", ((xn + c3) / 2, top + 0.145),
            ha="center", va="bottom")
    ym = (r1 - pad + r2 + h) / 2
    f.arrow([(c3, r1 - e), (c3, r2 + h + e)], "last beam", (c3 - 0.05, ym), ha="right")
    f.arrow([(xs[1] + w / 2, r1 - pad - 0.004), (xs[1] + w / 2, r2 + h + e)],
            "moved (any active phase)\nor STOP / other fault", (xs[1] + w / 2 + 0.05, ym))
    f.save()


def tpl_scan():
    """Fig. 8 for the course template: scanner_phase_t (prefix SCANNER_) in
    one subsystem panel. The sweep phases (POSITION, SETTLE, WAIT_ECHO) abort
    on movement; the rest phases (IDLE, COMPLETE, ABORTED) ignore it and all
    three accept the next request (the scanner never returns to IDLE)."""
    e = 0.012
    ink, grey = FLOW["hw"], "#5b6f82"
    x0, x1 = 0.12, TPL_W - 0.12                  # phase-group span
    # heights, bottom up: notes/legend, rest group, gap, sweep group, header
    yb, hb = 0.10, 1.12
    rh = 0.70                                     # rest states
    ry = yb + hb + 0.14                           # rest group bottom
    sy = ry + 0.06
    top_r = sy + rh + 0.06                        # rest group top
    wy = top_r + 0.44                             # sweep group bottom
    sh = 0.58                                     # sweep states
    py = wy + 0.12                                # red arrows leave the group edge
    st = py + sh
    yi, yo = st + 0.12, st + 0.37                 # next-beam and fine-pass loops
    ftop = st + 0.56                              # sweep group top
    head = 0.46
    H = round(ftop + 0.10 + head + 0.02, 2)
    f = Fig("scan", TPL_W, H)
    f.panel(0.02, 0.02, TPL_W - 0.04, H - 0.04, "OBSTACLE HANDLING",
            "(Scan task P6 · Vehicle P4 · Buddy 5)", role="orange", head=head)
    # notes (left) and legend (right) along the bottom
    kw = 2.70
    f.notes(x0, yb, x1 - x0 - kw - 0.14, hb, [
        "Only Vehicle P4 requests a scan, once the car is at rest.",
        "Not calibrated yet: request refused, Vehicle faults at 20 s.",
        "Unknown beams never count as clear; any one blocks BYPASS.",
        "Own servo/TRIG error or 20 s limit: hard fault; servo stays put.",
        "Any rest phase accepts a request; IDLE is never re-entered."])
    f.keybox(x1 - kw, yb, kw, hb, [
        ("start", None, "Initial state"), ("line", ("hw", False), "Transition"),
        ("line", ("fault", False), "Abort / fault"), ("group", None, "Phase group"),
        ("role", "orange", "Normal state"), ("role", "green", "Sweep complete"),
        ("role", "red", "Aborted state")], cols=2)
    kx = x1 - kw + 0.10                           # dashed swatch of "Phase group"
    f.swatches.append(f._rounded(kx, yb + hb - 0.34 - 3 * 0.19 - 0.055, 0.26, 0.11,
                                 "none", grey, lw=0.9, ls=(0, (4, 2)), r=0.02, z=3))

    def group(y, h, name, sub):
        """Dashed phase group with its name block at the left."""
        f.frames.append(f._rounded(x0, y, x1 - x0, h, "none", grey, lw=0.9,
                                   ls=(0, (4, 2)), r=0.08, z=1))
        f.label(0.22, y + h / 2 + 0.17, name, fontweight="bold", color="#455a64",
                linespacing=1.15)
        f.label(0.22, y + h / 2 - 0.17, sub, color="#444444", linespacing=1.15)

    # rest phases (lower row)
    group(ry, top_r - ry, "REST\nPHASES", "(movement\nignored)")
    ix, iw = 1.24, 1.06
    ax_, aw = 2.56, 1.66
    cx_ = 4.46
    f.start(1.06, sy + rh / 2)
    f.arrow([(1.124, sy + rh / 2), (ix - e, sy + rh / 2)], color="hw", lw=1.1)
    f.card(ix, sy, iw, rh, "IDLE", "(after start-up:\nawaits a request)",
           role="orange")
    f.card(ax_, sy, aw, rh, "ABORTED", "(no profile to Vehicle,\nservo centred)",
           role="red")
    f.card(cx_, sy, x1 - 0.08 - cx_, rh, "COMPLETE", "(profile ready, servo centred)",
           ["then Vehicle: BYPASS or FAULT"], role="green")
    # sweep phases (upper row)
    w = 1.30
    g = (x1 - 0.08 - 1.22 - 3 * w) / 2
    xs = [1.22 + i * (w + g) for i in range(3)]
    group(wy, ftop - wy, "SWEEP\nPHASES", "(movement\naborts)")
    f.card(xs[0], py, w, sh, "POSITION", "(next beam angle\nchosen)", role="orange")
    f.card(xs[1], py, w, sh, "SETTLE", "(servo settling)", role="orange")
    f.card(xs[2], py, w, sh, "WAIT_ECHO", "(wait up to 30 ms;\nno echo = unknown)",
           role="orange")
    ym = py + sh / 2
    for i, (ev, act) in enumerate([("2 ms step", "/ set servo"),
                                   ("180 ms", "/ TRIG pulse")]):
        a, b = xs[i] + w + e, xs[i + 1] - e
        f.arrow([(a, ym), (b, ym)], color="hw", lw=1.1)
        f.label((a + b) / 2, ym + 0.04, ev, ha="center", va="bottom", color=ink)
        f.label((a + b) / 2, ym - 0.04, act, ha="center", va="top", color=ink)
    # coarse and fine sweep: loops from WAIT_ECHO back to POSITION
    xi0, xi1 = xs[0] + w - 0.30, xs[2] + 0.30
    xo0, xo1 = xs[0] + 0.25, xs[2] + w - 0.25
    f.arrow([(xi1, st + e), (xi1, yi), (xi0, yi), (xi0, st + e)],
            "echo or 30 ms [beams left] / next beam: coarse +10° (−60…+60°), fine +2°",
            ((xi0 + xi1) / 2, yi + 0.03), ha="center", va="bottom", color="hw", lw=1.1)
    f.arrow([(xo1, st + e), (xo1, yo), (xo0, yo), (xo0, st + e)],
            "coarse pass done [any beam ≤ 0.60 m or unknown] / fine pass: 2° steps "
            "across them ±10°", ((xo0 + xo1) / 2, yo + 0.03), ha="center",
            va="bottom", color="hw", lw=1.1)
    # between the groups
    mid = (top_r + wy) / 2
    xr = (ix + iw + ax_) / 2                      # from the group, not from IDLE
    f.arrow([(xr, top_r + 0.004), (xr, py - e)], "scan request\n[calibrated, stationary]",
            (xr - 0.05, mid), ha="right", color="hw", lw=1.1)
    xm = 2.86
    f.arrow([(xm, wy - 0.004), (xm, sy + rh + e)], "moved\n/ hard SCAN fault",
            (xm + 0.05, mid), color="fault", lw=1.1)
    xt = 4.10
    f.arrow([(xt, wy - 0.004), (xt, sy + rh + e)], "STOP or other\ntask's fault",
            (xt + 0.05, mid), color="fault", lw=1.1)
    xl = xs[2] + 0.18
    f.arrow([(xl, py - e), (xl, sy + rh + e)], "echo or 30 ms\n[sweep done]",
            (xl + 0.05, mid), color="hw", lw=1.1)
    f.save()


def diagrams():
    """Draw both figure sets: compact (five-page build) and template."""
    global TEMPLATE
    for template, folder in ((False, FIG), (True, FIG_T)):
        folder.mkdir(exist_ok=True)
        TEMPLATE = template
        try:
            for name in FIG_NAMES:
                # The template set is drawn by tpl_<name>() in the course
                # template's style where one exists, else by fig_<name>().
                draw = globals().get("tpl_" + name) if template else None
                (draw or globals()["fig_" + name])()
        finally:
            TEMPLATE = False


FIG_NAMES = ("hardware", "interaction", "system", "comm", "motion", "optical",
             "terrain", "scan")


def ensure_figures():
    """Draw the figures only when one is missing (clean checkout)."""
    if not all((d / (n + ".png")).exists() for n in FIG_NAMES for d in (FIG, FIG_T)):
        diagrams()


def fig_size(name, template=False):
    """Physical size (in) of a generated figure, from its PNG header."""
    data = ((FIG_T if template else FIG) / (name + ".png")).read_bytes()[16:24]
    return int.from_bytes(data[:4], "big") / 220, int.from_bytes(data[4:], "big") / 220


# ------------------------------------------------------------ docx helpers --
UNIT = re.compile(r"(?<=[\d×])\s(?=(?:mm|m|ms|s|Hz|kHz|B|KiB|rad|V|µs|%|°|m/s)"
                  r"(?![\w/]))")
APOSTROPHE = re.compile(r"(?<=\w)'(?=\w)")
_CURLY = [False]  # report(): typographic apostrophes (run's -> run’s)


def _runs(par, text, size=None, bold=False, italic=False):
    """Add text with **bold** segments to a paragraph; units stay with numbers."""
    text = UNIT.sub("\u00a0", text)
    if _CURLY[0]:
        text = APOSTROPHE.sub("\u2019", text)
    for i, seg in enumerate(re.split(r"\*\*", text)):
        if not seg:
            continue
        run = par.add_run()
        for k, part in enumerate(seg.split("\u2011")):   # Word non-breaking hyphen
            if k:
                run._r.append(OxmlElement("w:noBreakHyphen"))
            if part:
                run._r.add_t(part)
        run.bold = bold or (i % 2 == 1)
        run.italic = italic
        if size:
            run.font.size = Pt(size)
    return par


def p(d, text, style=None, size=None, keep=False, after=None, before=None):
    par = d.add_paragraph(style=style)
    _runs(par, text, size)
    fmt = par.paragraph_format
    if keep:
        fmt.keep_with_next = True
    if after is not None:
        fmt.space_after = Pt(after)
    if before is not None:
        fmt.space_before = Pt(before)
    return par


HEADS = []        # (bookmark, heading text, level) in document order
FIGURES = []      # (bookmark, figure name) in document order
_BOOKMARKS = [0]
FIG_TITLES = {"hardware": "Hardware integration and pin layout",
              "interaction": "RTOS task interaction diagram",
              "system": "System-level statechart (Vehicle)",
              "comm": "Communications statechart (Buddy 1)",
              "motion": "Motion control statechart (Buddy 2)",
              "optical": "Optical navigation statechart (Buddy 3)",
              "terrain": "Terrain and hump height statechart (Buddy 4)",
              "scan": "Obstacle handling statechart (Buddy 5)"}


def _bookmark(par, name):
    """Wrap a paragraph's content in a hidden bookmark (a contents target)."""
    _BOOKMARKS[0] += 1
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(_BOOKMARKS[0]))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(_BOOKMARKS[0]))
    ppr = par._p.pPr
    if ppr is None:
        par._p.insert(0, start)
    else:
        ppr.addnext(start)
    par._p.append(end)


def h(d, text, level=1, mark=True):
    head, sep, tail = text.partition(RTOS_TAIL)
    par = d.add_heading(head, level)
    if sep:                          # 6.x: "Related RTOS task(s): …" regular
        par.add_run(sep + tail).bold = False
    par.paragraph_format.keep_with_next = True
    if not mark:
        return par
    name = f"_Toc_w6_{len(HEADS) + 1:02d}"
    _bookmark(par, name)
    HEADS.append((name, text, level))
    return par


def _figure_mark(par, name):
    mark = f"_Fig_w6_{len(FIGURES) + 1}"
    _bookmark(par, mark)
    FIGURES.append((mark, name))


def cap(d, text, keep=False, before=None):
    return p(d, text, "Caption", keep=keep, before=before)


def fig_caption(par, caption):
    """Figure caption: bold navy "Fig. N." label, then the caption text."""
    m = re.fullmatch(r"(Fig\. \d+\.)(.*)", caption, re.S)
    if not m:
        raise ValueError(f'caption must start "Fig. N.": {caption!r}')
    label, rest = m.groups()
    run = par.add_run(label)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)
    _runs(par, rest)
    return par


# Children that follow w:pBdr, w:shd, w:tcMar and the w:tblPr elements in
# the OOXML schema order (Word can reject out-of-order properties that
# LibreOffice tolerates).
_AFTER_PBDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku",
               "w:wordWrap", "w:overflowPunct", "w:topLinePunct",
               "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
               "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
               "w:mirrorIndents", "w:suppressOverlap", "w:jc",
               "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
               "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
               "w:pPrChange")
_AFTER_SHD = ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
              "w:vAlign", "w:hideMark", "w:headers", "w:cellIns", "w:cellDel",
              "w:cellMerge", "w:tcPrChange")
_TBL_TAIL = ("w:tblCaption", "w:tblDescription", "w:tblPrChange")
_AFTER_TCMAR = ("w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark",
                "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge",
                "w:tcPrChange")


def _shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().insert_element_before(shd, *_AFTER_SHD)


def _rule_below(par, size, space, color):
    """Bottom border of a paragraph (a rule under a title or header)."""
    par._p.get_or_add_pPr().insert_element_before(parse_xml(
        f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{size}" '
        f'w:space="{space}" w:color="{color}"/></w:pBdr>'), *_AFTER_PBDR)


def _cell_margins(cell, left, right):
    """Left/right margins of one cell in twips (top/bottom from the table)."""
    cell._tc.get_or_add_tcPr().insert_element_before(parse_xml(
        f'<w:tcMar {nsdecls("w")}><w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/></w:tcMar>'), *_AFTER_TCMAR)


def _table_props(t, widths, borders=True, margins=(4, 43, 4, 43)):
    """Fixed layout, width, cell margins and borders, each written once in
    the schema order of w:tblPr (tblW, jc, tblBorders, tblLayout,
    tblCellMar, tblLook)."""
    tblPr = t._tbl.tblPr

    def put(el, *after):
        for old in tblPr.findall(el.tag):
            tblPr.remove(old)
        tblPr.insert_element_before(el, *after, *_TBL_TAIL)

    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(int(round(sum(widths) * 1440))))
    tw.set(qn("w:type"), "dxa")
    put(tw, "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
        "w:tblLayout", "w:tblCellMar", "w:tblLook")
    bd = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + side)
        if borders:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), GRID)
        else:
            el.set(qn("w:val"), "nil")
        bd.append(el)
    put(bd, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    put(layout, "w:tblCellMar", "w:tblLook")
    mar = OxmlElement("w:tblCellMar")
    for side, val in zip(("top", "left", "bottom", "right"), margins):
        el = OxmlElement("w:" + side)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    put(mar, "w:tblLook")
    grid = t._tbl.tblGrid
    for gc, wd in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(round(wd * 1440))))


def _segments(par, parts, size=None, bold=False):
    """Add text and bound identity fields (a list of str and Field)."""
    for part in parts:
        if isinstance(part, Field):
            par._p.append(part.element())
        else:
            _runs(par, part, size, bold=bold)
    return par


def _cell_text(cell, text, size, bold=False, color=None):
    cell.text = ""
    par = cell.paragraphs[0]
    if isinstance(text, (list, tuple)):          # text with identity fields
        _segments(par, text, size, bold)
        lines = []
    else:
        lines = text.split("\n")
    for k, line in enumerate(lines):
        if k:
            par = cell.add_paragraph()
        _runs(par, line, size, bold=bold)
    for par in cell.paragraphs:
        if color:
            for run in par.runs:
                run.font.color.rgb = RGBColor.from_string(color)
        fmt = par.paragraph_format
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.line_spacing = LEAD


def table(d, headers, rows, widths, size=TABLE_PT, keep=False, span_last=False):
    t = d.add_table(rows=0, cols=len(widths))
    t.autofit = False
    _table_props(t, widths)
    allrows = ([headers] if headers else []) + list(rows)
    for r, values in enumerate(allrows):
        row = t.add_row()
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        if headers and r == 0:
            trPr.append(OxmlElement("w:tblHeader"))
        cells = row.cells
        if span_last and r == len(allrows) - 1:
            merged = cells[0].merge(cells[-1])
            merged.width = Inches(sum(widths))
            _cell_text(merged, values[0], size)
            continue
        for c, wd, txt in zip(cells, widths, values):
            c.width = Inches(wd)
            head = headers is not None and r == 0
            _cell_text(c, txt, size, bold=head, color="FFFFFF" if head else None)
            if head:                 # template header row: navy, white bold
                _shade(c, NAVY)
        if keep and r < len(allrows) - 1:
            for c in cells:
                for par in c.paragraphs:
                    par.paragraph_format.keep_with_next = True
    return t


def add_picture(run, path, width):
    """Inline picture with explicit zero wrap distances: Word treats missing
    distances as 0, LibreOffice as 0.125 in (it would shift every figure
    9 pt right and clip the subsystem charts)."""
    shape = run.add_picture(str(path), width=Inches(width))
    for side in ("distT", "distB", "distL", "distR"):
        shape._inline.set(side, "0")
    return shape


def picture(d, name, caption):
    wd, _ = fig_size(name)
    par = d.add_paragraph()
    par.paragraph_format.space_after = Pt(0)
    par.paragraph_format.line_spacing = 1.0     # exact leading would clip it
    par.paragraph_format.keep_with_next = True
    add_picture(par.add_run(), FIG / (name + ".png"), wd)
    cp = fig_caption(d.add_paragraph(style="Caption"), caption)
    _figure_mark(cp, name)
    return cp


def subsystem(d, heading, name, caption, text):
    """Heading, then chart (left, full size) beside its field text (right)."""
    h(d, heading, 2)
    wl = SIDE_FIG
    widths = [wl, WIDTH - wl]
    t = d.add_table(rows=1, cols=2)
    t.autofit = False
    _table_props(t, widths, borders=False, margins=(0, 29, 0, 29))
    row = t.rows[0]
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    left, right = row.cells
    left.width, right.width = Inches(widths[0]), Inches(widths[1])
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    _cell_margins(left, 0, 0)
    _cell_margins(right, 115, 0)                  # 5.75 pt chart-to-text gap
    par = left.paragraphs[0]
    par.paragraph_format.space_after = Pt(0)
    par.paragraph_format.line_spacing = 1.0     # exact leading would clip it
    add_picture(par.add_run(), FIG / (name + ".png"), fig_size(name)[0])
    cp = fig_caption(left.add_paragraph(style="Caption"), caption)
    cp.paragraph_format.space_after = Pt(0)
    _figure_mark(cp, name)
    right.text = ""
    _runs(right.paragraphs[0], text, TABLE_PT)
    fmt = right.paragraphs[0].paragraph_format
    fmt.space_after = Pt(0)
    fmt.line_spacing = LEAD
    return t


DRAFT_NOTE = ("Generated by Tools/build_week6_submission.py; identities, AI evidence "
              "and team decisions pending. Not a record of student sign-off or "
              "instructor approval.")
FINAL_NOTE = ("Generated by Tools/build_week6_submission.py from the team's entries. "
              "Not a record of instructor approval.")


def _clean_package(d, title, comments=DRAFT_NOTE, date=DOC_DATE):
    """Neutral metadata, fixed dates, no unused template media (W6-15)."""
    cp = d.core_properties
    cp.title = title
    cp.subject = "INF2004 Week 6 design review"
    cp.author = "INF2004 robot-car team"
    cp.last_modified_by = "INF2004 robot-car team"
    cp.comments = comments
    cp.keywords = "INF2004, RTOS, Week 6, design review"
    cp.category = ""
    cp.revision = 1
    cp.created = date
    cp.modified = date
    cp.last_printed = date
    used = set(re.findall(r'r:embed="(rId\d+)"', d.element.xml))
    for rid, rel in list(d.part.rels.items()):
        if rel.reltype.endswith("/image") and rid not in used:
            d.part.drop_rel(rid)
    for part in d.part.package.iter_parts():
        if str(part.partname) == "/docProps/app.xml":
            part._blob = (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/'
                '2006/extended-properties" xmlns:vt="http://schemas.openxmlformats.org/'
                'officeDocument/2006/docPropsVTypes"><Template>Normal</Template>'
                '<Application>Microsoft Office Word</Application><DocSecurity>0'
                '</DocSecurity><ScaleCrop>false</ScaleCrop><Company></Company>'
                '<LinksUpToDate>false</LinksUpToDate><SharedDoc>false</SharedDoc>'
                '<HyperlinksChanged>false</HyperlinksChanged><AppVersion>16.0000'
                '</AppVersion></Properties>').encode("utf-8")


def save_docx(d, path, date=DOC_DATE):
    """Save with fixed ZIP entry dates so identical content gives identical bytes."""
    buf = io.BytesIO()
    d.save(buf)
    with zipfile.ZipFile(buf) as src, zipfile.ZipFile(path, "w") as dst:
        for item in src.infolist():
            info = zipfile.ZipInfo(item.filename, date.timetuple()[:6])
            info.create_system = 0       # as on Windows: same bytes on any OS
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = item.external_attr
            dst.writestr(info, src.read(item.filename))


def _field(par, instr, cached):
    """A simple Word field (PAGE, SECTIONPAGES) with its cached result."""
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = cached
    r.append(t)
    fld.append(r)
    par._p.append(fld)


def document(header, footer, body_pt=BODY_PT, of_total=False, report_look=False,
             title_page=False):
    """A cleared copy of the course template with the page setup and styles.

    report_look (the Week 6 report only; the buddy forms keep the plain
    look): template heading hierarchy (navy Heading 1, lighter blue
    Heading 2) at the old line heights, a grey running header with a rule
    and a centred grey page footer.  title_page (final build): no running
    header on the first body page, which carries the title.
    """
    d = Document(TEMPLATE_DOCX)
    body = d._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    s = d.sections[0]
    s.page_width, s.page_height = Inches(PAGE_W), Inches(PAGE_H)
    s.top_margin = s.bottom_margin = Inches(TOPBOT)
    s.left_margin = s.right_margin = Inches(SIDE)
    # Header 13 pt + 11.4 pt line + rule ends ≈ 4.5 pt above the body text.
    s.header_distance = s.footer_distance = Inches(0.18)
    for name, sz in [("Normal", body_pt), ("Title", 14), ("Heading 1", 11.5),
                     ("Heading 2", 10.5), ("Caption", TABLE_PT)]:
        st = d.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(sz)
        st.font.italic = False
        st.font.color.rgb = RGBColor.from_string("111111")
        fmt = st.paragraph_format
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(2)
        fmt.line_spacing = LEAD if sz <= TABLE_PT else 1.0
        fmt.widow_control = True
        rpr = st.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attr), "Calibri")
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if fonts.get(qn(attr)) is not None:
                del fonts.attrib[qn(attr)]
    for name in ["Heading 1", "Heading 2"]:
        st = d.styles[name]
        st.font.color.rgb = RGBColor.from_string(BLUE)
        st.font.bold = True
        st.paragraph_format.space_before = Pt(3 if name == "Heading 1" else 2)
        st.paragraph_format.space_after = Pt(1)
        st.paragraph_format.keep_with_next = True
    if report_look:
        # Same line heights (14.0/12.8 pt) and spacing (3+1/2+1 pt) as the
        # plain 11.5/10.5 pt single-spaced headings, so nothing reflows in
        # Word or LibreOffice (they treat space at a page top and between
        # headings differently, so the split must stay as it was).
        for name, size, color, line, before, after in (
                ("Heading 1", 13, NAVY, 14.0, 3, 1),
                ("Heading 2", 11, SKY, 12.8, 2, 1)):
            fmt = d.styles[name].paragraph_format
            d.styles[name].font.size = Pt(size)
            d.styles[name].font.color.rgb = RGBColor.from_string(color)
            fmt.line_spacing = Pt(line)
            fmt.space_before, fmt.space_after = Pt(before), Pt(after)
    d.styles["Title"].font.bold = True
    d.styles["Title"].font.color.rgb = RGBColor.from_string(BLUE)
    d.styles["Caption"].font.bold = False
    d.styles["Caption"].font.color.rgb = RGBColor.from_string("333333")
    for border in list(d.styles.element.xpath(".//w:pBdr")):
        border.getparent().remove(border)
    hp = s.header.paragraphs[0]
    hp.text = ""
    hp.style = "Caption"
    _runs(hp, header)

    def page_footer(fp):
        fp.text = ""
        fp.style = "Caption"
        _runs(fp, footer + "Page ")
        _field(fp, "PAGE", "1")
        if of_total:                 # body pages only: the front matter is a
            _runs(fp, " of ")        # separate, unnumbered section
            _field(fp, "SECTIONPAGES", str(MAX_PAGES))
        return fp

    page_footer(s.footer.paragraphs[0])
    if report_look:
        for name in ("Header", "Footer"):
            st = d.styles[name]       # based on Normal: Calibri
            st.font.size = Pt(TABLE_PT)
            st.font.color.rgb = RGBColor.from_string(HEAD_FOOT)
            st.paragraph_format.line_spacing = LEAD
            st.paragraph_format.space_before = st.paragraph_format.space_after = Pt(0)
        hp.style = "Header"
        _rule_below(hp, 4, 1, GRID)
        footers = [s.footer.paragraphs[0]]
        if title_page:               # first body page: the title, no header
            s.different_first_page_header_footer = True
            s.first_page_header.is_linked_to_previous = False
            footers.append(page_footer(s.first_page_footer.paragraphs[0]))
        for fp in footers:
            fp.style = "Footer"
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return d


# ------------------------------------------------------------ front matter --
STYLES_XML = {
    "PlaceholderText": (
        '<w:style w:type="character" w:styleId="PlaceholderText">'
        '<w:name w:val="Placeholder Text"/><w:basedOn w:val="DefaultParagraphFont"/>'
        '<w:uiPriority w:val="99"/><w:semiHidden/><w:rPr><w:color w:val="808080"/>'
        '</w:rPr></w:style>'),
    "TOC1": (
        '<w:style w:type="paragraph" w:styleId="TOC1"><w:name w:val="toc 1"/>'
        '<w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:uiPriority w:val="39"/>'
        '<w:unhideWhenUsed/><w:pPr><w:tabs><w:tab w:val="right" w:leader="dot" '
        f'w:pos="{int(FRONT_WIDTH * 1440)}"/></w:tabs><w:spacing w:before="100" '
        'w:after="20" w:line="276" w:lineRule="auto"/></w:pPr><w:rPr><w:b/>'
        '<w:sz w:val="23"/></w:rPr></w:style>'),
    "TOC2": (
        '<w:style w:type="paragraph" w:styleId="TOC2"><w:name w:val="toc 2"/>'
        '<w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:uiPriority w:val="39"/>'
        '<w:unhideWhenUsed/><w:pPr><w:tabs><w:tab w:val="right" w:leader="dot" '
        f'w:pos="{int(FRONT_WIDTH * 1440)}"/></w:tabs><w:spacing w:before="0" '
        'w:after="20" w:line="276" w:lineRule="auto"/><w:ind w:left="360" '
        'w:hanging="0"/></w:pPr><w:rPr><w:sz w:val="22"/></w:rPr></w:style>'),
}


def _styles(d):
    """Add the placeholder and contents styles once (Word built-in names)."""
    have = set(d.styles.element.xpath("./w:style/@w:styleId"))
    for sid, xml in STYLES_XML.items():
        if sid not in have:
            d.styles.element.append(parse_xml(xml.replace(
                "<w:style ", f"<w:style {nsdecls('w')} ", 1)))


def _run(par, text, size, bold=False, italic=False, font=None, color=None):
    run = par.add_run(text)
    run.font.size = Pt(size)
    run.bold, run.italic = bold, italic
    if font:
        run.font.name = font
        fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn("w:cs"), font)
        fonts.set(qn("w:eastAsia"), font)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def _front_par(d, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=0, line=1.0):
    par = d.add_paragraph()
    fmt = par.paragraph_format
    fmt.alignment = align
    fmt.space_before, fmt.space_after = Pt(before), Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = False
    return par


def _cover_cell(cell, part, bold, fill, color=None):
    """One centred 12 pt cover-table cell (text or field)."""
    cell.text = ""
    par = cell.paragraphs[0]
    fmt = par.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt.space_before = fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    if isinstance(part, Field):
        par._p.append(part.element())
    else:
        _run(par, part, 12, bold=bold, font=COVER_FONT, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill:
        _shade(cell, fill)


def cover(d):
    """Cover page in the team's SIT cover layout (logo, programme, module,
    title, team, names and student IDs as bound fields), in the body's
    typeface and colours: a 3.2 in logo as on SIT student covers, the title
    in navy on two lines, and the content spread over the page."""
    par = _front_par(d, before=48, after=48)
    add_picture(par.add_run(), LOGO, 3.2)
    for line in DEGREE:
        _run(_front_par(d, after=3), line, 12, bold=True, font=COVER_FONT)
    _run(_front_par(d, before=6), MODULE, 12, font=COVER_FONT)
    par = _front_par(d, before=72, after=72, line=1.05)
    for k, line in enumerate(COVER_LINES):
        if k:
            par.add_run().add_break()
        _run(par, line, 24, bold=True, font=COVER_FONT, color=NAVY)
    par = _front_par(d, after=10)
    _run(par, "Done by: INF2004 Team ", 12, bold=True, font=COVER_FONT)
    par._p.append(Field("group_id", bold=True, size=12, font=COVER_FONT).element())
    widths = [3.62, FRONT_WIDTH - 3.62]
    t = d.add_table(rows=6, cols=2)
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _table_props(t, widths, borders=False, margins=(0, 100, 0, 100))
    bd = t._tbl.tblPr.find(qn("w:tblBorders"))
    for el in bd:
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), COVER_RULE)
    for r, row in enumerate(t.rows):
        tr = row._tr.get_or_add_trPr()
        tr.append(OxmlElement("w:cantSplit"))
        height = OxmlElement("w:trHeight")
        height.set(qn("w:val"), str(int((0.36 if r == 0 else 0.34) * 1440)))
        height.set(qn("w:hRule"), "atLeast")
        tr.append(height)
        fill = COVER_HEAD if r == 0 else (COVER_ALT if r % 2 == 0 else None)
        cells = row.cells
        for c, wd in zip(cells, widths):
            c.width = Inches(wd)
        if r == 0:
            _cover_cell(cells[0], "Student Name", True, fill, "FFFFFF")
            _cover_cell(cells[1], "Student ID", True, fill, "FFFFFF")
        else:
            _cover_cell(cells[0], Field(f"b{r}_name", size=12, font=COVER_FONT),
                        False, fill)
            _cover_cell(cells[1], Field(f"b{r}_id", size=12, font=COVER_FONT),
                        False, fill)


def _toc_entry(d, style, anchor, text, page, first=False, last=False):
    """One contents line: hyperlinked text, dot leader, PAGEREF page number."""
    par = d.add_paragraph()
    par._p.style = style                     # built-in "toc 1" / "toc 2"

    def run(xml):
        par._p.append(parse_xml(f'<w:r {nsdecls("w")}>{xml}</w:r>'))

    if first:
        run('<w:fldChar w:fldCharType="begin"/>')
        run('<w:instrText xml:space="preserve"> TOC \\o "1-2" \\h \\z \\u </w:instrText>')
        run('<w:fldChar w:fldCharType="separate"/>')
    link = parse_xml(f'<w:hyperlink {nsdecls("w")} w:anchor="{anchor}" w:history="1">'
                     f'<w:r><w:t xml:space="preserve">{escape(text)}</w:t></w:r>'
                     '<w:r><w:tab/></w:r>'
                     '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
                     f'<w:r><w:instrText xml:space="preserve"> PAGEREF {anchor} \\h '
                     '</w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r>'
                     f'<w:r><w:t>{page}</w:t></w:r>'
                     '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:hyperlink>')
    par._p.append(link)
    if last:
        run('<w:fldChar w:fldCharType="end"/>')
    return par


def toc_text(text):
    """Contents entry for a heading: a subsystem heading keeps its name and
    buddy, not the task list (the body heading carries the full text)."""
    return text.split(RTOS_TAIL)[0]


def toc_map():
    """Body page of each heading and figure from the last --update-toc."""
    if TOC_PAGES.exists():
        return json.loads(TOC_PAGES.read_text(encoding="utf-8"))
    return {"headings": {}, "figures": {}}


def contents(d):
    """Contents page: a real Word TOC field (cached from toc_pages.json, so
    the DOCX opens correct without a field update) and a list of figures."""
    pages = toc_map()
    par = _front_par(d, align=WD_ALIGN_PARAGRAPH.LEFT, after=6)
    _run(par, "Table of Contents", 14, bold=True, color=NAVY)
    par.paragraph_format.page_break_before = True
    for k, (mark, text, level) in enumerate(HEADS):
        _toc_entry(d, "TOC1" if level == 1 else "TOC2", mark, toc_text(text),
                   pages["headings"].get(text, "?"), first=(k == 0),
                   last=(k == len(HEADS) - 1))
    par = _front_par(d, align=WD_ALIGN_PARAGRAPH.LEFT, before=14, after=4)
    _run(par, "List of Figures", 14, bold=True, color=NAVY)
    for k, (mark, name) in enumerate(FIGURES, 1):
        _toc_entry(d, "TOC2", mark, f"Fig. {k}. {FIG_TITLES[name]}",
                   pages["figures"].get(name, "?"))
    par = _front_par(d, align=WD_ALIGN_PARAGRAPH.LEFT, before=14)
    _run(par, "Page numbers refer to the numbered report body (pages 1–"
         f"{MAX_PAGES}); the cover and this page are not counted.", 10,
         italic=True, color="555555")


def front_matter(d):
    """Insert the cover and contents pages as an unnumbered first section
    (no header or footer) before the body, whose numbering starts at 1."""
    body = d._element.body
    first = body[0]
    count = len(body)                  # the last child is the body sectPr
    _styles(d)
    cover(d)
    contents(d)
    # Section break: the front matter's own page setup, no header/footer.
    brk = d.add_paragraph()
    brk.paragraph_format.space_after = Pt(0)
    sect = copy.deepcopy(body.sectPr)
    for tag in ("w:headerReference", "w:footerReference", "w:pgNumType", "w:titlePg"):
        for el in sect.findall(qn(tag)):
            sect.remove(el)
    mar = sect.find(qn("w:pgMar"))
    for side, val in (("top", FRONT_TOP), ("bottom", FRONT_BOTTOM),
                      ("left", FRONT_SIDE), ("right", FRONT_SIDE)):
        mar.set(qn("w:" + side), str(int(round(val * 1440))))
    brk._p.get_or_add_pPr().append(sect)
    added = list(body)[count - 1:-1]
    for el in added:
        first.addprevious(el)
    num = OxmlElement("w:pgNumType")
    num.set(qn("w:start"), "1")
    body.sectPr.find(qn("w:cols")).addprevious(num)


# ------------------------------------------------------------------ report --
def report(final=False):
    HEADS.clear()
    FIGURES.clear()
    _CURLY[0] = True
    _BOOKMARKS[0] = _SDT_IDS[0] = 0
    if final:
        d = document("INF2004 · Week 6 Design Review · Autonomous Robot Car", "",
                     of_total=True, report_look=True, title_page=True)
    else:
        d = document("INF2004 · Week 6 Design Review · PRIVATE TEAM-REVIEW DRAFT: "
                     "identities, AI evidence and team decisions pending",
                     "Draft; no instructor approval, student sign-off or hardware "
                     "result is claimed  |  ", of_total=True, report_look=True)
    # Template title block: centred bold navy with an accent rule under it
    # (the blank end of body page 1, left because Fig. 1 cannot split,
    # absorbs the extra height).
    title = p(d, "INF2004 Week 6 Design Review — Autonomous Robot Car", "Title",
              after=4)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.bold = True
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor.from_string(NAVY)
    _rule_below(title, 8, 2, RULE_BLUE)
    table(d, ["Submission detail", "Student entry"], [
        ("GROUP ID", [Field("group_id")]),
        ("Public GitHub repository", [Field("repo_url")])],
        [1.70, WIDTH - 1.70])
    h(d, "Team Role and Subsystem Ownership")
    # --final is run only after the team confirms the Part 2 section 3 roles.
    table(d, ["Buddy", "Name · Student ID", "Primary subsystem / responsibility"
              + ("" if final else " (Part 2 §3 proposal)")], [
        ("B1", [Field("b1_name"), " · ", Field("b1_id")],
         "Comms: Comm, Telemetry, radio; CI/build coordinator"),
        ("B2", [Field("b2_name"), " · ", Field("b2_id")],
         "Motion: Motion task, motors, encoders; pin-table keeper"),
        ("B3", [Field("b3_name"), " · ", Field("b3_id")],
         "Optical: Line task, barcode; Vehicle navigation co-owner"),
        ("B4", [Field("b4_name"), " · ", Field("b4_id")],
         "Terrain: IMU task, tilt, motion class, hump height"),
        ("B5", [Field("b5_name"), " · ", Field("b5_id")],
         "Obstacles: Scan, Vehicle; Safety/app integration")],
        [0.43, 3.55, WIDTH - 3.98])

    h(d, "1. Requirements")
    if final:
        limits = (("No value has been measured and none of the validation trials "
                   "V01–V15 has run. T-numbers are task-board items and Q-numbers are "
                   "questions to the lecturer (both private lists). " if LONG_TEXT else
                   "Nothing is measured yet and trials V01–V15 have not run. T-numbers "
                   "are task-board items and Q-numbers are lecturer questions (private "
                   "lists). ")
                  + "**Team acceptance limits:** ")
    else:
        limits = ("**Other numeric pass limits await team agreement**; configured "
                  "values are unmeasured; trials V01–V15 are proposals. Limits or "
                  "explicit deferral: ")
    cap(d, ("Sources: Write-up, Briefing, Barcode/Line spec and the Week 4 lecture. "
            "Briefing p8's Week 10 bands are used together with the zero-collision, "
            "clearance and no-false-line-lock criteria (FR8); how the bands count "
            "towards marks is not yet confirmed. " if LONG_TEXT else
            "Sources: Write-up, Briefing, Barcode/Line spec, Week 4 lecture. Briefing "
            "p8's Week 10 bands are used with the zero-collision, clearance and "
            "no-false-line-lock criteria (FR8); their effect on marks is unconfirmed. ")
           + limits
           + E("fr_limits", "agreed limits or explicit deferral").rstrip(".")
           + ".")
    h(d, "1.1 Functional Requirements", 2)
    table(d, ["ID · owner", "Requirement: the robot shall …", "Acceptance / verification"], [
        ("FR1 · B3", "follow the 18 mm black line autonomously, robust to minor "
         "lighting, track and speed variation (FOLLOW, Line)",
         "Course trials at agreed lighting and speed; line-loss count"),
        ("FR2 · B3", "decode barcodes (≈ 141 mm long, 20 mm from the line) into "
         "LEFT/RIGHT/STRAIGHT/U-TURN; ignore unknown codes",
         "Labelled official corpus; wrong actions 0; codebook pending lecturer"),
        ("FR3 · B3,\u00a0B2", "execute the decoded action at the next junction: drive the "
         "edge-to-axle approach, then pivot (NAVIGATE)",
         "Junction trials: axle over the crossbar at the pivot"),
        ("FR4 · B2", "control speed (≤ 0.15 m/s), distance and turns with encoder "
         "feedback and PID (Motion)", "Measured distance and turn error after "
         "encoder calibration"),
        ("FR5 · B4", "detect humps, estimate height and report the run's highest "
         "peak (estimator accepts 5–300 mm)", "Known-height humps; reported peak "
         "vs ground truth"),
        ("FR6 · B4", "classify motion (turn rate = encoder yaw) and tilt; detect "
         "collision; tilt > 0.60 rad or collision → FAULT (IMU, Safety)",
         "Orientation, turning (encoder yaw rate, no gyroscope) and "
         "injected-impact cases"),
        ("FR7 · B5", "detect obstacles on or near the line; profile width, nearest "
         "range and clear side (±60°: 10° coarse, 2° fine)",
         "Briefing p8 A–D: detection and direction 100/95/90/80 %; width "
         "error ≤ 10/20/30/40 %"),
        ("FR8 · B5, B2, B3", "bypass on a permitted side, minimise deviation, reacquire "
         "the original line; ‘reverse and reattempt’ (Write-up p9): not "
         "implemented; " + ("see 8.2 decision (T-51)" if final else
                            "implement-or-descope decision pending with B5 "
                            "(T-51); 8.2 only proposes deferral"),
         "Zero collisions, clearance, no false-line lock; Briefing p8: reacquire "
         "100/95/90/80 %, bypass < 3/5/7/10 s (≈ 31–36 s estimated, "
         "unmeasured); heading-error trial"),
        ("FR9 · B1", "publish MQTT telemetry: status 5 Hz (the liveness signal), "
         "perception and diagnostics 1 Hz, ≤ 256 B frames; ‘online’ after SUBACK, "
         "Last Will ‘offline’; reconnect",
         "Broker trace with outage and reconnect; drop counters"),
        ("FR10 · B1", "accept only whitelisted, non-retained commands; remote STOP "
         "only; START/RESET stay local" + ("" if final else " (proposal)"),
         "Whitelist and retained tests; lecturer policy answer"),
        ("FR11 · All", "brake and inhibit on STOP or a critical fault; never restart "
         "without an eligible RESET, then a separate START",
         "STOP, stale-data and reset regressions; measured motor response"),
        ("FR12 · All", "detect the course finish: rule unresolved, not implemented",
         "Obtain the rule, then model and test")],
        [0.98, 3.50, WIDTH - 4.48])
    h(d, "1.2 Non-Functional Requirements", 2)
    table(d, ["ID", "Requirement: the system shall …", "Verification"], [
        ("NFR1", "use C, the Pico C SDK 2.2.0 and micro T-Kernel, following the "
         "course template GitHub repo (Write-up p1 gives no URL; we build into the "
         "lecturer's port sirfonzie/mtk3smp-rp2040 at 2e8e6ad, to be confirmed as "
         "that template)" if LONG_TEXT else
         "use C, Pico C SDK 2.2.0 and micro T-Kernel, following the course template "
         "repo (Write-up p1, no URL; lecturer's port sirfonzie/mtk3smp-rp2040, to be "
         "confirmed)",
         "Pinned versions; fresh builds"),
        ("NFR2", "conform to the BARR-C:2018 coding standard; deviations justified",
         "Automated style subset + review"),
        ("NFR3", "use bounded memory: 8 × 4 KiB task, 4 KiB init, 6 KiB radio stacks, bounded buffers, one MQTT "
         "client", "Map file; stack high-water"),
        ("NFR4", "run Safety/Line/Scan at 2 ms, Motion/Comm 5, Vehicle/IMU 10, "
         "Telemetry 200 ms", "Target jitter, sample gaps (T-36)"),
        ("NFR5", "zero motor drive ≤ 10 ms after the STOP latch (V02); STOP polled "
         "every 2 ms, rechecked at motor commit",
         "Host regressions; measured latency, stopping distance"),
        ("NFR6", "report state, values in fixed integer units (mm, mm/s, mrad, ms), "
         "calibration version and first fault code in telemetry, and a start-up "
         "refusal reason on UART0", "Schema tests; start-up log"),
        ("NFR7", "keep radio load from starving control: RX 4×32 B, STOP 1×4 B, "
         "TX 4×256 B (enqueue overflow counted); ≤ 1 publish/100 ms",
         "Queue/drop tests; loaded trials"),
        ("NFR8", "keep documents, diagrams, source and evidence on one revision",
         "§9 software identity; CI; team sign\u2011off")],
        [0.52, 4.75, WIDTH - 5.27])

    h(d, "2. Hardware Integration Pin Layout")
    picture(d, "hardware", "Fig. 1. Proposed wiring by GPIO label (not header pins); "
            "arrows show signal direction. Dated pin approval is required before wiring.")
    table(d, ["Device / module", "Type", "Signal name", "Sensor label", "Pico GPIO",
              "Interface", "Notes · RTOS task owner"], [
        ("Encoders (Grove 3, 4)", "Sensor", "ENC_L/R_A/B", "A/B", "4/5, 16/17",
         "GPIO IRQ", "ISR counts → Motion (B2); Line reads"),
        ("Motor driver", "Actuator", "MOT_L/R_A/B", "M1A–M2B", "8/9, 10/11",
         "PWM 20 kHz", "Motion (B2); fault latch brakes"),
        ("Scan servo (port 1)", "Actuator", "SERVO", "SIG", "12", "PWM 50 Hz",
         "Scan (B5); 1–2 ms, uncalibrated"),
        ("HC-SR04 (servo 3/4)", "Sensor", "TRIG/ECHO", "TRIG/ECHO", "14/15",
         "GPIO out/IRQ", "Scan (B5); 5 V; ECHO divider"),
        ("GY-511 (Grove 2)", "Sensor", "IMU_SDA/SCL", "SDA/SCL", "2/3",
         "I2C1, 0x19 (expected)", "IMU (B4); 100 Hz ODR; no gyroscope; "
         + ("address to be confirmed by an I2C bus scan at first power-up"
            if LONG_TEXT else "address: I2C scan at first power-up")),
        ("IR line L (Grove 5)", "Sensor", "IR_LEFT", "A0", "26", "ADC0",
         "Line (B3); GP26 contact; D0 open"),
        ("IR line R (Grove 6)", "Sensor", "IR_RIGHT", "A0", "27", "ADC1",
         "Line (B3); GP27 contact only"),
        ("IR barcode (Grove 7)", "Sensor", "IR_BARCODE", "A0", "28", "ADC2",
         "Line (B3); VBAT jumper open"),
        ("Buttons (on board)", "Input", "STOP, START/RESET", "on board", "20/21",
         "GPIO in", "Safety/Vehicle (B5); active-low (A8)"),
        ("Console (UART)", "Comms", "TX/RX", "TX/RX", "0/1", "UART0 115k",
         "Start-up refusal diagnostics (B1)"),
        ("CYW43 (Pico W)", "Comms", "WL_*", "on module", "23–25, 29",
         "SPI", "Radio owner P9 (B1)")],
        [1.30, 0.58, 0.88, 0.80, 0.66, 0.86, WIDTH - 5.08])
    cap(d, ("Supplies (lecturer's wiring slides, 24 Sep): power bank on VIN; encoders "
            "and IR modules on 3.3 V. The team kit photos (26 Sep) show a 5 V/2.1 A "
            "power bank and the classic 5 V HC\u2011SR04. The HC\u2011SR04 therefore "
            "takes 5 V from servo port 3, and ECHO reaches GP15 through a 2.2k/3.3k "
            "divider (≈\u00a03.0\u00a0V). The lecturer's slide 4 powers it from 3.3 V, "
            "so the pin approval request asks the lecturer to confirm this change. "
            "Grove contacts are named by their silkscreen GPIO label. Each IR cable "
            "uses A0 only; D0 and the contact next to 3V3 stay unconnected. Briefing p6 "
            "calls these sensors Left IR and Center IR (line) and Right IR (barcode); "
            "here GP26, GP27, GP28. Continuity checks before the pin request: B2 Grove "
            "3–4, B3 Grove 5–7, B4 Grove 2. The Barcode/Line spec labels the "
            "barcode-to-line gap 20 mm, but its photo, scaled from the 18 mm line, "
            "shows ≈ 26–27 mm: measure the printed gap on the course sheet before "
            "mounting the barcode sensor (T-07)." if LONG_TEXT else
            "Supplies (lecturer's slides, 24 Sep): power bank on VIN; encoders and IR "
            "modules on 3.3 V. Kit photos (26 Sep): 5 V/2.1 A bank; classic 5 V "
            "HC\u2011SR04 fed 5 V from servo port 3 (slide 4: 3.3 V; lecturer to confirm), "
            "ECHO to GP15 via a 2.2k/3.3k divider (≈\u00a03.0\u00a0V). IR cables use "
            "A0 only (D0 unconnected). Briefing p6's Left, Center and Right IR are "
            "GP26, GP27, GP28. The spec labels the barcode gap 20 mm, but its photo, "
            "scaled from the 18 mm line, shows ≈ 26–27 mm: measure it before mounting "
            "the sensor (T-07)."), before=3)

    h(d, "3. RTOS Task Table")
    table(d, ["Task · owner", "Purpose", "Inputs", "Outputs", "Trigger", "Priority"], [
        ("Safety · B5 coord.", "Stop and supervise", "STOP, terrain, heartbeats",
         "Fault latch, BRAKE, watchdog", "2 ms", "P1 highest"),
        ("Motion · B2", "Leased motion → PWM", "Request, encoder deltas",
         "PWM, pose, completion", "5 ms", "P2"),
        ("Line · B3", "Line, barcode, junction", "ADC0–2, encoder totals",
         "Steering, code, junction edge", "2 ms", "P3"),
        ("Vehicle · B5 (B3 nav.)", "Mission supervisor", "Snapshots, GP21, motion",
         "Motion/scan requests, state", "10 ms", "P4"),
        ("IMU · B4", "Tilt, motion, hump", "I2C1 sample, Motion path",
         "Tilt, peaks, collision flag", "10 ms", "P5"),
        ("Scan · B5", "Sweep and ranging", "Echo times, scan request",
         "Servo, TRIG, range profile", "2 ms", "P6"),
        ("Comm · B1", "Commands, STOP first", "RX + STOP buffers",
         "Remote STOP; others counted", "5 ms", "P7"),
        ("Telemetry · B1", "JSON telemetry", "Shared snapshot",
         "TX: status 5 Hz, others 1 Hz", "200 ms", "P8"),
        ("Radio owner · B1", "lwIP, CYW43, MQTT", "Wi-Fi, TX buffer",
         "Publish; RX, STOP buffers", "≥ 2 ms", "P9 lowest")],
        [1.32, 1.40, 1.52, 1.78, 0.52, WIDTH - 6.54])
    p(d, "**Priority and ownership justification:** Safety P1 pre-empts all mission "
         "work. Motion P2 keeps PID timing and re-checks STOP at the motor commit. "
         "Line P3 samples faster than Vehicle P4 consumes. IMU P5 is below control "
         + ("because each sample busy-waits ≈ 1.2 ms on I2C (at most ≈ 8 ms: four "
            "transfer phases, each limited to 2 ms); this delays only Scan, Comm, "
            "Telemetry and the radio owner. Scan P6 can sit below IMU because the "
            "GPIO ISR times ECHO: a Scan step delayed by up to 8 ms still reads the "
            "same pulse, well inside the 150 ms range-age limit. " if LONG_TEXT else
            "(each I2C sample busy-waits ≈ 1.2 ms, ≈ 8 ms at worst). Scan P6 can wait "
            "behind IMU because the ISR times ECHO. ")
         + ("Comm P7 sits below the control and sensing tasks because the GP20 "
            "button, polled every 2 ms by Safety P1, is the primary STOP. Comm "
            "outranks Telemetry P8, so telemetry cannot delay the STOP dequeue. "
            if LONG_TEXT else
            "Comm P7 outranks Telemetry P8, so telemetry cannot delay the STOP "
            "dequeue. ")
         + ("A remote STOP still enters its buffer through the P9 radio owner, which "
            "P1–P8 all pre-empt. The watchdog "
            "is fed only while P2–P7 are fresh. Freshness checks use wrap-safe ages: "
            "a stamp that a higher-priority task publishes after a check starts "
            "counts as age 0, not stale. Telemetry and the radio owner are "
            "unsupervised (remote STOP best-effort). " if LONG_TEXT else
            "The watchdog is fed only while P2–P7 are fresh. Telemetry and the radio "
            "owner, which receives remote STOP, are unsupervised (STOP best-effort). ")
         + "Priorities do not bound "
         "interrupt or critical-section latency. Team acceptance: "
         + E("s3_acceptance", "team acceptance").rstrip(".") + ".", before=3)

    h(d, "4. Task Interaction Diagram")
    picture(d, "interaction", "Fig. 2. " + (
            "Nine RTOS tasks, the GPIO ISR, hardware, the shared-data store and the "
            "three radio message buffers. BRAKE is drawn from Safety; any task that "
            "latches a fault brakes the same way." if LONG_TEXT else
            "Nine RTOS tasks, GPIO ISR, hardware, shared data, radio buffers. BRAKE "
            "is drawn from Safety; every fault latch brakes alike."))
    p(d, "**Main communication paths:** Only Vehicle issues motion and scan requests; "
         "only Motion writes drive PWM (STOP and fault inhibit rechecked in its PRIMASK write); any task that "
         "latches a fault (robot_app_fault), and Vehicle while in FAULT, can only force "
         "BRAKE. "
         "The GPIO ISR only counts encoder edges and times ECHO: Motion "
         "consumes deltas, Line reads non-consuming totals, Scan consumes echo times, "
         "IMU uses Motion's path. " + (
         "Single core: the build enforces SMP=0 (make guard, #error in app_main.c), "
         "so core 1 never runs and short PRIMASK copies exclude every task and the "
         "ISR; no mutex, semaphore or event group. " if LONG_TEXT else
         "Shared data are copied in short PRIMASK sections (single core, SMP=0 "
         "enforced: no mutex, semaphore or event group). ") + "Message buffers "
         "isolate the radio: "
         "RX 4×32 B, STOP 1×4 B drained first by Comm, TX 4×256 B.")

    h(d, "5. System-Level Statechart")
    picture(d, "system", "Fig. 3. Vehicle mission states (vehicle_state_t names) with "
            "Safety precedence; guards are design intent, not measured behaviour.")
    if LONG_TEXT:                    # template box: three paragraphs
        states = (
            "START needs the start-up gates, a valid line and a fresh clear range. "
            "RANGE_WAIT halts for at most 1 s and waits for a fresh range reading "
            "taken at rest: obstacle → SCAN; clear range and a valid line → FOLLOW; "
            "non-clear range or the timeout → FAULT. NAVIGATE (at a later junction) "
            "drives the edge-to-axle approach, then pivots. STOP in any state, "
            "including INIT and IDLE, "
            "enters FAULT with zero drive.\n\n"
            "**Soft faults** allow an eligible local RESET (healthy, stationary, "
            "gates confirmed) → IDLE, then a separate START. They are local or "
            "remote STOP and tilt unavailable (unusable tilt, tilt older than "
            "100 ms, or an IMU sample PENDING over 50 ms). Every FAULT that the "
            "Vehicle supervisor enters is also soft, including a RANGE_WAIT, SCAN, "
            "REACQUIRE or pending-code timeout, REACQUIRE's 0.6 m drive ending "
            "without the line, a lost line, a lost or non-clear range, no qualified "
            "route, or a new code while one is pending. In RANGE_WAIT, a new code "
            "or a junction edge for a pending code also causes a FAULT.\n\n"
            "**Hard faults** latch until reboot and win over soft faults. They are: "
            "valid tilt over 0.60 rad, a collision, an encoder error, a wheel "
            "stall, an IMU INVALID sample during a run (I2C error, saturation or "
            "inconsistent status), a stale task owner, a motion fault (lease "
            "expiry, timeout, rejected request or sensor fault), an optical sensor "
            "fault, a second code before Vehicle takes the first, a scan abort or "
            "the scan task's 20 s limit, and a supervisor step failure. A "
            "motor-output fault is also hard once calibration is confirmed; before "
            "that, the motors are braked without a fault. An IMU overrun is not a "
            "fault. There is no finish state (FR12).")
    else:
        states = (
            "START needs start-up gates, a valid line and a fresh clear range. "
            "RANGE_WAIT halts ≤ 1 s for a fresh range taken at rest: obstacle → "
            "SCAN; clear range and line → FOLLOW; non-clear or timeout → FAULT. NAVIGATE "
            "(later junction) drives the edge-to-axle approach, then pivots. STOP in "
            "any state, INIT and IDLE too, enters FAULT with zero drive. **Soft "
            "faults** allow a local RESET (if healthy, stationary, gates "
            "confirmed) → IDLE, then a new START. They are local or remote STOP, "
            "unavailable tilt (unusable, > 100 ms old, IMU PENDING > 50 ms) and "
            "every Vehicle FAULT, e.g. RANGE_WAIT, SCAN, REACQUIRE (0.6 m or 10 s) "
            "or pending-code limits, lost line, lost or non-clear range, no route, "
            "a new code while one is pending, and in RANGE_WAIT any code or a "
            "pending code's junction edge. "
            "**Hard faults** latch until reboot and win over soft: valid tilt "
            "> 0.60 rad, collision, encoder, IMU INVALID during a run (I2C error, "
            "saturation, inconsistent status), stale task "
            "owner, wheel stall, motion lease expiry, timeout, rejected request or "
            "sensor fault, optical sensor, a second code before Vehicle takes the "
            "first, scan abort or scan-task 20 s limit, motor output once "
            "calibration is confirmed (before that, braked without a fault), "
            "supervisor step. No finish state (FR12).")
    p(d, "**Main states and assumptions:** " + states)

    h(d, "6. Subsystem Statecharts")
    cap(d, "One chart per buddy (Week 4 lecture 01:44:40); heading: Subsystem name, "
           "Related RTOS task(s); other fields beside the chart.")

    def fields(purpose, reason, states):
        return (f"**Purpose:** {purpose} **Reason a statechart is needed:** {reason} "
                f"**States, events and fault handling:** {states}")

    subsystem(d, "6.1 Communications — Buddy 1 · Related RTOS task(s): Comm P7, "
                 "Telemetry P8, radio owner P9", "comm",
              "Fig. 4. " + (
                  "Communications: radio setup, Wi-Fi join and MQTT session (SETUP "
                  "and UNAVAILABLE come from the radio owner's start-up loop; the "
                  "other modes come from net_mqtt.c flags)." if LONG_TEXT else
                  "Radio owner setup; Wi-Fi join, MQTT (net_mqtt.c flags)."), fields(
                  "non-blocking MQTT telemetry and command intake.",
                  "setup, join, session and retry need bounded recovery.",
                  ("the third failed radio setup → UNAVAILABLE until reboot. The MQTT "
                   "session counts as up only after the broker's SUBACK (CONNECTING → "
                   "SUBSCRIBED). Wi-Fi link loss → JOINING. A refused or lost session "
                   "retries with a 1 s to 30 s doubling backoff. " if LONG_TEXT else
                   "third failed setup → UNAVAILABLE until reboot; session up only "
                   "after SUBACK (CONNECTING → SUBSCRIBED); link loss → JOINING; "
                   "refused or lost session: 1 → 30 s doubling backoff. ")
                  + "QoS 0 and "
                  "the unsupervised radio owner make remote STOP best-effort; GP20 is "
                  "the primary STOP."))
    subsystem(d, "6.2 Motion Control — Buddy 2 · Related RTOS task(s): Motion P2 "
                 "(fault latches may force BRAKE)", "motion",
              "Fig. 5. Motion task states (motion_state_t).", fields(
                  "leased requests → PID wheel speed, odometry, completion.",
                  "renewal, completion and inhibition each decide differently "
                  "whether a request may keep driving the wheels or drive them "
                  "again." if LONG_TEXT else
                  "renewal, completion and inhibition differ in whether a request "
                  "may keep driving or drive again.",
                  "renewal extends the running ID's 50 ms lease; a completed ID never "
                  "replays; STOP, expiry, timeout or sensor fault → INHIBITED (hard "
                  "unless STOP). No acceleration ramp; " + (
                      "integral limited to ±0.35 duty and not integrated further "
                      "into saturation; " if LONG_TEXT else
                      "integral within ±0.35 duty, anti-windup; ")
                  + ("a wheel driven at 0.15 duty or more with no encoder count for "
                     "500 ms latches a stall fault (threshold unmeasured)." if LONG_TEXT
                     else "stall: duty ≥ 0.15, no encoder count for 500 ms "
                     "(unmeasured).")))
    subsystem(d, "6.3 Optical Navigation — Buddy 3 · Related RTOS task(s): Line P3, "
                 "Vehicle P4 (with B5), Motion P2", "optical",
              "Fig. 6. " + (
                  "Line following, pending barcode action and junction navigation "
                  "(line_ctrl, barcode and Vehicle states)." if LONG_TEXT else
                  "Line following, pending barcode action, junction navigation."),
              fields(
                  "follow the line, decode barcodes, act at the next junction.",
                  "a code must wait for a later, debounced junction edge.",
                  "rejected and unknown barcode frames are counted, with no action; a "
                  "frame still open when FOLLOW ends counts as rejected. A second code, "
                  "the 10 s pending limit or LOST → FAULT; REACQUIRE faults after 0.6 m "
                  "without the verified line or at 10 s. NAVIGATE drives 0.060 m "
                  "(unmeasured), then pivots (STRAIGHT drives the approach only). "
                  "Codebook empty until confirmed." if LONG_TEXT else
                  "rejected and unknown frames are only counted, as is a frame still "
                  "open when FOLLOW ends (rejected). A second code, 10 s pending limit "
                  "or LOST → FAULT; REACQUIRE faults after 0.6 m without the verified "
                  "line or at 10 s. NAVIGATE drives 0.060 m (unmeasured), then pivots "
                  "unless STRAIGHT. Codebook empty until confirmed."))
    subsystem(d, "6.4 Terrain and Hump Height — Buddy 4 · Related RTOS task(s): IMU P5, "
                 "Safety P1, Telemetry P8", "terrain",
              "Fig. 7. Hump phases, IMU sample status"
              + (" and motion classes." if LONG_TEXT else "."), fields(
                  "tilt, motion class, hump height, run peak.",
                  "fresh, pending and invalid samples recover differently.",
                  ("an overrun (a new sample overwrote an unread one) is FRESH: the "
                   "newest sample is used. A sample PENDING over 50 ms, or tilt that "
                   "is unusable or older than 100 ms, gives a soft stop. An INVALID "
                   "sample (I2C error, saturation or inconsistent status), a valid "
                   "tilt over 0.60 rad or a collision gives a hard fault. "
                   if LONG_TEXT else
                   "an overrun keeps the newest sample (FRESH). PENDING > 50 ms, or "
                   "unusable or > 100 ms-old tilt: soft stop. INVALID (I2C error, "
                   "saturation, inconsistent status), valid tilt > 0.60 rad or "
                   "collision: hard fault. ") +
                  "Collision's encoder deceleration is count‑quantised "
                  "(T‑53). Height = Σ Δs·sin(pitch), not Briefing p7's "
                  "h ≈ v²/2g (needs gravity-free vertical velocity; no "
                  "gyro to track tilt over the bump)" + (". " if LONG_TEXT else "; ") + (
                      "Briefing p7's IMU mounting near the centre of gravity is "
                      "assumption A5." if LONG_TEXT else
                      "IMU mounting: A5.")))
    subsystem(d, "6.5 Obstacle Handling — Buddy 5 · Related RTOS task(s): Scan P6, "
                 "Vehicle P4, Motion P2, Line P3", "scan",
              "Fig. 8. Scan task phases (scanner_phase_t).", fields(
                  "stationary profile, bypass, original-line recovery.",
                  "settle, echo wait and route permission need sequencing.",
                  "movement mid-scan → ABORTED (hard SCAN fault); no qualified "
                  "route or 20 s timeout → FAULT; "
                  "unknown echo is never clear. **Limitation:** REACQUIRE drives "
                  "≤ 0.6 m straight on the scan-entry heading, no lateral search: "
                  "≈ 1–2° heading or turn error can miss the 18 mm sensor pair (safe "
                  "FAULT; redesign T-52)."))

    h(d, "7. Assumptions")
    table(d, ["ID", "Assumption", "Risk if wrong", "Validation method", "Owner", "Status"], [
        ("A1", "Pico W radio and broker work in the lab; remote STOP with local "
         "START is allowed", "No telemetry; policy conflict",
         "Broker trial with outage (T-35); dated lecturer answer", "B1", "Open"),
        ("A2", "Two-channel encoders on the drive motors (else the kit's "
         "single-channel slot sensors, supported by a firmware build option); counts "
         "per turn, signs, 32.5 mm wheel radius, 140 mm track" if LONG_TEXT else
         "A/B encoders (else slot sensors); counts, signs, 32.5 mm radius, 140 mm "
         "track",
         "Wrong distance, turns, bar widths", "Hand-count edges, measure wheels (T-30)",
         "B2", "Open"),
        ("A3", "Three IR spots resolve the 18 mm line and narrowest bar at 0.08 m/s",
         "Lost line; rejected frames", "Spot photos, calibration, resolution study "
         "(T-06, T-32)", "B3", "Open"),
        ("A4", "Codebook, next-junction rule, 0.060 m approach (edge detection "
         "to axle over crossbar centre)",
         "Wrong or misplaced turn; a U-turn needs a junction, else FAULT",
         "Lecturer answer (Q1–Q3); measure edge-to-crossbar-centre travel",
         "B3, B2", "Open"),
        ("A5", "IMU flat near the centre of gravity (Briefing p7); Σ Δs·sin(pitch) "
         "gives height", "False tilt fault; wrong peak",
         "Six-face calibration, known humps, Q5 (T-33)", "B4", "Open"),
        ("A6", "Obstacle depth ≤ 0.30 m; swept corridor clear; straight reacquire "
         "meets line", "Collision; FAULT after missed line",
         "Footprint; both-side and heading-error trials (T-34)", "B5", "Open"),
        ("A7", "Task periods, 4 KiB stacks and 1.2 ms I2C reads fit the target",
         "Stale data; watchdog reset", "Jitter, sample gaps, stack high-water (T-36)",
         "All", "Open"),
        ("A8", "GP20/GP21 buttons read low when pressed (internal pull-ups)",
         "STOP never registers", "Bench: GP20 low only when pressed, before "
         "motor power; press test before every run", "B5", "Open")],
        [0.30, 2.52, 1.50, 2.00, 0.47, WIDTH - 6.79])

    h(d, "8. AI-Supported Design Evidence")
    h(d, "8.1 Pre-Submission AI Ideation", 2)
    if final:
        cap(d, "At the team coordinator's request, Claude Code ran one AI session "
               "per subsystem (Claude Opus 5.5, 25 Sep 2026), each with a prompt "
               "written for that subsystem. A second AI session checked every "
               "suggestion against the firmware and drafted the row; each buddy owns "
               "the decision in their row. The prompts, replies and checks are private "
               "team records (not in the public repository), available on request and "
               "at Tribunal 1. Adoptions are post-freeze work." if LONG_TEXT
               else "One AI session per subsystem, checked by a second (Claude Opus "
               "5.5, 25 Sep 2026); each buddy owns their row's decision.")
    else:
        cap(d, "Pending: each buddy enters their own AI session (no AI findings below). "
               "Suggested review topics only: B1 telemetry vs remote STOP; B2 averaged "
               "speed vs raw path; B3 later-junction trigger; B4 pending vs invalid "
               "samples; B5 unknown range, hidden depth.")
    rows = []
    for i in range(1, 6):
        rows.append((f"**B{i}:** " + E(f"ai_b{i}_prompt", "prompt / export"),
                     E(f"ai_b{i}_useful", "useful idea"),
                     E(f"ai_b{i}_weak", "weak idea"),
                     E(f"ai_b{i}_decision", "decision + reason")))
    table(d, ["AI prompt or discussion used", "Useful AI ideas",
              "Weak or unrealistic AI ideas", "Team decision"], rows,
          [1.80, 1.80, 1.80, WIDTH - 5.40])
    h(d, "8.2 Team-Reviewed Architecture Freeze", 2)
    table(d, ["AI or initial idea", "Team decision", "Engineering reason",
              "Affected artefact"], [
        ("One PWM owner (Motion), one requester (Vehicle)",
         E("freeze_1", "decision"), "No conflicting motor commands", "§§3–4, Fig. 5"),
        ("Unknown or stale range halts, never clear",
         E("freeze_2", "decision"), "Absent echo is not free space", "Figs. 3, 8; A6"),
        ("Averaged speed window, raw-count distance", E("freeze_3", "decision"),
         "Less quantisation, no lost path", "Figs. 5, 7"),
        ("Act at a later junction; reset needs a new START",
         E("freeze_4", "decision"), "No premature turn or restart", "Figs. 3, 6"),
        ("Defer ‘reverse and reattempt’ (Write-up p9, Q6)",
         E("freeze_5", "decision"), "No rear ranging to prove clear", "FR8; A6")],
        [1.72, 3.22, 1.58, WIDTH - 6.52])
    h(d, "8.3 Final AI-Vetted Gap Check", 2)
    audits = ("Rows 1–4 come from additional AI repository audits (21 Sep prompt "
              "“Perform a thorough final audit of our INF2004 project…”; 23 Sep)")
    if final:
        head = "**AI reviewer prompt used** (template, adapted): "
        tail = ("Rows 1–4: earlier AI repository audits (21–23 Sep; private "
                "records); rows 5–6 and the last row: the team's own run of this "
                "prompt (25 Sep). Team Response options: ")
    else:
        head = ("**AI reviewer prompt used** (template, adapted; the team must still "
                "run it on the frozen design): ")
        tail = (audits + ", not that prompt; responses await the team's choice. "
                "Team Response options: ")
    cap(d, head + "“Act as a strict embedded systems reviewer for our RTOS robot car "
           "design. Check the requirements, pin layout, RTOS task table, task "
           "interaction diagram, statecharts, assumptions and AI evidence. Identify "
           "missing requirements, unclear interfaces, untested assumptions, weak safety "
           "mitigations, timing risks, task conflicts and shared-resource problems. Do "
           "not rewrite the design; list only gaps that the team must review and "
           "justify.” " + tail + ", ".join(RESPONSES) + ".")

    def response(i, proposal):
        return E(f"gap{i}_response", "proposed: " + proposal, choices=RESPONSES)

    table(d, ["AI-identified gap", "Team response", "Action taken", "Evidence or reason"], [
        ("IMU overrun (two samples per 10 ms poll) latched a reboot-only fault",
         response(1, "Fixed"), "Overrun = newest FRESH sample; I2C error, "
         "saturation stay hard", "Host tests; ODR bench choice (B4)"),
        ("NAVIGATE pivoted before the axle reached the junction",
         response(2, "Partially fixed"), "Approach 0.060 m, then pivot",
         "Host tests; offset and turn trial open"),
        ("Newer shared stamp read as ≈ 49 days stale", response(3, "Fixed"),
         "Wrap-safe age (newer stamp = 0)", "Host regression scenarios"),
        ("Straight reacquire and bypass time vs Briefing p8 bands",
         response(4, "Deferred"), "Redesign (T-50 time, T-52 search) in "
         "Week 10; limitation in §6.5", "Estimate ≈ 31–36 s vs band D < 10 s"),
        ("Accelerometer-only tilt and collision: stops or humps may trip faults",
         response(5, "Needs validation"), (
             "A transient makes tilt invalid: a soft stop, not a hard fault (logged "
             "in T-67); collision windowing T-53" if LONG_TEXT else
             "Transients invalidate tilt: soft stop (T-67); collision windowing T-53"),
         (
             "Collision trips on one 10 ms sample; tilt on one sample of 80 "
             "ms-filtered pitch or roll; false-alarm hump, start and stop trials in T-33"
             if LONG_TEXT else "Collision and tilt each trip on one sample; "
             "trials T-33")),
        ("Fig. 3 drew range-loss exits only from FOLLOW and RANGE_WAIT; obstacle "
         "depth (A6) unmeasured",
         response(6, "Needs validation"), "NAVIGATE, REACQUIRE, BYPASS fault on lost "
         "range; A6 in T-34", "vehicle.c range checks; Q6"),
        (("**Team acceptance and own run:** " if final else
          "**Team acceptance of these responses, and gaps from the team's own run of "
          "the prompt (export reference):** ") + E("gap_acceptance",
                                                  "acceptance, own-run gaps, export"),)],
        [2.20, 1.25, 2.12, WIDTH - 5.57], span_last=True)

    h(d, "9. Design Consistency")
    table(d, ["Check", "Student response"], [
        ("Which task owns each hardware device?", "Motion: motor PWM (any fault latch "
         "can force BRAKE). ISR: encoder edges, ECHO. Line: IR ADC0–2. IMU: GY-511. "
         "Scan: servo, TRIG. Safety: STOP GP20, watchdog. Vehicle: GP21. Radio: CYW43. "
         "Initial task: UART0 console (start-up refusal only)."),
        ("Which task has the highest priority, and why?", "Safety P1: STOP, tilt, "
         "collision and stale-owner checks must pre-empt all mission work within 2 ms; "
         "the radio owner is lowest."),
        ("Which subsystem statechart corresponds to which RTOS task or subsystem?",
         "Fig. 3 Vehicle under Safety; Fig. 4 Comm, Telemetry, radio (B1); Fig. 5 Motion "
         "(B2); Fig. 6 Line, Vehicle navigation (B3); Fig. 7 IMU (B4); Fig. 8 Scan (B5)."),
        ("How are functional requirements reflected in the task and statechart design?",
         "FR1–3 Line, FOLLOW/NAVIGATE; FR4 Motion; FR5–6 IMU, Safety; FR7–8 Scan, "
         "SCAN/BYPASS/REACQUIRE; FR9–10 Comm, Telemetry; FR11 FAULT; FR12 open."),
        ("What changes are still expected after the Week 6 review?", "Codebook, finish, "
         "turn rules; measured calibration, junction approach, timing, stacks; 8.1 "
         "adoptions; " + (
             "reacquire and bypass-time redesign; the requirement change announced "
             "at Tribunal 1 (Week 8), implemented by Week 13 (T-40)." if LONG_TEXT else
             "reacquire/bypass redesign; the Tribunal 1 requirement change, due by "
             "Week 13 (T-40).")),
        ("Review record", E("review_record", "date, five reviewers, freeze, exceptions"))],
        [2.60, WIDTH - 2.60])
    cap(d, ("Supporting records (AI sessions, audits, test logs) are private team "
            "files, available on request; no instructor approval" if LONG_TEXT else
            "Supporting records are private team files, on request; no instructor "
            "approval")
           + (" is claimed. " if final else
              ", student sign-off or hardware result is claimed. ")
           + software_status(final), before=3)
    team_store(d)
    if COVER:
        front_matter(d)
    else:
        _styles(d)                   # grey placeholder style for blank fields
    return d


# ------------------------------------------------------------------- forms --
def buddy_forms():
    _CURLY[0] = False
    d = document("INF2004 · private buddy review form: uncompleted",
                 "Not a record of review, answers or sign-off  |  ", body_pt=10.5)
    for i in range(1, 6):
        source = (OUT / "Buddy_Reviews" / f"Buddy_{i}.md").read_text(encoding="utf-8")
        for block in source.split("\n\n"):
            block = block.strip()
            if not block:
                continue
            block = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", block)
            block = block.replace("`", "")
            if block.startswith("# "):
                par = p(d, block[2:], "Title")
                par.paragraph_format.page_break_before = i > 1
            elif block.startswith("## "):
                h(d, block[3:], 2, mark=False)
            elif re.match(r"(\d+\.|-) ", block):
                for ln in block.splitlines():
                    p(d, ln)
            else:
                p(d, " ".join(block.splitlines()))
    _clean_package(d, "INF2004 buddy review forms (uncompleted)", date=FORMS_DATE)
    save_docx(d, FORMS, date=FORMS_DATE)
    print(FORMS)


# ------------------------------------------------------------------- check --
HEADINGS = ["Submission detail", "GROUP ID", "Team Role and Subsystem Ownership",
            "1. Requirements", "1.1 Functional Requirements",
            "1.2 Non-Functional Requirements", "2. Hardware Integration Pin Layout",
            "3. RTOS Task Table", "Priority and ownership justification",
            "4. Task Interaction Diagram", "Main communication paths",
            "5. System-Level Statechart", "Main states and assumptions",
            "6. Subsystem Statecharts", "6.1 Communications", "6.2 Motion Control",
            "6.3 Optical Navigation", "6.4 Terrain and Hump Height",
            "6.5 Obstacle Handling", "Subsystem name", "Related RTOS task(s)",
            "Purpose:", "Reason a statechart is needed:",
            "States, events and fault handling:",
            "7. Assumptions", "Risk if wrong", "Validation method",
            "8. AI-Supported Design Evidence", "8.1 Pre-Submission AI Ideation",
            "Weak or unrealistic AI ideas", "8.2 Team-Reviewed Architecture Freeze",
            "8.3 Final AI-Vetted Gap Check", "AI reviewer prompt used", "Team Response",
            "9. Design Consistency", "Which task owns each hardware device?",
            "Which task has the highest priority, and why?",
            "Which subsystem statechart corresponds",
            "How are functional requirements reflected",
            "What changes are still expected after the Week 6 review?"]


HEADING_RGB = {int(NAVY, 16), int(SKY, 16)}


def _norm(text):
    return re.sub(r"\s+", " ", text).strip()


def front_pages(doc):
    """Front pages of a render: 2 when page 1 is the SIT cover, else 0."""
    return FRONT_PAGES if len(doc) and "BACHELOR OF ENGINEERING" in \
        _norm(doc[0].get_text()) else 0


def rendered_pages(doc):
    """Body page (1-based, after any front matter) of each heading and figure
    caption in a Word render: headings are the lines in the Heading 1/2
    colours that start in bold, captions the lines that start "Fig. N."."""
    heads, figs = {}, {}
    fp = front_pages(doc)
    for n, page in enumerate(doc, 1):
        body_page = n - fp
        if body_page < 1:
            continue
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                spans = [s for s in line["spans"] if s["text"].strip()]
                if not spans:
                    continue
                text = _norm("".join(s["text"] for s in line["spans"]))
                if spans[0]["flags"] & 16 and all(s["color"] in HEADING_RGB
                                                  for s in spans):
                    heads.setdefault(text, body_page)
                m = re.match(r"Fig\. (\d+)\. ", text)
                if m:
                    figs.setdefault(int(m.group(1)), body_page)
    return heads, figs


def toc_from_pdf(doc):
    """{"headings": {text: page}, "figures": {name: page}} for the current
    HEADS and FIGURES (report() must have run), or raise on a missing one."""
    heads, figs = rendered_pages(doc)
    out = {"headings": {}, "figures": {}}
    missing = []
    for _, text, _ in HEADS:
        page = next((pg for line, pg in heads.items()
                     if line.startswith(_norm(text)[:60])), None)
        if page is None:
            missing.append(text)
        out["headings"][text] = page
    for k, (_, name) in enumerate(FIGURES, 1):
        if k not in figs:
            missing.append(f"Fig. {k}")
        out["figures"][name] = figs.get(k)
    return out, missing


def check_pdf(path, final=False, max_pages=MAX_PAGES, blank_fields=False):
    """Print and return the number of violated page-contract checks."""
    import fitz  # PyMuPDF, only needed here
    doc = fitz.open(str(path))
    bad = []
    W, H = PAGE_W * 72, PAGE_H * 72
    sizes, text = [], []
    for n, page in enumerate(doc, 1):
        if abs(page.rect.width - W) > 2 or abs(page.rect.height - H) > 2:
            bad.append(f"page {n} is not A4: {page.rect}")
        body = fitz.Rect(SIDE * 72 - 1.5, 0.15 * 72, W - SIDE * 72 + 1.5, H - 0.15 * 72)
        for block in page.get_text("rawdict")["blocks"]:
            for line in block.get("lines", []):
                for span in line["spans"]:
                    chars = [c for c in span["chars"] if c["c"].strip()]
                    if not chars:
                        continue
                    word = "".join(c["c"] for c in span["chars"])[:40]
                    sizes.append(span["size"])
                    if span["size"] < MIN_PT:   # strict: Word writes 9.5 pt as 9.48
                        bad.append(f"page {n}: {span['size']:.2f} pt {word!r}")
                    # Trailing spaces may hang into the margin; glyphs may not.
                    if not all(body.contains(fitz.Rect(c["bbox"])) for c in chars):
                        bad.append(f"page {n}: text outside body {word!r}")
        for img in page.get_image_info():
            r = fitz.Rect(img["bbox"])
            if not body.contains(r):
                bad.append(f"page {n}: image outside body {r}")
        text.append(page.get_text())
    fp = front_pages(doc)                  # 2 with --cover, else 0
    front, rest = text[:fp], text[fp:]
    if len(doc) <= fp:
        bad.append(f"{len(doc)} pages: no body after the cover and contents")
    if len(rest) > max_pages:
        bad.append(f"{len(rest)} body pages > {max_pages}")
    toc_page = ""
    if fp:
        cover_text = _norm(front[0])
        for need in ("BACHELOR OF ENGINEERING", MODULE, COVER_TITLE, "Done by",
                     "Student Name", "Student ID"):
            if need not in cover_text:
                bad.append(f"cover lacks {need!r}")
        if len(doc[0].get_image_info()) != 1:
            bad.append("cover must hold exactly one image (the SIT logo)")
        toc_page = _norm(front[1]) if len(front) > 1 else ""
        if "Table of Contents" not in toc_page or "List of Figures" not in toc_page:
            bad.append("page 2 is not the contents page")
    elif rest and not _norm(rest[0]).lstrip().startswith("INF2004 Week 6 Design Review"):
        bad.append("page 1 does not start with the report title")
    flat = _norm(" ".join(rest))
    for head in HEADINGS:
        if head.lower() not in flat.lower():
            bad.append(f"missing heading/field: {head!r}")
    for n, page in enumerate(doc, 1):
        if n <= fp:
            continue
        foot = re.findall(r"Page (\d+) of (\d+)", page.get_text())
        if not foot or int(foot[-1][0]) != n - fp \
                or int(foot[-1][1]) != len(rest):
            bad.append(f"page {n}: footer is not 'Page {n - fp} of {len(rest)}'")
    widths = sorted(round(fitz.Rect(i["bbox"]).width / 72, 2)
                    for pg in list(doc)[fp:] for i in pg.get_image_info())
    expected = sorted(round(fig_size(nm)[0], 2) for nm in FIG_NAMES)
    if len(widths) != len(expected) or any(abs(a - b) > 0.02
                                           for a, b in zip(widths, expected)):
        bad.append(f"figure widths {widths} != generated {expected}")
    # Contents: the page numbers printed (cached in the DOCX from
    # toc_pages.json) must equal where Word actually put each heading.
    report(final=final)                       # HEADS / FIGURES of this build
    actual, missing = toc_from_pdf(doc)
    bad += [f"heading or figure not found in the body: {m!r}" for m in missing]
    printed = toc_map()
    for kind in ("headings", "figures") if fp else ():
        for key, page in actual[kind].items():
            if page is not None and printed[kind].get(key) != page:
                bad.append(f"contents gives {key!r} page {printed[kind].get(key)}, "
                           f"rendered on body page {page} (run --update-toc)")
    for _, text_h, _ in HEADS if fp else ():
        if _norm(toc_text(text_h)) not in toc_page:
            bad.append(f"contents lacks {toc_text(text_h)!r}")
    ident = flat[flat.find("Software identity"):] if "Software identity" in flat else ""
    if not ident:
        bad.append("software identity caption missing")
    blanks = sorted(set(re.findall(PLACEHOLDER + r"(B\d (?:name|student ID)|group ID|"
                                   r"public repository URL)", _norm(" ".join(text)))))
    if final and blanks:
        (print if blank_fields else bad.append)(
            "blank identity fields (fill them in Word before submitting): "
            + ", ".join(blanks))
    if final:
        for marker, why in (("[ENTER", "pending [ENTER] entries"),
                            ("DRAFT", "draft wording"),
                            ("FILL-TEST", "fill-test dummy text (capacity test only)")):
            if marker in flat:
                bad.append(f"final report still contains {why}")
        # Draft-only qualifiers ("(proposal)", "Part 2 §3 proposal"); a genuine
        # entry that uses the word must be reworded before submission.
        for m in re.finditer(r"\bproposal\b", flat, re.IGNORECASE):
            bad.append("final report still contains 'proposal': …"
                       + flat[max(0, m.start() - 40):m.end() + 20] + "…")
        if "PENDING" in ident.upper():
            bad.append("final report prints an unfinalised software identity")
        if not re.search(r"digest [0-9a-f]{12} ?… ?[0-9a-f]{6};", ident):
            bad.append("final report lacks a 64-hex build-input digest")
    print(f"{path}: {len(doc)} pages ({fp} front, {len(rest)} body); "
          f"min span {min(sizes):.2f} pt; {len(widths)} figures; {len(bad)} problems")
    for b in bad:
        print("  " + b)
    return len(bad)


# -------------------------------------------------------------------- main --
def main(argv=None):
    global ENTRIES, FILL, COVER
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--final", action="store_true",
                    help="omit draft wording; refuse while entries are missing")
    ap.add_argument("--entries", type=Path, help="JSON map of genuine team entries")
    ap.add_argument("--list-entries", action="store_true")
    ap.add_argument("--fill-test", type=Path, metavar="OUT.docx")
    ap.add_argument("--check-pdf", type=Path, metavar="PDF")
    ap.add_argument("--update-toc", type=Path, metavar="PDF",
                    help="record heading and figure pages of a Word render")
    ap.add_argument("--cover", action="store_true",
                    help="add the SIT cover and contents pages before the "
                         "template's five pages")
    ap.add_argument("--blank-fields", action="store_true",
                    help="--final/--check-pdf: identity fields may stay blank "
                         "(filled in Word before submission)")
    args = ap.parse_args(argv)
    COVER = args.cover
    if args.entries:
        ENTRIES = json.loads(args.entries.read_text(encoding="utf-8"))
        if not isinstance(ENTRIES, dict):
            raise SystemExit(f"{args.entries}: expected a JSON object of key: text")
        # Earlier key (Part 2 handovers, T-14): bN_identity = "name · student ID".
        for i in range(1, 6):
            legacy = str(ENTRIES.pop(f"b{i}_identity", "")).strip()
            if legacy and not (ENTRIES.get(f"b{i}_name") or ENTRIES.get(f"b{i}_id")):
                name, _, sid = legacy.rpartition("·") if "·" in legacy else \
                    legacy.rpartition(" ")
                ENTRIES[f"b{i}_name"], ENTRIES[f"b{i}_id"] = name.strip(), sid.strip()
    if args.check_pdf:
        return 1 if check_pdf(args.check_pdf, final=args.final,
                              blank_fields=args.blank_fields) else 0
    if args.update_toc:
        import fitz  # PyMuPDF, only needed here
        report(final=args.final)
        pages, missing = toc_from_pdf(fitz.open(str(args.update_toc)))
        if missing:
            raise SystemExit("--update-toc: not found in the render: " + ", ".join(missing))
        TOC_PAGES.write_text(json.dumps(pages, indent=2, ensure_ascii=False) + "\n",
                             encoding="utf-8", newline="\n")
        print(f"{TOC_PAGES}: {len(pages['headings'])} headings, "
              f"{len(pages['figures'])} figures; rebuild and render again")
        return 0
    ensure_figures()

    def pending(allow_blank_fields=False):
        return [k for k in USED if not str(ENTRIES.get(k, "")).strip()
                and not (allow_blank_fields and k in FIELDS)]

    if args.list_entries:
        report(final=False)
        for key, hint in USED.items():
            state = "pending" if key in pending() else "entered"
            opts = (" — one of: " + ", ".join(CHOICE[key])) if key in CHOICE else ""
            print(f"{key}\t{state}\t{hint}{opts}")
        return 0
    if args.fill_test:
        out = args.fill_test.resolve()
        if ROOT in out.parents:
            raise SystemExit("--fill-test output must be outside the repository")
        FILL = True
        d = report(final=args.final)
        _clean_package(d, "FILL TEST: dummy entries, not a report",
                       "Capacity test with dummy entries; never submit or commit.")
        out.parent.mkdir(parents=True, exist_ok=True)
        save_docx(d, out)
        print(f"{out} ({len(USED)} entries filled with dummy text)")
        return 0
    if args.final:
        # Refuse before writing anything: every entry and the identity file.
        report(final=False)
        todo = pending(args.blank_fields)
        problems = ["pending entries: " + ", ".join(todo)] if todo else []
        try:
            software_status(final=True)
        except SystemExit as exc:
            problems.append(str(exc).replace("--final refused: ", ""))
        if problems:
            raise SystemExit("--final refused; " + "; ".join(problems))
        USED.clear()
    diagrams()
    d = report(final=args.final)
    missing = pending()
    _clean_package(d, "INF2004 Week 6 Design Review",
                   FINAL_NOTE if args.final else DRAFT_NOTE)
    REPORT.parent.mkdir(exist_ok=True)
    save_docx(d, REPORT)
    print(REPORT)
    buddy_forms()
    named = [k for k in FIELDS if k.endswith(("_name", "_id")) and
             str(ENTRIES.get(k, "")).strip()]
    if named:
        # Names and student IDs: private repository only (team decision of
        # 25 September 2026), never a public repository or public copy.
        print(f"NOTE: {REPORT.relative_to(ROOT).as_posix()} holds names and "
              "student IDs (SHA-256 "
              f"{hashlib.sha256(REPORT.read_bytes()).hexdigest()}); commit it "
              "only to the private teck1341/embedded repository.")
    unknown = sorted(set(ENTRIES) - set(USED))
    if unknown:
        print("warning: unknown entry keys ignored: " + ", ".join(unknown))
    blank = [k for k in missing if k in FIELDS]
    print(f"{len(USED)} entries, {len(missing)} pending"
          + (f" ({len(blank)} identity fields to fill in Word)" if blank else ""))
    return 0


if __name__ == "__main__":
    sys.exit(main())
